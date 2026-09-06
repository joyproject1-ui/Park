# -*- coding: utf-8 -*-
"""Cpk 계산 파일(경향분석 Sheet, HLF-QC-126-08 한쪽 규격 / -09 양쪽 규격)을 openpyxl 만으로 그린다.

전년도 결재본의 .xls 서식을 Excel(COM)·LibreOffice 로 채우는 길이 모두 막혔을 때의 마지막 대비책 —
담당자 2026-09-06: "퀴노비드는 10 Lot 이상이라서 이 폴더에 Cpk 도 작성이 되었어야 하는데 안 됐네."
Excel 연결이 어떤 까닭으로든 터지면 파일이 아예 없게 되어 있었다. 이 길은 외부 프로그램 없이
서식의 칸·수식·꺾은선 그래프를 그대로 그린다(한림 서식 SS-QA-04/05 Ver_1.0 배치).

build(dst, values, cells, today) — cells: {"C4": 제품명, "C5": 공정, "K5": 시험항목, "N6": 하한, "P6": 상한 …}
  N6 와 P6 가 모두 있으면 양쪽 규격(Bilateral), 하나만 있으면 한쪽 규격(Unilateral).
"""
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FIRST_DATA_ROW, ROWS = 10, 35
YELLOW = PatternFill("solid", fgColor="FFFF00")
AMBER = PatternFill("solid", fgColor="FFCC00")
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="굴림", size=10)
BOLD = Font(name="굴림", size=10, bold=True)
BIG = Font(name="굴림", size=16)
BIG_BOLD = Font(name="굴림", size=16, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="bottom")
LEFT = Alignment(horizontal="left", vertical="bottom")


def _put(ws, ref, value, font=FONT, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    return c


def _blank(v):
    return v is None or v == "" or str(v).strip().upper() == "N/A"


def build(dst, values, cells, today=None):
    """dst(.xlsx) 에 Cpk 계산 파일을 만든다. 돌려주는 값은 'Bilateral' 또는 'Unilateral'."""
    cells = dict(cells or {})
    if today:
        cells["K4"] = today
    lo, hi = cells.get("N6"), cells.get("P6")
    if not _blank(lo) and not _blank(hi):
        kind = "Bilateral"
    else:
        kind = "Unilateral"
        if not _blank(lo) and _blank(cells.get("O6")):
            cells["O6"] = lo                           # 한쪽 규격 서식은 Min 이 O6 이다
        cells.pop("N6", None)
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    for col, width in (("A", 5.7), ("B", 15.7), ("C", 8.7), ("I", 12.7), ("J", 8.7), ("N", 8.7), ("Q", 10.7)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 43.5
    ws.merge_cells("E1:P1")
    if kind == "Bilateral":
        _put(ws, "E1", "제품품질평가 경향분석 Sheet (양쪽 규격 용)\nTrend Analysis Sheet for Product Quality Review (Bilateral specification)",
             BIG, align=CENTER)
    else:
        _put(ws, "E1", "제품품질평가 경향분석 Sheet (한쪽 규격 용)\nTrend Analysis Sheet for Product Quality Review (Unilateral specification)",
             BIG, align=CENTER)
    # 머리 — 제품명·공정·일자·시험항목·규격
    for a, b in (("A4", "B4"), ("C4", "G4"), ("I4", "J4"), ("K4", "L4"), ("A5", "B5"), ("C5", "G5"), ("I5", "J5"),
                 ("K5", "L5"), ("A6", "B6"), ("C6", "G6"), ("I6", "J6"), ("A8", "B8"), ("I11", "P12")):
        ws.merge_cells("%s:%s" % (a, b))
    _put(ws, "A4", "제품명 Product : ", BOLD, align=RIGHT)
    _put(ws, "C4", cells.get("C4") or "", fill=YELLOW, align=LEFT)
    _put(ws, "I4", "일자 Date : ", BOLD, align=RIGHT)
    _put(ws, "K4", cells.get("K4") or "", fill=YELLOW, align=LEFT)
    _put(ws, "A5", "공정명 Process : ", BOLD, align=RIGHT)
    _put(ws, "C5", cells.get("C5") or "", fill=YELLOW, align=LEFT)
    _put(ws, "I5", "시험항목 Test item : ", BOLD, align=RIGHT)
    _put(ws, "K5", cells.get("K5") or "", fill=YELLOW, align=LEFT)
    if kind == "Bilateral":
        ws.merge_cells("N4:P4")
        _put(ws, "N4", "Specification", BOLD, align=CENTER)
        for ref, text in (("N5", "Min"), ("O5", "Nominal(M)"), ("P5", "Max")):
            _put(ws, ref, text, align=CENTER, border=BOX)
        _put(ws, "N6", lo, fill=YELLOW, align=CENTER, border=BOX)
        _put(ws, "O6", cells.get("O6") if not _blank(cells.get("O6")) else "N/A", fill=YELLOW, align=CENTER, border=BOX)
        _put(ws, "P6", hi, fill=YELLOW, align=CENTER, border=BOX)
    else:
        ws.merge_cells("O4:P4")
        _put(ws, "O4", "Specification", BOLD, align=CENTER)
        for ref, text in (("O5", "Min"), ("P5", "Max")):
            _put(ws, ref, text, align=CENTER, border=BOX)
        _put(ws, "O6", cells.get("O6") if not _blank(cells.get("O6")) else None, fill=YELLOW, align=CENTER, border=BOX)
        _put(ws, "P6", hi if not _blank(hi) else None, fill=YELLOW, align=CENTER, border=BOX)
    # 통계
    _put(ws, "A8", "Data", align=CENTER, border=BOX)
    _put(ws, "F8", "Std Deviation (σ) :", align=RIGHT)
    _put(ws, "G8", "=STDEV(B%d:B%d)" % (FIRST_DATA_ROW, FIRST_DATA_ROW + ROWS - 1), align=CENTER)
    _put(ws, "A9", "No.", align=CENTER, border=BOX)
    _put(ws, "B9", "결과값 Result", align=CENTER, border=BOX)
    _put(ws, "F9", "Mean (X) :", align=RIGHT)
    _put(ws, "G9", "=AVERAGE(B%d:B%d)" % (FIRST_DATA_ROW, FIRST_DATA_ROW + ROWS - 1), align=CENTER)
    if kind == "Bilateral":
        _put(ws, "O8", "Short-term Capability (Cp) :", align=RIGHT)
        _put(ws, "P8", "=(P6-N6)/(6*G8)", align=CENTER)
        _put(ws, "O9", "Short-term Centered Capability - including Global Variation (Cpk) :", fill=AMBER, align=RIGHT)
        _put(ws, "P9", "=MIN((P6-G9)/(3*G8),(G9-N6)/(3*G8))", fill=AMBER, align=CENTER)
        _put(ws, "F10", "Upper Control Limit (UCL) :", align=RIGHT)
        _put(ws, "G10", "=G9+(3*G8)", align=CENTER)
        _put(ws, "F11", "Lower Control Limit (LCL) :", align=RIGHT)
        _put(ws, "G11", "=G9-(3*G8)", align=CENTER)
        heads = ("Mean", "UCL", "LCL", "USL", "LSL")
    else:
        _put(ws, "O8", '=IF(COUNTBLANK(O6:P6)=2,"Define Min or Max Specification :",IF(COUNTBLANK(O6)=1,'
                       '"Short-term Capability (Cpu) :",IF(COUNTBLANK(P6)=1,"Short-term Capability (Cpl) :",'
                       '"Define Min or Max Specification :")))', align=RIGHT)
        _put(ws, "P8", '=IF(COUNTBLANK(O6:P6)=2,"",IF(COUNTBLANK(P6)=1,(G9-O6)/(3*G8),IF(COUNTBLANK(O6)=1,'
                       '(P6-G9)/(3*G8),"")))', align=CENTER)
        _put(ws, "O9", "Short-term Centered Capability - including Global Variation (Cpk) :", fill=AMBER, align=RIGHT)
        _put(ws, "P9", "=P8", fill=AMBER, align=CENTER)
        _put(ws, "F10", '=IF(COUNTBLANK(O6:P6)=2,"Define Min or Max Specification :",IF(COUNTBLANK(O6)=1,'
                        '"Upper Control Limit (UCL) :",IF(COUNTBLANK(P6)=1,"Lower Control Limit (LCL) :",'
                        '"Define Min or Max Specification :")))', align=RIGHT)
        _put(ws, "G10", '=IF(COUNTBLANK(O6:P6)=2,"",IF(COUNTBLANK(P6)=1,G9-(3*G8),G9+(3*G8)))', align=CENTER)
        heads = ("Mean",
                 '=IF(COUNTBLANK(O6:P6)=2,"",IF(COUNTBLANK(O6)=1,"UCL",IF(COUNTBLANK(P6)=1,"LCL","")))',
                 '=IF(COUNTBLANK(O6:P6)=2,"",IF(COUNTBLANK(O6)=1,"USL",IF(COUNTBLANK(P6)=1,"LSL","")))')
    _put(ws, "I11", '=IF(P9>=1,"공정 능력 충분 Sufficient Process Capability","공정 능력 부족 Insufficient Process Capability")',
         BIG_BOLD, align=CENTER)
    # 그래프가 쓰는 계산 열 R.. — 서식과 같은 자리
    for k, head in enumerate(heads):
        _put(ws, "%s9" % get_column_letter(18 + k), head, align=CENTER)
    for i in range(ROWS):
        r = FIRST_DATA_ROW + i
        _put(ws, "A%d" % r, i + 1, fill=YELLOW, align=CENTER, border=BOX)
        v = values[i] if i < len(values) else None
        _put(ws, "B%d" % r, v, fill=YELLOW, align=CENTER, border=BOX)
        if v is None:
            ws["B%d" % r].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN, diagonal=THIN, diagonalUp=True)
        if kind == "Bilateral":
            forms = ("=G9", "=G10", "=G11", "=P$6", "=N$6") if i == 0 else \
                    ("=R%d" % (r - 1), "=S%d" % (r - 1), "=T%d" % (r - 1), "=P$6", "=N$6")
        else:
            forms = ("=G9", "=G10", "=MAX(O$6:P$6)") if i == 0 else \
                    ("=R%d" % (r - 1), "=S%d" % (r - 1), "=MAX(O$6:P$6)")
        for k, f in enumerate(forms):
            ws["%s%d" % (get_column_letter(18 + k), r)] = f
    for ref in ("G8", "G9", "G10", "G11", "P8", "P9"):
        ws[ref].number_format = "0.00"
    ws.merge_cells("A45:P45")
    _put(ws, "A45", "제품품질평가 경향분석 Sheet(%s 규격 용)(%s(Ver_1.0))%sHLF-QC-126-%s/Rev.000"
         % (("양쪽", "SS-QA-05", " " * 120, "09") if kind == "Bilateral" else ("한쪽", "SS-QA-04", " " * 120, "08")),
         align=LEFT)
    # 꺾은선 그래프 — 결과값과 평균·관리 한계·규격선. 자료 줄 수만큼만 잡는다.
    last = FIRST_DATA_ROW + max(len(values), 1) - 1
    chart = LineChart()
    chart.width, chart.height = 15, 7.5
    chart.legend.position = "r"
    colors = ("000080", "ff6600", "969696", "ff9900", "666699", "339966") if kind == "Bilateral" \
        else ("000080", "008000", "ff9900", "ff0000")
    cols = [2] + [18 + k for k in range(len(heads))]
    for k, col in enumerate(cols):
        ref = Reference(ws, min_col=col, min_row=9, max_row=last)
        chart.add_data(ref, titles_from_data=True)
    for k, s in enumerate(chart.series):
        s.graphicalProperties.line.solidFill = colors[k % len(colors)]
        s.marker = Marker(symbol="diamond" if k == 0 else "none")
        s.smooth = False
    chart.set_categories(Reference(ws, min_col=1, min_row=FIRST_DATA_ROW, max_row=last))
    ws.add_chart(chart, "C13")
    ws.print_area = "A1:P45"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(dst)
    return kind
