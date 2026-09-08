# -*- coding: utf-8 -*-
"""안연고제 결재본(퀴노비드안연고 계열)에 값을 채우는 절차.

퀴노비드안연고 2026 PQR 을 담당자와 함께 만들며 확정한 채움 규칙을 그대로 옮겼다.
표는 항 제목과 머리행 낱말로 찾고, 값은 ProductData(판독 결과)에서 가져온다.
판독하지 못한 항(손글씨 안정성 등)은 결재본 값을 두고 issues 에 '확인 필요' 로 남긴다.
"""
import copy
import datetime as _dt
import os
import re
import statistics

from docx.oxml.ns import qn

from . import carry as CARRY
from . import detail92 as D
from . import docedit as E
from . import lotcode, qc
from .readers import masters as masters_mod
from .locate import find_tables, find_para, _text, outline
from .ooxml_order import get_or_add

STORE = ("25±2°C,", "60±5%RH")


def _num(s):
    m = re.search(r"-?\d+(?:\.\d+)?", str(s or ""))
    return float(m.group()) if m else None


def _rng(s):
    m = re.findall(r"\d+(?:\.\d+)?", str(s or ""))
    return (float(m[0]), float(m[1])) if len(m) >= 2 else None


def avg(vals, digits=2):
    return ("%%.%df" % digits) % (sum(vals) / len(vals))


DEFAULT_LIMITS = {"particle": 75.0, "assay": (90.0, 110.0), "metal": 50.0}   # 퀴노비드안연고 값 — 9.1 을 못 읽을 때만


def parse_limits(spec_texts):
    """9.1 시험결과표의 허용기준 글에서 Cpk 한계를 읽는다. {'particle': usl, 'assay': (lsl, usl), 'metal': usl}

    spec_texts: {'particle': '75 ㎛ 이하', 'assay': '90.0 ~ 110.0%', 'metal': '… 합계 50개 이하 …'}.
    못 읽은 항목은 넣지 않는다 — 부르는 쪽이 기본값을 쓰고 문의 목록에 남긴다.
    """
    out = {}
    t = spec_texts.get("particle") or ""
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:㎛|um|µm|μm)", t)
    if m:
        out["particle"] = float(m.group(1))
    t = spec_texts.get("assay") or ""
    m = re.search(r"(\d+(?:\.\d+)?)\s*[~∼～-]\s*(\d+(?:\.\d+)?)\s*%", t)
    if m:
        out["assay"] = (float(m.group(1)), float(m.group(2)))
    t = spec_texts.get("metal") or ""
    m = re.search(r"(\d+)\s*개\s*이하", t)
    if m:
        out["metal"] = float(m.group(1))
    return out


def cpk_uni(vals, usl):
    if usl is None or not vals:          # 허용기준을 못 읽었으면 Cpk 를 내지 않는다 (담당자 2026-09-08)
        return None
    m = sum(vals) / len(vals)
    s = statistics.stdev(vals) if len(vals) > 1 else 0
    return (usl - m) / (3 * s) if s else None


def cpk_bi(vals, lsl, usl):
    if lsl is None or usl is None or not vals:
        return None
    m = sum(vals) / len(vals)
    s = statistics.stdev(vals) if len(vals) > 1 else 0
    return min((usl - m) / (3 * s), (m - lsl) / (3 * s)) if s else None


def _tables(document, prefix):
    return [document.tables[i] for i in find_tables(document, prefix)]


def _note_numbers_under(document, prefix):
    """prefix 항(9.2 …) 안에 이미 있는 각주 번호들 — 양식의 '1) 모든 Lot 의 …' 같은 것."""
    numbers = set()
    body = document.element.body
    inside = False
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
        if not text:
            continue
        m = re.match(r"^(\d{1,2}(?:\.\d+)*)[.\s]", text)
        if m and len(text) < 80:
            inside = m.group(1).startswith(prefix)
            continue
        if inside:
            n = re.match(r"^(\d)\)", text)
            if n:
                numbers.add(int(n.group(1)))
    return numbers


def _mark_header(table, word, number):
    """머리행 칸(예: '개개')의 글 끝에 각주 번호 'n)' 를 붙인다 — 뒤에서 윗첨자로 바뀐다."""
    for row in table.rows[:3]:
        for cell in row.cells:
            if re.sub(r"\s+", "", cell.text) == word:
                paragraphs = [p for p in cell.paragraphs if p.text.strip()]
                if not paragraphs or not paragraphs[-1].runs:
                    return False
                paragraphs[-1].runs[-1].text += "%d)" % number
                return True
    return False


def _by_header(tables, *words):
    for t in tables:
        head = re.sub(r"\s+", "", _text(t._tbl.findall(qn("w:tr"))[0]))
        if all(w in head for w in words):
            return t
    return None


ASSAY_PART = re.compile(r"\s*\[([^\]]+)\]")


def assay_component(crit_text):
    """9.1 허용기준 칸 앞머리의 [성분] 이름. 없으면 None.

    주성분이 둘 이상인 제품(디겐타안연고: 플루오로메톨론·겐타마이신황산염)은 함량 줄이
    성분별로 나뉘고, 어느 줄이 어느 성분인지는 이 표시로만 알 수 있다.
    """
    m = ASSAY_PART.match(crit_text or "")
    return m.group(1).strip() if m else None


def _avg_text(values, fmt="%.2f", unit_on_avg=True):
    """질량·용량 평균 칸의 글 — 값이 모두 같으면 평균·범위를 적지 않고 그 값만 적는다.

    2026 결재본의 주석 그대로다: "모든 Lot 의 시험결과가 (…) 동일하여 최댓값, 최솟값, 평균은
    별도로 작성하지 않음." 3 Lot 이 모두 4.1 g 인데 'Av. 4.1 (4.1 ~ 4.1g)' 은 읽기 나쁘다.
    """
    if len(set(values)) == 1:
        return (fmt + " g") % values[0]
    head = "Av. " + fmt + ("g" if unit_on_avg else "")
    return (head + "\n(" + fmt + " ~ " + fmt + "g)") % (
        sum(values) / len(values), min(values), max(values))


# 안연고 라인에서 디겐타안연고만 건열멸균기를 더 쓴다 — 다른 연고 보고서에서는 뺀다
# (담당자 2026-09). 설비명·관리번호 어느 쪽으로도 알아본다.
DRY_HEAT = re.compile(r"건열\s*멸균|DAE5030")


def drop_equipment(table, pattern, keep=(), name=""):
    """설비 표에서 이 제품이 쓰지 않는 설비 줄을 지운다. 지운 대수를 돌려준다.

    설비 한 대가 두 줄(문서번호 줄 + 완료일 줄)이라 두 줄을 함께 지운다. keep 에 든 제품
    이름이면 그대로 둔다.
    """
    if any(k and k in (name or "") for k in keep):
        return 0
    trs = table._tbl.findall(qn("w:tr"))
    hit = [i for i, tr in enumerate(trs) if pattern.search(_text(tr))]
    if not hit:
        return 0
    # 설비 한 대가 여러 줄(문서번호 줄 + 완료일 줄)이라 연번이 같은 줄을 함께 지운다
    drop, machines = set(), set()
    for i in hit:
        drop.add(i)
        no = _text(trs[i].findall(qn("w:tc"))[0]).strip()
        machines.add(no or i)
        for j in range(i + 1, len(trs)):
            if _text(trs[j].findall(qn("w:tc"))[0]).strip() not in ("", no):
                break
            drop.add(j)
    for i in sorted(drop, reverse=True):
        trs[i].getparent().remove(trs[i])
    return len(machines)


def _deviation_lines(dev):
    """11항 일탈사항 칸의 글 — 2026 결재본 차림새로 요점만.

        [조제 작업 중 장비 이상 건]
        * 일탈 내용
        조제 작업 진행 중 MAINMIXER 의 스크래퍼가 작동되지 않아 조제 작업 중단함.
        * 일탈 원인
        원인 미상

    담당자 지적(2026-09): "표 안의 일탈사항은 더 요약해서 요점만." 예전에는 일탈보고서의
    설명을 통째로 옮겨 시각 기록('1) 조제 시작 시간 …')까지 실렸다. 말은 지어내지 않고
    원본에서 **고르고 다듬기만** 한다 — 제목은 제품명·Lot 을 떼고, 내용은 시각·번호 기록을
    뺀 문장만, 원인은 일탈보고서의 '원인' 칸을 그대로 쓴다.
    """
    title = re.sub(r"^\S+\s*\([A-Z0-9]+\)\s*", "", dev.get("title") or "").strip()
    out = [("[%s]" % title, "none")] if title else []
    body = []
    for line in (dev.get("description") or "").split("\n"):
        line = line.strip()
        if not line or re.match(r"^\d\)", line):      # '1) 조제 시작 시간 …' 같은 시각 기록은 뺀다
            continue
        body.append(re.sub(r"^\d\.\s*", "", line))
    if body:
        out.append(("* 일탈 내용", "none"))
        out.append((" ".join(body), "none"))
    cause = re.sub(r"^\d\.\s*", "", (dev.get("cause") or "").replace("\n", " ")).strip()
    if cause:
        out.append(("* 일탈 원인", "none"))
        out.append((cause, "none"))
    return out


def _quarter(day):
    return (day.month - 1) // 3 + 1


def _small(para_el):
    for r in para_el.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = r.makeelement(qn("w:rPr"), {}); r.insert(0, rpr)
        for tag in ("sz", "szCs"):
            get_or_add(rpr, tag).set(qn("w:val"), "18")


GROUP_WIDE = re.compile(r"전\s*제품|전\s*품목|일괄|전\s*라인")


def change_covers(cc, name, parts=()):
    """이 변경관리가 이 제품에 해당하는가 — 해당하면 True.

    변경요청서의 '관련 제품' 은 제품 이름을 하나하나 적기도 하지만, '안연고 전 제품',
    '무균라인에서 생산하는 전 제품' 처럼 무리로 적는 일이 더 잦다. 또 주성분을 건드리는
    변경은 그 성분을 쓰는 제품 모두에 해당한다(디겐타안연고의 플루오로메톨론 멸균 온도
    변경 건은 '관련 제품' 이 후메론점안액으로만 적혀 있었지만 디겐타에도 해당한다 —
    담당자 2026-09: "변경관리는 모두 디겐타안연고에 해당이 되는 내용이야").
    """
    where = re.sub(r"\s+", "", (cc.get("products") or "") + " " + (cc.get("title") or ""))
    if not name:
        return True
    if name[:4] in where:
        return True
    if any(part and re.sub(r"\s+", "", part) in where for part in parts):
        return True                       # 주성분을 건드리는 변경
    # '안연고 전 제품', '무균라인에서 생산하는 전 제품' 처럼 무리로 적은 것은 다 해당한다.
    # 어느 무리에 드는지는 서류만 봐서 알 수 없고(안연고는 무균라인에서 만든다), 12항 폴더에
    # 그 변경요청서를 넣은 사람이 담당자다 — 넣었다는 것 자체가 해당한다는 뜻이다.
    return bool(GROUP_WIDE.search(cc.get("products") or ""))


def _latest_pair(entry, a="IQ", b="OQ"):
    """IQ 와 OQ 의 가장 최근 문서가 같은 한 문서(IOQ…)면 (문서, 일자), 아니면 None."""
    def latest(kind):
        got = [(d, dt) for d, dt in entry.get(kind, []) if d and dt]
        return max(got, key=lambda x: re.sub(r"\D", "", x[1])[:8]) if got else None
    la, lb = latest(a), latest(b)
    if la and lb and la[0].split("(")[0].strip() == lb[0].split("(")[0].strip():
        return la
    if la and not lb and la[0].upper().startswith("IOQ"):
        return la
    return None


def blank_qualification_cells(table, lookup):
    """채우고 난 뒤의 검토 — 마스터파일에 문서가 있는데도 빈 IQ·OQ·PQ 칸. [(관리번호, 종류)].

    담당자 2026-09-06: "여러 번 말한 부분이니까 다음부터는 실수하지 않게 작성할 때 꼭 검토해 줘".
    """
    rows = table.rows
    width = E.grid_width(table)
    kind_col = {}
    for row in rows[:4]:
        for ci, cell in E.grid_cells(row, width).items():
            head = E.cell_text(cell).strip().upper()
            if head in ("IQ", "OQ", "PQ"):
                kind_col[head] = ci
        if kind_col:
            break
    out = []
    for ri, row in enumerate(rows):
        cells = E.raw_cells(row)
        mid = E.cell_text(cells[1]).strip() if len(cells) > 1 else ""
        if not re.match(r"^[A-Z]{3}\d{4}", mid) or mid not in lookup or ri + 1 >= len(rows):
            continue
        grid = E.grid_cells(row, width)
        for kind, col in kind_col.items():
            if not [1 for d, dt in lookup[mid].get(kind, []) if d and dt]:
                continue                                    # 마스터에도 없다 — 비워 두는 것이 맞다
            if col not in grid:
                continue                                    # 옆 칸(IQ)과 합쳐져 거기에 적혔다
            if not E.cell_text(grid[col]).strip():
                out.append((mid, kind))
    return out


def _crit_key(text):
    """9.1 기준 글을 견주기 위한 열쇠 — '허가)/자가)' 머리, '1)' 번호, 띄어쓰기, 끝 마침표를 뺀다."""
    t = D.PREFIX.sub("", text or "")
    t = re.sub(r"^\s*\d\)\s*", "", t)
    return re.sub(r"[^0-9A-Za-z가-힣]", "", t)          # 띄어쓰기·기호('~'·'-'·괄호) 차이는 무시한다


def _as_result(crit):
    """기준 문장을 결과형으로 — '…을 나타낸다.' → '…을 나타냄', '…동일해야 한다.' → '…동일함'."""
    t = re.sub(r"^\s*\d\)\s*", "", D.PREFIX.sub("", crit or "")).strip().rstrip(".")
    t = re.sub(r"나타낸다$", "나타냄", t)
    t = re.sub(r"(이어야|해야)\s*한다$", "함", t)
    t = re.sub(r"(생긴다|된다)$", lambda m: {"생긴다": "생김", "된다": "됨"}[m.group(1)], t)
    return t


def _prior_91(data, idx=0):
    """전년도 결재본 9.1 표(idx: 0 내수 / 1 수출)의 {기준 열쇠: 결과 글}."""
    olds = (getattr(data, "prev_sections_all", None) or {}).get("9.1") or []
    grid = olds[idx] if len(olds) > idx else None
    prior = {}
    for row_ in (grid or [])[1:]:
        crit_ = next((t for t in row_ if D.PREFIX.search(t or "")), None)
        res_ = next((t for t in reversed(row_) if (t or "").strip()), "")
        if crit_ and res_ and res_ != crit_:
            prior[_crit_key(crit_)] = res_.strip()
    return prior


def _ident_result(lots, rec, prior, crit_text):
    """확인 시험 결과 — 올해 성적서가 모두 '확인시험 적합' 이면 전년도 결재본 문안(같은 기준 줄) 또는 기준 문장의
    결과형을 쓴다. 성적서에 확인시험 판정이 없거나 부적합이 있으면 None."""
    oks = [rec(l, "924").get("ident_ok") for l in lots]
    if not lots or any(v is not True for v in oks):
        return None
    old = prior.get(_crit_key(crit_text))
    if old and not re.search(r"\d", old):
        return old
    return _as_result(crit_text)


CODE = re.compile(r"^[A-Z]{1,3}\d{3,5}")           # 관리번호: RBO101 · EPP116 · P17039


def _has_codes(table):
    return any(CODE.match(E.cell_text(E.raw_cells(r)[1]).strip())
               for r in table.rows[1:] if len(E.raw_cells(r)) > 1)


def _seed_rows_by_header(table, old_grid):
    """머리행 이름이 같은 열끼리 전년도 결재본의 자료 줄(관리번호가 있는 줄)을 옮겨 세운다. 연번은 새로 매기고
    '특이사항' 줄은 건드리지 않는다. 옮긴 줄 수를 돌려준다."""
    if not old_grid or len(old_grid) < 2:
        return 0
    width = E.grid_width(table)
    new_head = {CARRY.squeeze(E.cell_text(c)): i for i, c in E.grid_cells(table.rows[0], width).items()}
    old_head = {CARRY.squeeze(t): j for j, t in enumerate(old_grid[0]) if CARRY.squeeze(t)}
    same = {i: old_head[name] for name, i in new_head.items() if name in old_head}
    code_j = next((j for name, j in old_head.items() if "관리번호" in name or "코드" in name), None)
    if code_j is None or len(same) < 2:
        return 0
    rows = [r for r in old_grid[1:] if code_j < len(r) and CODE.match((r[code_j] or "").strip())]
    if not rows:
        return 0
    last = len(table.rows) - 1
    if CARRY.squeeze(E.cell_text(E.raw_cells(table.rows[-1])[0])).startswith("특이사항"):
        last -= 1
    f, l = E.fit_rows(table, 1, last, len(rows))
    for k, r in enumerate(rows):
        cells = E.grid_cells(table.rows[f + k], width)
        for i, cell in cells.items():
            E.clear_diag(cell)
            E.set_vmerge(cell, False)
            if i == 0:
                E.set_cell(cell, str(k + 1))
                continue
            j = same.get(i)
            text = (r[j] if j is not None and j < len(r) else "") or ""
            E.set_cell(cell, *text.split("\n"))
    return len(rows)


def _plain_stage(text):
    """공정 이름에서 괄호 설명과 빈칸을 뗀 이름 — '포장 (베트남)' → '포장'."""
    return re.sub(r"[\s]", "", re.sub(r"[(（][^)）]*[)）]", "", str(text or "")))


def _yield_columns(table):
    """7항 표의 수율 열 — [(그리드 열, 공정 이름)]. '공정' 칸 오른쪽부터 '비고' 앞까지.

    제품에 따라 포장이 내수·베트남으로 갈린다(담당자 2026-09-07). 열을 셋으로 박아 두면
    베트남 열이 통째로 빠지고, 그 시장으로 포장한 Lot 이 '확인 필요' 로 남는다.
    """
    width = E.grid_width(table)
    for row in table.rows[:3]:
        cells = E.grid_cells(row, width)
        texts = {i: D.squeeze(E.cell_text(c)) for i, c in cells.items()}
        head = [i for i, t in texts.items() if t == "공정"]
        if not head:
            continue
        got, seen = [], set()
        for i in sorted(texts):
            name = texts[i]
            if i <= head[0] or not name or "비고" in name or name in seen:
                continue
            seen.add(name)
            got.append((i, name))
        if got:
            return got
    return [(2, "조제"), (3, "충전"), (4, "포장")]


def _yield_first_row(table):
    """7항 표에서 자료(Lot)가 시작하는 줄 — 기준 줄 다음.

    담당자 2026-09-08: "LWY201 수율 값이 기준에 적혀 있네." 머리가 세 줄인 표(중요공정 /
    공정·조제·충전·포장 / 내수·온누리)가 있어 자리를 넷째 줄로 박아 두면 첫 Lot 이 기준 줄에 찍힌다.
    """
    width = E.grid_width(table)
    last_head = 0
    for i, row in enumerate(table.rows[:5]):
        texts = [E.cell_text(c) for c in E.grid_cells(row, width).values()]
        if any(("±" in t or "이상" in t or "이하" in t) for t in texts):
            last_head = i                       # 기준 줄
        elif any(D.squeeze(t) in ("공정", "기준", "LotNo.", "연번") for t in texts):
            last_head = max(last_head, i)       # 이름 줄
    return min(last_head + 1, max(1, len(table.rows) - 4))


def _yield_value(vals, name):
    """표 열 이름으로 수율현황표 값 찾기 — 이름이 그대로 있으면 그것을, 없으면 괄호를 뗀 이름으로."""
    if name in vals:
        return vals[name]
    base = _plain_stage(name)
    same = [v for k, v in vals.items() if _plain_stage(k) == base]
    return same[0] if len(same) == 1 else None


def _other_market(columns, values, k):
    """이 열은 비었는데 같은 공정의 다른 시장 열에는 값이 있는가 — 그러면 '확인 필요' 가 아니라 사선."""
    base = _plain_stage(columns[k][1])
    return any(values[j] is not None and _plain_stage(name) == base
               for j, (_gi, name) in enumerate(columns) if j != k)


def _last_data_row(table):
    """마지막 자료 줄 = '특이사항' 줄 바로 위. 특이사항이 없으면 마지막 줄.

    담당자 2026-09-07: "특이사항은 연번 4로 적는 게 아니고 13.2항 표 연번 3 아래 바로 기재해야지."
    자료 줄이 모자라면 마지막 자료 줄을 복제해 늘리는데, 그 자리를 '끝에서 두 번째' 로 박아 두면
    특이사항 줄이 복제되어 그 글이 연번 줄 안으로 들어간다.
    """
    for i in range(len(table.rows) - 1, 0, -1):
        cells = E.raw_cells(table.rows[i])
        if cells and CARRY.squeeze(E.cell_text(cells[0])).startswith("특이사항"):
            return max(1, i - 1)
    return max(1, len(table.rows) - 1)


def _mark_carried_cells(table, old_grid):
    """전년도 결재본에서 옮겨 온 값이 그대로 남은 칸을 노랑으로 표시한다 → 표시한 칸 수.

    담당자 2026-09-07: "지금 기록서가 없어 내용이 없거나 이상하면 전년도 PQR 정보로 가져오고 노랑 MARK 해."
    올린 자료로 갱신된 칸은 값이 달라져 표시되지 않는다 — 노랑은 '아직 작년 것' 이라는 뜻이다.
    """
    if not old_grid or len(old_grid) < 2 or not table.rows:
        return 0
    width = E.grid_width(table)
    new_head = {CARRY.squeeze(E.cell_text(c)): i for i, c in E.grid_cells(table.rows[0], width).items()}
    old_head = {CARRY.squeeze(t): j for j, t in enumerate(old_grid[0]) if CARRY.squeeze(t)}
    code_i = next((i for name, i in new_head.items() if "관리번호" in name or "코드" in name), None)
    code_j = next((j for name, j in old_head.items() if "관리번호" in name or "코드" in name), None)
    if code_i is None or code_j is None:
        return 0
    by_code = {}
    for r in old_grid[1:]:
        code = (r[code_j] if code_j < len(r) else "").strip()
        if CODE.match(code):
            by_code.setdefault(code, r)
    n = 0
    for row in table.rows[1:]:
        cells = E.grid_cells(row, width)
        cell = cells.get(code_i)
        old = by_code.get(E.cell_text(cell).strip()) if cell is not None else None
        if not old:
            continue
        for name, i in new_head.items():
            # 관리번호·품명은 해마다 같은 것이 정상이라 노랑으로 칠하지 않는다
            # (담당자 2026-09-08: "8.1.1 히아루론산나트륨은 왜 노랑 마크인지?")
            if i in (0, code_i) or not name or any(w in name for w in ("자재명", "원료명", "품명")):
                continue
            j = old_head.get(name)
            here = cells.get(i)
            want = (old[j] if j is not None and j < len(old) else "").strip()
            if here is None or not want:
                continue
            if E.cell_text(here).strip() == want:
                E.highlight_cell(here)
                n += 1
    return n


def _base_name(text):
    """원/자재명에서 괄호 설명·빈칸을 뗀 이름 — '케이스(내수)' → '케이스', '케 이 스' → '케이스'."""
    return re.sub(r"[\s·]", "", re.sub(r"[(（][^)）]*[)）]", "", str(text or "")))


def _one_letter_off(a, b):
    """관리번호가 한 자만 다른가 — 결재본의 오기(P38033)와 공 기록서(P38003)를 알아보려고."""
    return len(a) == len(b) and sum(1 for x, y in zip(a, b) if x != y) == 1


EQUIPMENT_CODE = re.compile(r"^(DA|DE|FA|HA|HC|HE|AI|CW)[A-Z]?\d{3,5}")


def _is_equipment(code):
    """관리번호가 설비인가 — 8.1.3 은 부원료·포장자재 표다.

    담당자 2026-09-08: "8.1.3 부원료 및 포장자재인데 고압증기 멸균기를 적는 것은 적절하지 않아,
    고압증기 멸균기는 생산 장비야." 공 기록서의 자재 목록에 설비 관리번호(DAE5024 …)가 섞여 온다.
    설비는 10.2 적격성 표가 다룬다.
    """
    return bool(EQUIPMENT_CODE.match(str(code or "").strip().upper()))


def _add_material_rows(table, materials, fixed=None):
    """표에 없는 관리번호의 원/자재를 줄로 보탠다 — 관리번호·원/자재명·규격만, 나머지 칸은 비운다(노랑).
    표가 비어 있으면(빈 공양식) 그 줄들이 표가 된다. 보탠 줄 수를 돌려준다.

    표의 관리번호가 공 기록서와 한 자만 다르고 이름이 같으면 결재본의 오기로 보고 기록서 번호로
    고친다 — 같은 자재가 두 줄이 되지 않게 (담당자 2026-09-07: "케이스는 내수 케이스야, 기존
    P38033 관리번호 오기라서 P38003 으로 수정해 줘"). 고친 것은 fixed 목록에 (옛, 새)로 담는다.
    """
    materials = [m for m in (materials or []) if not _is_equipment(m.get("code"))]
    if not materials:
        return 0
    width = E.grid_width(table)
    head = {CARRY.squeeze(E.cell_text(c)): i for i, c in E.grid_cells(table.rows[0], width).items()}
    col = lambda *words: next((i for name, i in head.items() if any(w in name for w in words)), None)
    c_code, c_name, c_spec = col("관리번호", "코드"), col("원/자재명", "원자재명", "자재명", "원료명"), col("규격")
    if c_code is None:
        return 0
    have = {}
    rows = []
    for ri, row in enumerate(table.rows[1:], 1):
        cells = E.grid_cells(row, width)
        text = E.cell_text(cells[c_code]).strip() if c_code in cells else ""
        if CODE.match(text):
            have[text] = (ri, E.cell_text(cells[c_name]).strip() if (c_name is not None and c_name in cells) else "")
            rows.append(ri)
        elif not CARRY.squeeze(E.cell_text(E.raw_cells(row)[0])).startswith("특이사항"):
            rows.append(ri)                                  # 빈 줄(공양식)
    new = []
    for m in materials:
        if m["code"] in have:
            continue
        # 한 자만 다른 관리번호에 같은 이름이 붙어 있으면 그 줄의 번호가 오기다 — 줄을 보태지 않고 고친다
        같은것 = [(code, ri) for code, (ri, nm) in have.items()
                if _one_letter_off(code, m["code"]) and _base_name(nm)
                and (_base_name(m["name"]) in _base_name(nm) or _base_name(nm) in _base_name(m["name"]))]
        if len(같은것) == 1:
            code, ri = 같은것[0]
            cells = E.grid_cells(table.rows[ri], width)
            if c_code in cells:
                E.set_cell(cells[c_code], m["code"])
                E.highlight_cell(cells[c_code])
                have[m["code"]] = have.pop(code)
                if fixed is not None:
                    fixed.append((code, m["code"]))
                continue
        new.append(m)
    if not new:
        return 0
    filled = [ri for ri in rows if CODE.match(E.cell_text(E.grid_cells(table.rows[ri], width)[c_code]).strip())]
    last = rows[-1] if rows else 0
    f, l = E.fit_rows(table, 1, last, len(filled) + len(new)) if rows else (1, len(new))
    for k, m in enumerate(new):
        cells = E.grid_cells(table.rows[f + len(filled) + k], width)
        for i, cell in cells.items():
            E.clear_diag(cell); E.set_vmerge(cell, False)
            E.set_cell(cell, "")
        if 0 in cells:
            E.set_cell(cells[0], str(len(filled) + k + 1))
        if c_code in cells:
            E.set_cell(cells[c_code], m["code"])
        if c_name is not None and c_name in cells:
            E.set_cell(cells[c_name], m["name"]); E.highlight_cell(cells[c_name])
        if c_spec is not None and c_spec in cells:
            E.set_cell(cells[c_spec], m.get("spec") or ("자사규격" if m["code"].startswith("P") else ""))
    return len(new)


def _has_equipment(table):
    return any(re.match(r"^[A-Z]{3}\d{4}", E.cell_text(E.raw_cells(r)[1]).strip())
               for r in table.rows if len(E.raw_cells(r)) > 1)


def _seed_equipment_rows(table, rows):
    """빈 설비 표에 전년도 결재본의 설비 줄(관리번호·설비명·문서번호·완료일)을 세운다 — 설비 하나가 두 줄.
    문서·완료일은 뒤에 update_qualification 이 마스터파일로 갱신한다."""
    if not rows:
        return 0
    width = E.grid_width(table)
    kind_col = {}
    for row in table.rows[:4]:
        for ci, cell in E.grid_cells(row, width).items():
            head = E.cell_text(cell).strip().upper()
            if head in ("IQ", "OQ", "PQ"):
                kind_col[head] = ci
        if kind_col:
            break
    first = 4 if len(table.rows) > 4 else len(table.rows) - 1
    f, l = E.fit_rows(table, first, len(table.rows) - 1, 2 * len(rows))
    for i, r in enumerate(rows):
        for half in (0, 1):
            grid = E.grid_cells(table.rows[f + 2 * i + half], width)
            head = half == 0
            for ci, cell in grid.items():
                E.clear_diag(cell)
                if ci in (0, 1, 2):
                    text = [str(i + 1), r["mid"], r["name"]][ci] if head else ""
                    E.set_cell(cell, *(text.split("\n") if text else [""]))
                    E.set_vmerge(cell, "restart" if head else None)
                elif ci in kind_col.values():
                    kind = next(k for k, c in kind_col.items() if c == ci)
                    doc, day = r["docs"].get(kind, ("", ""))
                    E.set_cell(cell, doc if head else day.replace(". ", ".").replace(" ", ""))
                    E.set_vmerge(cell, False)
                else:
                    E.set_cell(cell, "")
                    E.set_vmerge(cell, False)
    return len(rows)


def _spans(cell):
    """이 칸이 덮는 그리드 열 수 — 1 이면 저 혼자, 2 이상이면 옆 칸과 합쳐져 있다."""
    pr = cell._tc.find(qn("w:tcPr"))
    el = pr.find(qn("w:gridSpan")) if pr is not None else None
    return int(el.get(qn("w:val"))) if el is not None else 1


def update_qualification(table, lookup):
    """관리번호 행마다 IQ·OQ·PQ 열의 (문서번호 행 · 완료일 행) 을 마스터 값으로 채운다.

    빈 서식(공양식)의 IQ·OQ 칸은 사선만 그어져 있고 비어 있다 — 마스터파일에서 읽어
    채우고 사선을 지운다(담당자 2026-09). 이미 값이 있으면 마스터가 더 최근일 때만 바꾼다.
    """
    rows = table.rows
    width = E.grid_width(table)
    kind_col = {}
    for row in rows[:4]:
        for ci, cell in E.grid_cells(row, width).items():
            head = E.cell_text(cell).strip().upper()
            if head in ("IQ", "OQ", "PQ"):
                kind_col[head] = ci
        if kind_col:
            break
    if not kind_col:
        return 0
    n = 0
    for ri, row in enumerate(rows):
        cells = E.raw_cells(row)
        mid = E.cell_text(cells[1]).strip() if len(cells) > 1 else ""
        if not re.match(r"^[A-Z]{3}\d{4}", mid) or mid not in lookup:
            continue
        if ri + 1 >= len(rows):
            continue
        # 열 번호는 그리드 기준이다 — IQ·OQ 를 한 칸에 합쳐 적은 줄(IOQ…)이 있으면
        # 자리로 세었을 때 옆 칸(PQ)을 덮어쓴다.
        doc_grid, date_grid = E.grid_cells(rows[ri], width), E.grid_cells(rows[ri + 1], width)
        # IQ·OQ 를 하나로 합친 문서(IOQ20-UT-HEA5029-R)는 두 칸을 합쳐 한 번만 적는다 — 결재본 관행.
        # 공양식의 IQ·OQ 칸이 비어 있으면(담당자 PC 2026-09-06: 10.4·10.5 IOQ 칸이 사선만) 여기서
        # 칸을 합쳐 채우고, 반대로 합쳐진 칸에 IQ·OQ 문서가 따로 있으면 칸을 나눠 각각 적는다.
        ioq = _latest_pair(lookup[mid], "IQ", "OQ")
        if ioq is not None and "IQ" in kind_col and "OQ" in kind_col:
            ci, co = kind_col["IQ"], kind_col["OQ"]
            for grid, row_ in ((doc_grid, rows[ri]), (date_grid, rows[ri + 1])):
                if ci in grid and co in grid and _spans(grid[ci]) == 1:
                    E.merge_right(grid[ci], grid[co])
            doc_grid, date_grid = E.grid_cells(rows[ri], width), E.grid_cells(rows[ri + 1], width)
        elif ioq is None and "IQ" in kind_col and "OQ" in kind_col:
            ci, co = kind_col["IQ"], kind_col["OQ"]
            if any(lookup[mid].get(k) for k in ("IQ", "OQ")):
                for grid in (doc_grid, date_grid):
                    if ci in grid and co not in grid and _spans(grid[ci]) > 1 and not E.cell_text(grid[ci]).strip():
                        E.split_span(grid[ci])          # 비어 있는 합친 칸만 나눈다 — 적힌 값은 건드리지 않는다
                doc_grid, date_grid = E.grid_cells(rows[ri], width), E.grid_cells(rows[ri + 1], width)
        for kind, col in kind_col.items():
            got = [(d, dt) for d, dt in lookup[mid].get(kind, []) if dt and d]
            if not got:
                continue
            latest = max(got, key=lambda x: re.sub(r"\D", "", x[1])[:8])
            doc_cells, date_cells = doc_grid, date_grid
            if col not in doc_cells or col not in date_cells:
                continue                    # 그 줄에서는 옆 칸과 합쳐져 있다 — IQ 칸에서 함께 적는다
            if (_spans(doc_cells[col]) > 1 or _spans(date_cells[col]) > 1) and not (kind == "IQ" and ioq is not None):
                continue                    # 합쳐진 칸인데 합친 문서가 아니다 — 그대로 둔다
            if kind == "IQ" and ioq is not None:
                latest = ioq
            old_doc = E.cell_text(doc_cells[col]).strip()
            old_date = E.cell_text(date_cells[col]).strip()
            doc = latest[0].split("(")[0].strip()
            new_date = latest[1].replace(". ", ".").replace(" ", "")[:10]
            blank = not old_doc or not old_date      # 문서번호나 완료일 한쪽만 비어도 마스터로 채운다
            newer = (re.sub(r"\D", "", new_date) > re.sub(r"\D", "", old_date)
                     and doc != old_doc.split("(")[0].strip())
            if not (blank or newer):
                continue
            E.set_cell(doc_cells[col], doc)
            E.set_cell(date_cells[col], new_date)
            E.clear_diag(doc_cells[col]); E.clear_diag(date_cells[col])
            n += 1
    return n


LINE = "연고"          # 이 조리법은 연고·안연고 라인 전용이다
FLOOR = "1"            # 연고 라인은 1층에 있다 (담당자 2026-09-07: "내가 검토한 결과 연고라인은 1층에 위치해 있어")

_FLOOR = re.compile(r"(\d)\s*층")


def _floors(name):
    """설비명에 적힌 층 — '주사용수 제조 시스템 (2t - 2층)' → {"2"}. 없으면 빈 집합."""
    return set(_FLOOR.findall(str(name or "")))


def _kind_name(name):
    """설비명에서 괄호 설명을 뗀 이름 — '주사용수 제조 시스템 (2t - 2층)' → '주사용수제조시스템'."""
    return re.sub(r"[\s·]", "", re.sub(r"[(（][^)）]*[)）]", "", str(name or "")))


def use_our_floor(table, support, floor=FLOOR):
    """다른 층 설비로 적힌 줄을 같은 종류의 우리 층 설비로 바꾼다. [(옛 관리번호, 새 관리번호)].

    제조용수(10.4)는 층마다 시스템이 따로 있고, 적격성평가 현황표의 설비명에 층이 적혀 있다
    ('주사용수 제조 시스템 (1.5t - 1층, 냉주사용수)' · '(2t - 2층)'). 연고 라인은 1층이라 2층
    설비를 실으면 안 된다 — 전년도 결재본이 2층 설비로 적혀 있어도 바로잡는다
    (담당자 2026-09-07: "제조용수 정보는 왼쪽 내용대로 작성하면 돼").
    """
    swapped = []
    for row in table.rows:
        cells = E.raw_cells(row)
        if len(cells) < 2:
            continue
        mid = E.cell_text(cells[1]).strip()
        one = support.get(mid)
        if not one:
            continue
        floors = _floors(one.get("name"))
        if not floors or floor in floors:
            continue                       # 층이 안 적혔거나 이미 우리 층 설비다
        kind = _kind_name(one.get("name"))
        best, best_score = "", -1
        for mid2, two in sorted(support.items()):
            if mid2 == mid or _kind_name(two.get("name")) != kind:
                continue
            if floor not in _floors(two.get("name")):
                continue
            score = (2 if (two.get("system") or "") == (one.get("system") or "") else 0) \
                + (1 if mid2[:3] == mid[:3] else 0)
            if score > best_score:
                best, best_score = mid2, score
        if not best:
            continue
        E.set_cell(cells[1], best)
        swapped.append((mid, best))
    return swapped


def fill(document, data, product, period, today=None, log=None):
    log = log or (lambda *a: None)
    issues = []
    today = today or _dt.date.today()
    if isinstance(today, str):
        today = _dt.date(*[int(x) for x in re.findall(r"\d+", today)[:3]])
    write_year = today.year
    dom, exp = data.domestic, data.export
    # 평가 대상 연도는 제조번호에서 읽는다 — 세 번째 글자가 제조 연도이고 PQR 은 그 다음 해 것이다
    # (OGY301 → 2025년 제조 → 2026년 PQR). 화면에서 넘어온 기간보다 자료가 앞선다.
    year_from, _pqr_year = lotcode.years(dom + exp, today)
    if year_from is None:
        year_from = lotcode.year_of(period.get("from"))
    else:
        odd = lotcode.odd_lots(dom + exp, year_from, today)
        if odd:
            issues.append(("6", ", ".join(odd),
                           "제조 연도가 다른 Lot 이 섞여 있습니다(%d년 것으로 평가함) — 확인하세요" % year_from))
    mfg = {l: d for l, d, _ in data.lots}
    name = product.get("name") or ""

    # ---------- 머리글 ----------
    # 머리글은 칸의 위치가 아니라 라벨로 찾는다. 전년도 결재본(HLF-QC-126-01)은 문서번호·
    # 작성일자·제품명 행이 있지만, EDMS 서식(E-HLF-32)은 문서번호·Rev. No.·Page 뿐이라
    # 세 번째 행이 쪽 번호 필드다 — 위치로 쓰면 거기에 날짜를 덮어쓴다.
    hdr = document.sections[0].header.tables[0]
    new_no, full_name = "", name
    for row in hdr.rows:
        cells = E.raw_cells(row)
        labels = "".join(E.cell_text(c) for c in cells[:-1])
        target = cells[-1]
        if "문서번호" in labels:
            old_no = E.cell_text(target).strip()
            if re.search(r"PQR\d{2}-", old_no):
                new_no = re.sub(r"PQR\d{2}-", "PQR%02d-" % (write_year % 100), old_no)
                E.set_cell(target, new_no)        # EDMS 서식은 비어 있다 — EDMS 가 번호를 준다
        elif "작성일자" in labels:
            E.set_cell(target, today.strftime("%Y.%m.%d"))
        elif "제품명" in labels:
            # 결론(16항)에는 머리글의 정식 제품명(성분명까지)을 쓴다 — 마스터의 짧은 이름이 아니라.
            full_name = E.cell_text(target).strip() or name
    log("머리글: %s / %s" % (new_no or "(EDMS 서식 — 문서번호 비움)", today.strftime("%Y.%m.%d")))

    # ---------- 3항 대상 제품 ----------
    t3 = _tables(document, "3.")
    if t3:
        # 비고 칸은 빈 줄끼리 이어 붙여 병합하고 사선 하나만 긋는다 (담당자 지시 2026-09)
        # — 줄마다 'N/A' 를 적거나 사선을 여러 개 긋지 않는다.
        blocks, rows_ = E.merge_empty_runs(t3[0], "비고")
        if blocks:
            log("3항 비고: 빈 칸 %d줄을 %d 묶음으로 합치고 사선" % (rows_, blocks))
        # 결론(16항)의 제품명은 정식 이름(성분명까지)이다 — 머리글이 비어 있으면 여기서 가져온다.
        for row in t3[0].rows[1:]:
            cells = E.raw_cells(row)
            if len(cells) >= 3 and "제품명" in E.cell_text(cells[1]):
                got = " ".join(E.cell_text(cells[2]).split())
                if got and (not full_name or full_name == name):
                    full_name = got
                break

    # ---------- 5항 책임과 권한 ----------
    # 담당자 지시(2026-09): "품질보증 1팀은 AQA 팀으로 변경해줘." EDMS 서식과 2026 결재본
    # 모두 'AQA팀 담당' · 'AQA팀 팀장' 이다. 전년도 결재본은 옛 이름이라 그대로 물려받는다.
    # '품질보증부서장' 은 팀이 아니므로 건드리지 않는다.
    t5 = _tables(document, "5.")
    if t5:
        renamed = 0
        for row in t5[0].rows[1:]:
            for cell in E.raw_cells(row):
                text = E.cell_text(cell)
                new_text = re.sub(r"품질보증\s*\d*\s*팀", "AQA팀", text)
                if new_text != text:
                    E.set_cell(cell, *new_text.split("\n"))
                    renamed += 1
        if renamed:
            log("5항: 품질보증n팀 → AQA팀 %d칸" % renamed)

    # ---------- 4항 ----------
    # 1항 '목적' 도 "제품품질평가는" 으로 시작하므로 평가 기간 문장만 집어 찾는다.
    p = find_para(document, "월까지 생산된")
    if p is not None and "월" in p.text:
        # 2026 결재본 문안: "제품품질평가는 2025년도 1월 ~ 12월까지 생산된 해당제품에 대하여 평가를
        # 실시하며, 'QC-126 제품품질평가 규정'에 따라 2 그룹으로 선정되어 차년도 3분기 내에 완료한다."
        # — 평가 대상 연도는 '년도', 마감은 연도를 적지 않고 '차년도 N분기' 로 쓴다.
        text = re.sub(r"\d{4}\s*년도?\s*1월", "%d년도 1월" % year_from, p.text)
        due = None
        try:
            due = _dt.date(*[int(x) for x in re.findall(r"\d+", str(product.get("due") or ""))[:3]])
        except (TypeError, ValueError):
            due = None
        q = _quarter(due) if due else _quarter(today)
        text = re.sub(r"(\d{4}\s*년도|차년도)\s*(상반기|하반기|\d\s*분기)", "차년도 %d분기" % q, text)
        group = str(product.get("group") or "").strip()
        if group and re.search(r"\S+\s*그룹으로 선정", text):
            text = re.sub(r"\S+\s*그룹으로 선정", "%s 그룹으로 선정" % group, text)
        E.set_para_text(p, text)

    # ---------- 6항 제조내역 ----------
    olds6 = (getattr(data, "prev_sections_all", None) or {}).get("6") or []

    def fill_mfg(table, lots, old_grid=None, batch=None):
        if len(table.rows) < 2:
            return
        keep = [E.cell_text(c) for c in E.raw_cells(table.rows[1])]      # 배치 크기·포장 단위는 결재본 값
        # 빈 공양식이면 제조단위·포장단위 칸이 비어 있다 — 전년도 결재본 같은 표(6.1 내수 / 6.2 수출)에서 가장
        # 많이 적힌 값을 쓴다. 해마다 같은 값이다(담당자 2026-09-06: "제조단위와 포장단위가 공란인 이유가 있나?")
        batch = batch or {}
        for k, key in ((3, "batch_size"), (4, "pack_unit")):
            if k < len(keep) and not keep[k].strip() and batch.get(key):
                keep[k] = batch[key]                         # 6항에 올린 공 기록서(제조·포장) 값 (담당자 2026-09-06)
                log("6항: %s 를 공 기록서 값(%s)으로 채움" % ("제조단위" if k == 3 else "포장단위", keep[k]))
        carried = set()                              # 전년도에서 가져온 칸 — 노랑으로 표시한다
        for k in (3, 4):
            if k < len(keep) and not keep[k].strip() and old_grid:
                seen = {}
                for row in old_grid[1:]:
                    v = (row[k] if k < len(row) else "").strip()
                    if v and not v.startswith("특이사항"):
                        seen[v] = seen.get(v, 0) + 1
                if seen:
                    keep[k] = max(seen.items(), key=lambda kv: kv[1])[0]
                    carried.add(k)
                    log("6항: %s 를 전년도 결재본 값(%s)으로 채움" % ("제조단위" if k == 3 else "포장단위", keep[k]))
                    issues.append(("6", "제조단위" if k == 3 else "포장단위",
                                   "올해 공 기록서가 없어 전년도 결재본 값(%s)을 옮겼습니다(노랑) — "
                                   "이 제품의 제조·충전·포장 기록서를 6항에 올리면 올해 값으로 채웁니다" % keep[k]))
        f, l = E.fit_rows(table, 1, len(table.rows) - 1, max(1, len(lots)))
        for i, lot in enumerate(lots):
            c = E.raw_cells(table.rows[f + i])
            E.set_cell(c[0], str(i + 1)); E.set_cell(c[1], lot); E.set_cell(c[2], mfg.get(lot, ""))
            for k in range(3, len(c)):
                E.set_cell(c[k], keep[k] if k < len(keep) else "")
                if k in carried:
                    E.highlight_cell(c[k])
            if len(c) > 5:
                E.set_cell(c[5], "■ 적합 □ 부적합")
            if len(c) > 6:
                E.set_cell(c[6], "")                  # 비고는 비워 사선 하나로 합친다 — 공양식의 'N/A' 를 줄마다 옮기지 않는다
    t6 = _tables(document, "6.")
    if t6:
        fill_mfg(t6[0], dom, olds6[0] if olds6 else None, getattr(data, "batch", None))
        if len(t6) > 1:
            fill_mfg(t6[1], exp, olds6[1] if len(olds6) > 1 else None, getattr(data, "batch_exp", None))

    # ---------- 7항 수율 ----------
    def yield_specs(table):
        """기준 행('95.0% 이상' · '91.0 ± 4.0%' · '98 ± 2%') → [(lo, hi) or None]"""
        specs = []
        for row in table.rows[:3]:
            cells = [E.cell_text(c) for c in E.raw_cells(row)]
            if any("이상" in c or "±" in c for c in cells):
                for c in cells[-4:-1] if len(cells) >= 5 else cells:
                    if "±" in c:
                        m, d = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", c)[:2]]
                        specs.append((m - d, m + d))
                    elif "이상" in c:
                        specs.append((_num(c), None))
                    else:
                        specs.append(None)
                break
        return specs

    yield_out, yield_dev = [], []
    dev_lots = {d.get("lot") for d in data.deviations if "수율" in (d.get("title") or "")}
    t7 = _tables(document, "7.")

    def put_sheet_specs(table, stages, is_dom=True):
        """올해 수율현황표에 적힌 기준을 결재본의 기준 행에 옮긴다.

        기준은 해가 바뀌며 개정된다(디겐타안연고 충전: 96.0 ± 3.5% → 86.5 ± 6.5%). 전년도
        결재본의 기준을 그대로 두면 표에 지난해 기준이 적히고, 멀쩡한 Lot 이 모두
        '기준 벗어남' 으로 잡힌다.
        """
        # 수율현황표에도 결재본 기준 줄에도 기준이 없는 공정은 6항 공 기록서(제조·충전·포장)의 수율 기준을 쓴다
        # (담당자 2026-09-06). 수출용 표는 '(수출용)' 기록서가 있을 때만.
        batch = getattr(data, "batch" if is_dom else "batch_exp", None) or {}
        fallback = batch.get("yield_specs") or {}
        if not data.yield_specs and not fallback:
            return
        for row in table.rows[:3]:
            cells = E.raw_cells(row)
            texts = [E.cell_text(c) for c in cells]
            if not any("이상" in t or "±" in t or "기준" in t for t in texts):
                continue
            targets = cells[-4:-1] if len(cells) >= 5 else cells
            for k, stage in enumerate(stages):
                spec = data.yield_specs.get(stage)
                if spec and k < len(targets):
                    E.set_cell(targets[k], spec)
                elif k < len(targets) and fallback.get(stage) and not E.cell_text(targets[k]).strip():
                    log("7항: %s 수율 기준을 공 기록서 값(%s)으로 채움" % (stage, fallback[stage]))
                    E.set_cell(targets[k], fallback[stage])
            return

    def fill_yield(table, lots, is_dom):
        # 공정 열은 표에서 읽는다 — 제품에 따라 '포장' 이 내수·베트남으로 갈린다
        # (담당자 2026-09-07: "ELYN01 과 ELYN02 는 내수가 아니고 베트남 포장했네").
        columns = _yield_columns(table)
        stages = tuple(name for _, name in columns)
        put_sheet_specs(table, stages, is_dom)
        specs = yield_specs(table)
        첫줄 = _yield_first_row(table)
        f, l = E.fit_rows(table, 첫줄, len(table.rows) - 4, max(1, len(lots)))
        width = E.grid_width(table)
        vals = {}
        for i, lot in enumerate(lots):
            row = table.rows[f + i]
            r = E.raw_cells(row)
            cells = E.grid_cells(row, width)
            y = data.yields.get(lot, {})
            v = [_yield_value(y, name) for _, name in columns]
            vals[lot] = v
            E.set_cell(r[0], str(i + 1)); E.set_cell(r[1], lot)
            for k, (gi, name) in enumerate(columns):
                cell = cells.get(gi)
                if cell is None:
                    continue
                if v[k] is not None:
                    E.clear_diag(cell)
                    E.set_cell(cell, v[k])
                elif _other_market(columns, v, k):
                    # 그 Lot 을 그 시장으로 포장하지 않았다 — 빈 칸에 사선 (담당자 수기본과 같게)
                    E.set_cell(cell, "")
                    E.add_diag(cell)
                else:
                    E.clear_diag(cell)
                    E.set_cell(cell, "확인 필요")
                    issues.append(("7", lot, "%s 수율 값이 수율현황표에 없음" % name))
            out = False
            for k in range(len(columns)):
                sp = specs[k] if k < len(specs) else None
                if sp and v[k] is not None:
                    lo, hi = sp
                    fv = float(v[k])
                    if (lo is not None and fv < lo) or (hi is not None and fv > hi):
                        out = True
            if len(r) > 5:                      # 비고: 행마다 따로 (병합을 풀고) 주석을 단다
                E.set_vmerge(r[5], False)
                E.clear_diag(r[5])
                E.set_cell(r[5], "")
            if out:
                yield_out.append(lot)
                if lot in dev_lots:
                    yield_dev.append(lot)
                if len(r) > 5:
                    # 기준을 벗어난 Lot 은 모두 '1)' — 표 아래 각주 한 줄이 설명한다(담당자 2026-09-06: "비고에는
                    # 1), 2) 가 아니고 모두 1) 로 작성되어야 표 하단 문구로 설명되는 것 아닌지"). 11항에 일탈
                    # 기록이 없는 Lot 은 노랑으로 남겨 대조하게 한다.
                    E.set_cell(r[5], "1)")
                    if lot not in dev_lots:
                        E.highlight_cell(r[5])
        # 최댓값·최솟값·평균 — 세 줄 모두 채우고 그 칸의 사선은 지운다 (담당자 지시 2026-09).
        # raw_cells 로 그 행이 실제로 가진 칸을 쓴다: .cells 는 세로 병합을 하나로 합쳐 돌려주어
        # 세 줄이 같은 칸을 가리키고, 마지막에 쓴 평균만 남는다(최댓값 자리에 평균이 찍혔다).
        summary = [(l + 1, lambda xs: "%.2f" % max(xs)),
                   (l + 2, lambda xs: "%.2f" % min(xs)),
                   (l + 3, avg)]
        # 칸은 그리드 열 번호로 짚는다 — 요약 행은 '최댓값' 칸이 연번·Lot No. 두 열을 덮어
        # 자리로 세면 값이 한 칸씩 밀린다(조제 자리에 사선, 비고 자리에 포장 수율).
        for ri, _fn in summary:
            if ri < len(table.rows):
                cells = E.grid_cells(table.rows[ri], width)
                for gi, _name in columns:
                    cell = cells.get(gi)
                    if cell is not None:
                        E.set_vmerge(cell, False)
                        E.clear_diag(cell)
        # 열마다 따로 센다 — 그 시장으로 포장하지 않은 Lot 은 그 열의 셈에서만 빠진다
        for ri, fn in summary:
            if ri >= len(table.rows):
                continue
            cells = E.grid_cells(table.rows[ri], width)
            for k, (gi, _name) in enumerate(columns):
                cell = cells.get(gi)
                if cell is None:
                    continue
                xs = [float(vals[lt][k]) for lt in lots
                      if lt not in yield_out and vals[lt][k] is not None]
                if xs:
                    E.set_cell(cell, fn(xs))
                else:
                    E.set_cell(cell, "")
                    E.add_diag(cell)
    if t7:
        fill_yield(t7[0], dom, True)
        if len(t7) > 1 and exp:
            fill_yield(t7[1], exp, False)
    note = (find_para(document, "수율의 최대, 최소, 평균") or find_para(document, "1) 11. 일탈")
            or find_para(document, "일탈 관련 기록 참고") or find_para(document, "수율 일탈로"))
    if note is not None:
        if yield_out:
            # 각주 한 줄 — 공양식 문안('1) 충전 수율 일탈로 11항 일탈 관련 기록 참고.')에 일탈 문서번호와
            # 결재본 관행('수율의 최대, 최소, 평균 계산에서 제외함')을 붙인다
            docs = [next((d["doc_no"] for d in data.deviations if d.get("lot") == l), "") for l in yield_dev]
            docs = [x for x in docs if x]
            stage_word = {"조제": "조제", "충전": "충전", "포장": "포장"}
            txt = "1) 충전 수율 일탈로 11항 일탈 관련 기록%s 참고. (수율의 최대, 최소, 평균 계산에서 제외함.)" % (
                "(%s)" % ", ".join(dict.fromkeys(docs)) if docs else "")
            others = [l for l in yield_out if l not in yield_dev]
            if others:
                issues.append(("7", ", ".join(others), "수율이 자가기준을 벗어났으나 11항에 일탈 기록이 없음 — 비고 '1)' 를 노랑으로 두었으니 일탈 기록을 확인하세요"))
            E.set_para_text(note, txt)
            _small(note._p)
        else:
            E.set_para_text(note, "")
    log("7항: 이탈 %s / 일탈 %s" % (yield_out, yield_dev))

    # ---------- 8.1 공급업체 평가 (8.1.1 주원료 · 8.1.2 공급망 · 8.1.3 부원료·포장자재) ----------
    def _norm_date(v):
        m = re.findall(r"\d+", str(v or ""))
        return "%s.%02d.%02d" % (m[0], int(m[1]), int(m[2])) if len(m) >= 3 else ""

    def _cols(table, *groups):
        """{이름: 열 번호} — 머리행 낱말로 짚는다. 열 차례는 제품·연도마다 다르다."""
        head = [re.sub(r"\s+", "", E.cell_text(c)) for c in E.raw_cells(table.rows[0])]
        out = {}
        for name, words in groups:
            out[name] = next((i for i, h in enumerate(head) if any(w in h for w in words)), None)
        return out

    def _key(v):
        return re.sub(r"[^a-z0-9가-힣]", "", str(v or "").lower())

    def _supplier(code, docno="", vendor=""):
        """공급업체 목록에서 그 줄 — 원료코드 → 평가문서번호 → 업체명 차례로 짚는다.

        자재 목록에는 자재코드 칸이 없어(2026 Rev.26) 코드로는 못 찾는다. 전년도 결재본에서
        이어받은 평가문서번호(VAR-P-Linhardt 등)가 그 줄을 가리키는 열쇠가 된다.
        """
        both = list(data.suppliers_raw) + list(data.suppliers_mat)
        code, docno, vendor = (code or "").strip(), _key(docno), _key(vendor)
        tests = [lambda r_: bool(code) and code == str(r_.get("원료코드") or "").strip(),
                 lambda r_: bool(docno) and _key(r_.get("문서번호")) == docno,
                 lambda r_: bool(docno) and _key(r_.get("문서번호")) and
                 (_key(r_.get("문서번호")).startswith(docno) or docno.startswith(_key(r_.get("문서번호")))),
                 lambda r_: bool(vendor) and len(vendor) >= 4 and _key(r_.get("공급업체명")) and
                 (vendor in _key(r_.get("공급업체명")) or _key(r_.get("공급업체명")).startswith(vendor))]
        for k, test in enumerate(tests):
            got = [r_ for r_ in both if test(r_)]
            if len(got) > 1:                      # 여럿이면 가장 최근에 평가한 줄
                got = [max(got, key=lambda r_: _norm_date(r_.get("평가승인일")))]
            if got:
                return got[0], (k == 0)           # 원료코드로 찾았는지
        return None, False

    def _put(cells, col, text):
        if col is None or col >= len(cells) or not text:
            return 0
        if E.cell_text(cells[col]).strip() == str(text).strip():
            return 0
        E.set_cell(cells[col], *str(text).split("\n"))
        return 1

    def _company(text):
        """'FUAN … CO., LTD / China' → 'FUAN … CO., LTD' (나라 이름은 뗀다)."""
        return re.split(r"\s*/\s*", str(text or ""))[0].strip()

    upd81, grade_odd = 0, []
    # 빈 공양식의 8.1.1·8.1.2(·8.1.3)는 관리번호 줄이 없다 — 전년도 결재본의 줄(관리번호·원/자재명·규격·제조원 …)을
    # 먼저 세우고, 아래에서 올린 공급업체 목록·공급망 마스터로 문서번호·완료일·업체를 갱신한다
    # (담당자 2026-09-06: "주원료명 모르면 작년 PQR 에서 가져오고 해당 항 첨부파일로 최신 정보를 가져오면 돼").
    olds81 = getattr(data, "prev_sections_all", None) or {}
    seeded81 = []                                    # [(항, 표, 전년도 표)] — 갱신이 끝난 뒤 남은 칸을 노랑으로
    for prefix in ("8.1.1", "8.1.2", "8.1.3"):
        grids = olds81.get(prefix) or []
        for i, tb in enumerate(_tables(document, prefix)):
            if i < len(grids) and not _has_codes(tb):
                n = _seed_rows_by_header(tb, grids[i])
                if n:
                    log("%s: 전년도 결재본에서 %d줄을 세움 — 올린 자료로 갱신" % (prefix, n))
                    seeded81.append((prefix, tb, grids[i]))
    # 6항에 올린 공 기록서(제조·충전·포장)의 원/자재 — 코드가 R 로 시작하면 주원료(8.1.1), 나머지 부원료·포장자재(8.1.3).
    # 표에 없는 코드만 줄을 보태고, 제조원·평가문서·완료일은 아래에서 공급업체 목록으로 채운다 (담당자 2026-09-06).
    batch_mats = {}
    for src in (getattr(data, "batch", None), getattr(data, "batch_exp", None)):
        for group, mats in ((src or {}).get("materials") or {}).items():
            have = {m["code"] for m in batch_mats.get(group, [])}
            batch_mats.setdefault(group, []).extend(m for m in mats if m["code"] not in have)
    if batch_mats:
        for prefix, groups in (("8.1.1", ("주원료",)), ("8.1.3", ("부원료", "포장자재"))):
            tables = _tables(document, prefix)
            if not tables:
                continue
            for gi, group in enumerate(groups):
                tb = tables[gi] if len(tables) > gi else tables[0]
                오기 = []
                n = _add_material_rows(tb, batch_mats.get(group) or [], 오기)
                for 옛, 새 in 오기:
                    log("%s: 관리번호 %s → %s (공 기록서와 한 자 차이라 오기로 봄)" % (prefix, 옛, 새))
                    issues.append((prefix, "%s → %s" % (옛, 새),
                                   "표의 관리번호가 공 기록서와 한 자만 달라 오기로 보고 고쳤습니다(노랑) — 확인 필요"))
                if n:
                    log("%s: 공 기록서의 %s %d줄을 보탬 — 공급업체 목록으로 갱신" % (prefix, group, n))
                    issues.append((prefix, ", ".join(m["code"] for m in batch_mats.get(group) or []),
                                   "공 기록서에 있는 %s 를 표에 보탰습니다(노랑) — 규격·제조원·평가 문서를 확인하세요" % group))
    for prefix in ("8.1.1", "8.1.3"):
        for tb in _tables(document, prefix):
            col = _cols(tb, ("code", ("관리번호", "코드")), ("name", ("원/자재명", "원자재명", "자재명", "원료명")),
                        ("maker", ("제조원", "제조소")), ("doc", ("문서번호", "평가문서")),
                        ("result", ("평가결과",)), ("day", ("완료일", "승인일")), ("spec", ("규격",)))
            if col["code"] is None:
                continue
            for row in tb.rows[1:]:
                cells = E.raw_cells(row)
                if len(cells) <= col["code"]:
                    continue
                take = lambda k: E.cell_text(cells[col[k]]).strip() if col[k] is not None and col[k] < len(cells) else ""
                got, by_code = _supplier(take("code"), take("doc"), take("maker"))
                if not got:
                    continue
                grade = str(got.get("평가등급") or "").strip().upper()
                if by_code or not take("maker"):
                    # 자재는 코드가 아니라 평가문서번호로 찾으므로, 목록의 업체명이 결재본의
                    # 제조원 이름과 다를 수 있다(‘린하르트 GmbH (Pausa) Linhardt GmbH (Pausa)’).
                    # 코드로 정확히 찾았을 때만 덮어쓰고, 아니면 쓰여 있는 이름을 둔다.
                    upd81 += _put(cells, col["maker"], _company(got.get("공급업체명")))
                upd81 += _put(cells, col["doc"], (got.get("문서번호") or "").strip())
                upd81 += _put(cells, col["day"], _norm_date(got.get("평가승인일")))
                if grade in ("A", "B"):
                    upd81 += _put(cells, col["result"], "적합")
                elif grade:
                    grade_odd.append("%s(%s등급)" % (E.cell_text(cells[col["code"]]).strip(), grade))
                # 포장자재(P 코드)의 규격은 자사규격이다 — 원료처럼 공정서(KP·USP …)가 없다
                # (담당자 2026-09: "설명서, 케이스는 자사규격인데").
                code_text = take("code")
                if code_text.upper().startswith("P") and col["spec"] is not None \
                        and col["spec"] < len(cells) and not take("spec"):
                    upd81 += _put(cells, col["spec"], "자사규격")
                if col["name"] is not None and col["name"] < len(cells) and not E.cell_text(cells[col["name"]]).strip():
                    korean = re.split(r"(?=[A-Za-z])", str(got.get("공급되는 품목") or ""), 1)[0].strip()
                    upd81 += _put(cells, col["name"], korean)
    for tb in _tables(document, "8.1.2"):
        col = _cols(tb, ("code", ("관리번호", "코드")), ("maker", ("제조업체", "제조 업체")),
                    ("supply", ("공급", "납품")), ("other", ("기타",)))
        if col["code"] is None:
            continue
        for row in tb.rows[1:]:
            cells = E.raw_cells(row)
            if len(cells) <= col["code"]:
                continue
            chain = data.api_chain.get(E.cell_text(cells[col["code"]]).strip())
            if not chain:
                continue
            maker = _company(chain.get("manufacturer"))
            links = [_company(x) for x in (chain.get("chain") or [])]
            # '제조소 납품' 은 제조소에서 바로 받는다는 뜻 — 납품 업체 칸에 제조소를 적는다
            # '제조소 납품' 은 제조소에서 바로 받는다는 뜻 — 업체 이름이 아니다
            links = [maker if x.startswith("제조소") else x for x in links]
            upd81 += _put(cells, col["maker"], maker)
            upd81 += _put(cells, col["supply"], links[0] if links else "")
            upd81 += _put(cells, col["other"], ", ".join(links[1:]))
    if grade_odd:
        issues.append(("8.1", ", ".join(sorted(set(grade_odd))),
                       "공급업체 평가등급이 A·B 가 아니어서 평가결과 칸을 비웠습니다 — 확인해 적으세요"))
    log("8.1 공급업체 정보 채움: %d 칸" % upd81)
    # 올린 자료로 갱신되지 않아 아직 전년도 값인 칸은 노랑으로 — 담당자가 확인해야 할 자리다
    # (담당자 2026-09-07: "전년도 PQR 정보로 가져오고 노랑 MARK 해").
    남은 = {}
    for prefix, tb, grid in seeded81:
        got = _mark_carried_cells(tb, grid)
        남은[prefix] = 남은.get(prefix, 0) + got
    for prefix, got in sorted(남은.items()):
        if got:
            log("%s: 전년도 값 그대로인 칸 %d개를 노랑으로 표시" % (prefix, got))
            issues.append((prefix, "", "공 기록서·공급업체 목록으로 갱신되지 않아 전년도 결재본 값을 그대로 "
                                       "옮긴 칸이 %d개 있습니다(노랑) — 올해 자료로 확인하세요" % got))

    # ---------- 8.2 시험성적 ----------
    def group_fill(table, records):
        """records: [(원료 코드, 원료명, 시험번호, [Lot ...]), ...] — Lot 하나가 한 줄.

        2026 결재본 차림새: **연번은 줄마다** 매기고, **원료 코드·원료명은 같은 원료끼리**,
        시험 성적 번호는 같은 번호끼리 세로로 합친다. 결과(적합)와 적용 Lot 은 줄마다 적는다.
        담당자 지적: "같은 원료 및 코드는 셀병합해줘."
        """
        rows = [(code, item, test, lot) for code, item, test, lots in records for lot in lots]
        if not rows:
            return
        f, l = E.fit_rows(table, 1, len(table.rows) - 1, len(rows))
        prev = None
        for i, (code, item, test, lot) in enumerate(rows):
            r = E.raw_cells(table.rows[f + i])
            # 열마다 '같은 묶음' 의 기준이 다르다 — 원료명은 원료가 바뀌면, 시험번호는 원료나
            # 번호가 바뀌면 새로 시작한다.
            keys = (code, (code, item), (code, test))
            E.set_cell(r[0], str(i + 1)); E.set_vmerge(r[0], False)
            for k, value in ((1, code), (2, item), (3, test)):
                head = prev is None or keys[k - 1] != prev[k - 1]
                E.set_cell(r[k], value if head else "")
                E.set_vmerge(r[k], "restart" if head else None)
            E.set_cell(r[4], lot); E.set_vmerge(r[4], False)
            E.set_cell(r[5], "적합"); E.set_vmerge(r[5], False)
            prev = keys

    def material_names(table):
        """결재본 8.2 표에 적힌 {원료 코드: 원료명} — 코드에 붙일 한글 이름은 여기에만 있다."""
        out = {}
        for row in table.rows[1:]:
            cells = E.raw_cells(row)
            if len(cells) > 2:
                code = E.cell_text(cells[1]).strip()
                if code:
                    out.setdefault(code, E.cell_text(cells[2]).strip())
        return out

    t821 = _tables(document, "8.2.1")
    if t821 and data.raw_tests:
        # 주원료가 둘 이상인 제품이 있다(디겐타안연고: 겐타마이신황산염·플루오로메톨론). 예전에는
        # 결재본 첫 행의 코드 하나만 채워 나머지 주원료가 표에서 통째로 빠졌다.
        names = material_names(t821[0])
        for tbl, lots in ((t821[0], dom), (t821[1] if len(t821) > 1 else None, exp)):
            if tbl is None:
                continue
            # 표에 이름이 없으면 소제목('8.2.1.1 포비돈 K-25')의 이름을 쓴다 — 담당자 2026-09-07:
            # "원료명에 왜 사선이야, 포비돈 K-25 라고 기재하면 되잖아"
            titled = _heading_material_name(document, tbl)
            recs = []
            for code, test, ls in data.raw_tests:
                mine = [l for l in lots if l in ls]
                if mine:
                    recs.append((code, names.get(code) or titled, test, mine))
            missing = sorted({c for c, item, _t, _l in recs if not item})
            if missing:
                issues.append(("8.2.1", ", ".join(missing),
                               "전년도 결재본에 없는 원료 코드입니다 — 원료명을 확인해 적으세요"))
            group_fill(tbl, recs)
    t822 = _tables(document, "8.2.2")
    if t822 and data.pkg_tests:
        tbl = t822[0]
        base_rows = {E.cell_text(E.raw_cells(r)[1]).strip(): E.cell_text(E.raw_cells(r)[2]).strip()
                     for r in tbl.rows[1:] if len(E.raw_cells(r)) > 2}
        recs = []
        for code, test, ls in data.pkg_tests:
            ls = [x for x in ls if x in dom + exp]
            if ls:
                item = base_rows.get(code) or ("튜브 (수출용)" if ls[0] in exp else "튜브 (내수용)")
                recs.append((code, item, test, ls))
        group_fill(tbl, recs)                  # 1차 포장 자재도 같은 차림새다 (담당자 지적)

    # ---------- 9항 ----------
    def spec_of(t91, words):
        """그 시험항목의 허용기준 글. '자가)·허가)' 가 붙은 칸을 찾는다 — 칸 수는 줄마다 다르다
        (금속성이물 줄은 기준 뒤에 '합계·개개' 칸이 하나 더 붙어 cells[-2] 가 기준이 아니다)."""
        for row in t91.rows:
            cells = [E.cell_text(c) for c in E.raw_cells(row)]
            if any(all(w in c for w in words) for c in cells[:2]):
                return next((c for c in cells if D.PREFIX.search(c or "")), cells[-2])
        return ""

    def has_row(t91, words):
        """그 시험항목 줄이 9.1 표에 있는가 — 없으면 이 제품에 없는 시험이다."""
        for row in t91.rows:
            cells = [E.cell_text(c) for c in E.raw_cells(row)]
            if any(all(w in c for w in words) for c in cells[:2]):
                return True
        return False

    def rec(lot, key):
        return (data.coa.get(lot) or {}).get(key) or {}

    # Cpk 한계는 이 제품의 9.1 허용기준에서 읽는다 — 퀴노비드 숫자를 다른 제품에 쓰지 않는다
    limits = dict(DEFAULT_LIMITS)
    # 성분별 함량 규격은 완제 성적서(9.2.4)에 성분마다 적혀 있다 — {성분: (하한, 상한)}
    parts = {}
    for lot in dom + exp:
        for a in (rec(lot, "924").get("assays") or []):
            try:
                if a.get("part") and a["part"] not in parts:
                    parts[a["part"]] = (float(a["lo"]), float(a["hi"]))
            except (TypeError, ValueError, KeyError):
                pass
    limits["assay_parts"] = parts
    t91_all = _tables(document, "9.1")
    if t91_all:
        found = parse_limits({"particle": spec_of(t91_all[0], ("입자도",)),
                              "assay": spec_of(t91_all[0], ("함량",)),
                              "metal": spec_of(t91_all[0], ("금속성",))})
        limits.update(found)
        # 이 제품에 없는 시험(점안제의 금속성이물 등)까지 '못 읽었다' 고 알리면 잔소리가 된다 —
        # 표에 줄이 있는 시험만 본다. 그리고 못 읽은 한계는 **기본값을 쓰지 않고** 그 항목 Cpk 를
        # 건너뛴다 — 퀴노비드안연고 숫자가 다른 제형으로 새면 안 된다 (담당자 2026-09-08).
        낱말 = {"particle": ("입자도",), "assay": ("함량",), "metal": ("금속성",)}
        이름 = {"particle": "입자도", "assay": "함량", "metal": "금속성이물"}
        missing = [k for k in ("particle", "assay", "metal")
                   if k not in found and has_row(t91_all[0], 낱말[k])]
        for k in ("particle", "assay", "metal"):
            if k not in found:
                limits[k] = (None, None) if k == "assay" else None
        if missing:
            issues.append(("9.1", "", "허용기준에서 %s 한계를 읽지 못해 그 항목은 Cpk 를 계산하지 않았습니다 — "
                                      "9.1 표의 허용기준 칸을 채우고 재작성하세요"
                           % ", ".join(이름[k] for k in missing)))
    log("Cpk 한계: %s" % limits)

    def app(lots):
        for l in lots:
            for k in ("924", "923", "922"):
                a = rec(l, k).get("appearance")
                if a:
                    return a
        return ""

    def assay_by_part(lots):
        """{성분 이름: [Lot 별 함량]} — 주성분이 둘 이상인 제품은 성적서에 성분마다 함량 줄이 있다."""
        out = {}
        for lot in lots:
            for a in rec(lot, "924").get("assays") or []:
                value = _num(a.get("value"))
                if a.get("part") and value is not None:
                    out.setdefault(a["part"], []).append(float(value))
        return out

    def numbers(lots):
        g = lambda k, key: [float(_num(rec(l, k).get(key))) for l in lots if _num(rec(l, k).get(key)) is not None]
        return {"parts": assay_by_part(lots),
                "ms": g("924", "metal_total"), "mi": g("924", "metal_each"), "pt": g("924", "particle"),
                "pa": g("924", "mass_avg"), "pi": g("924", "mass_each_min"), "ct": g("924", "assay"),
                "fa": g("923", "mass_avg"),
                "flo": [_rng(rec(l, "923").get("mass_each"))[0] for l in lots if _rng(rec(l, "923").get("mass_each"))],
                "fhi": [_rng(rec(l, "923").get("mass_each"))[1] for l in lots if _rng(rec(l, "923").get("mass_each"))]}

    def bio_text(lot):
        b = rec(lot, "922").get("bioburden") or ""
        m = re.match(r"(\d+)\s*(CFU|FU)?/g\s*(미만|이하)?", b)
        return ("%s %s" % (m.group(1), m.group(3) or "미만")) if m else (b or "확인 필요")

    # 각주 번호는 실제로 다는 것만 세어 매긴다 — 생균수 각주가 없는 제품에서 1) 없이 2) 로
    # 시작하면 안 된다(디겐타안연고 2026).
    odd_bio = [(l, rec(l, "922").get("bioburden")) for l in dom + exp
               if bio_text(l) != "10 미만" and rec(l, "922")]
    note_bio = 1 if odd_bio else None
    note_mass = 2 if odd_bio else 1

    def fill_91(t91, lots, is_dom):
        n = numbers(lots)
        if not n["ct"]:
            issues.append(("9", "", "완제 성적서에서 함량을 읽지 못함")); return n
        res = {}
        current_item = ""
        assay_parts = []          # 주성분이 둘 이상인 제품(디겐타안연고: 플루오로메톨론·겐타마이신황산염)
        # 값이 없는 글 결과(확인 시험 '검액은 표준액과 동일한 주피크 유지시간을 나타냄' 등)는 전년도 결재본의
        # 같은 기준 줄에서 옮긴다 — 숫자가 든 결과는 옮기지 않는다(담당자 2026-09-06: "9.1.1 시험결과들은 왜 사선인지?")
        prior = _prior_91(data, 0 if is_dom else 1)
        carried = 0
        for ri, row in enumerate(t91.rows):
            cells = E.raw_cells(row)
            if ri == 0 or len(cells) < 3:
                continue
            label = re.sub(r"\s+", "", E.cell_text(cells[-3])) if len(cells) >= 4 else ""
            crit_i = next((i for i, c in enumerate(cells) if D.PREFIX.search(E.cell_text(c) or "")), None)
            crit_text = E.cell_text(cells[crit_i]) if crit_i is not None else E.cell_text(cells[-2])
            if "평균" not in label and "개개" not in label:
                # 칸이 넷뿐인 표(디겐타안연고)는 평균·개개 칸이 따로 없고 허용기준 글에 적혀 있다
                # — "허가) 평균 : 표시량(4.0 g) 이상" / "허가) 개개 : 3.60 g 이상".
                if re.search(r"평균\s*[:：]", crit_text):
                    label += "평균"
                elif re.search(r"개개\s*[:：]", crit_text):
                    label += "개개"
            named = re.sub(r"\s+", "", "".join(E.cell_text(c) for c in cells[:-2]))
            named = re.sub(r"(평균|개개)$", "", named)
            if named:
                current_item = named
            item = current_item
            val = None
            if "성상" in item:
                val = app(lots)
            elif "생균수" in item:
                vals = [bio_text(l) for l in lots]
                common = max(set(vals), key=vals.count)
                m = re.match(r"(\d+)\s*(미만|이하)", common)
                val = ("%s CFU/g %s" % (m.group(1), m.group(2))) if m else common
                if any(v != common for v in vals) and note_bio:
                    val += "%d)" % note_bio
            elif "평균" in label and n["fa"] and "질량" in item and not n.get("_fill_avg"):
                val = _avg_text(n["fa"], "%.1f" if is_dom else "%.2f", unit_on_avg=not is_dom)
                n["_fill_avg"] = True
            elif "개개" in label and n["flo"] and "질량" in item and not n.get("_fill_each"):
                val = "%.2f ~ %.2fg" % (min(n["flo"]), max(n["fhi"])); n["_fill_each"] = True
            elif "확인" in item and (is_dom or not re.search(r"[123]\)", E.cell_text(cells[-2]))):
                val = _ident_result(lots, rec, prior, crit_text)      # 올해 성적서 '확인시험 적합' + 전년도 문안
            elif "확인" in item and not is_dom:
                crit = E.cell_text(cells[-2])
                if "1)" in crit:
                    vs = {rec(l, "924").get("ident_color") for l in lots} - {None}
                    val = sorted(vs)[0] if vs else None
                elif "2)" in crit:
                    vs = {rec(l, "924").get("ident_precip") for l in lots} - {None}
                    val = sorted(vs)[0] if vs else None
                elif "3)" in crit:
                    pairs = [re.findall(r"(\d{3})nm", rec(l, "924").get("uv_max") or "") for l in lots]
                    pairs = [p_ for p_ in pairs if len(p_) == 2]
                    if pairs:
                        a = sorted(int(p_[0]) for p_ in pairs); b = sorted(int(p_[1]) for p_ in pairs)
                        fmt = lambda xs: ("%d ~ %dnm" % (xs[0], xs[-1])) if xs[0] != xs[-1] else "%dnm" % xs[0]
                        val = "%s 및 %s에서 흡수극대를 나타냄" % (fmt(a), fmt(b))
            elif "튜브인쇄" in item:
                val = "인쇄상태가 양호하며 제조번호 및 사용기한의 압인상태가 명확히 식별 가능함"
            elif "기밀도" in item:
                leaks = {rec(l, "924").get("leak") for l in lots} - {None}
                val = "메틸렌블루시액 침투 없음" if not n.get("_leak") else (sorted(leaks)[0] if leaks else "메틸렌블루시액 침투 없이 양호")
                n["_leak"] = True
            elif "금속성이물" in item and n["ms"]:
                # 모든 Lot 이 같은 값이면 평균·범위를 적지 않는다 — 2026 결재본의 주석 그대로
                # ("모든 Lot 의 시험결과가 0개(매)으로 동일하여 최댓값, 최솟값, 평균은 별도로 작성하지 않음").
                same = len(set(n["ms"])) == 1 and len(set(n["mi"])) == 1
                val = ("50 ㎛ 이상\n: Av. %.2f개(%.0f ~ %.0f개)\n개개 중 8개 초과 : %.0f매" % (sum(n["ms"]) / len(n["ms"]), min(n["ms"]), max(n["ms"]), max(n["mi"]))) if (len(lots) >= 3 and not same) else \
                      ("50 ㎛ 이상 : %.0f개\n개개 중 8개 초과 : %.0f매" % (max(n["ms"]), max(n["mi"])))
            elif "입자도" in item and n["pt"]:
                val = "Av. %.2f㎛ 이하\n(%.2f ~ %.2f㎛ 이하)" % (sum(n["pt"]) / len(n["pt"]), min(n["pt"]), max(n["pt"]))
            elif "함량" in item and n["ct"]:
                # 성분 이름이 적힌 줄이면 그 성분의 값만 쓴다 — 성분마다 규격도 결과도 다르다.
                part = assay_component(E.cell_text(cells[-2]))
                vals = n["parts"].get(part) if part else None
                if part and not vals:
                    vals = next((v for k, v in n["parts"].items() if k in part or part in k), None)
                if part and vals is None:
                    assay_parts.append(part)                      # 성분별 값을 못 찾았다 — 짚어 준다
                use = vals or n["ct"]
                val = "Av. %.1f%%\n(%.1f ~ %.1f%%)" % (sum(use) / len(use), min(use), max(use))
            elif "평균" in label and n["pa"] and "질량" in item:
                val = _avg_text(n["pa"], "%.2f")
            elif "개개" in label and n["pi"] and "질량" in item:
                val = "%.2f ~ %.2fg 이상%d)" % (min(n["pi"]), max(n["pi"]), note_mass)
            elif "무균" in item:
                st = {rec(l, "924").get("sterility") for l in lots} - {None}
                val = sorted(st)[0] if st else "음성"
            elif "포장규격" in item:
                val = "각 규격에 적합함"
            if val is None and crit_text.strip():
                old = prior.get(_crit_key(crit_text))
                if old and not re.search(r"\d", old):
                    val = old
                    carried += 1
            if val is not None:
                target = cells[-1]
                # 결과 칸 왼쪽에 글 없는 쪽칸이 붙어 있으면(기준 칸이 아니면) 합쳐서 사선이 생기지 않게 한다
                if crit_i is not None and len(cells) - 2 > crit_i and not E.cell_text(cells[-2]).strip():
                    E.clear_diag(cells[-2])
                    E.merge_right(cells[-2], cells[-1])
                    target = cells[-2]
                E.clear_diag(target)
                E.set_cell(target, *val.split("\n"))
                res[ri] = val
        if carried:
            log("9.1항: 글 결과 %d줄은 전년도 결재본에서 옮김(확인 시험 등)" % carried)
        if assay_parts:
            # 성적서에서 그 성분의 함량을 찾지 못했다 — 다른 성분 값이 들어가 있으니 짚는다.
            issues.append(("9.1", ", ".join(assay_parts),
                           "이 성분의 함량을 성적서에서 찾지 못해 다른 값이 들어가 있습니다 — "
                           "원본에서 확인해 적으세요"))
        return n
    t91 = _tables(document, "9.1")
    # 2026 양식(공정별 9.2 표)이면 9.1·9.2 를 모두 머리글·허용기준으로 짚어 채운다.
    rules = D.criteria(t91[0]) if t91 else []
    pairs = D.tables_92(document)
    # 수출 Lot 이 있어도 머리글 기반으로 채운다 — 내수용·수출용 표를 제목으로 갈라 각각 채운다(2026-09-06 퀴노비드:
    # 자리 기반으로 떨어지면 열이 밀려 9.2 의 확인·질량 칸이 비거나 옆 칸에 들어갔다)
    labelled = bool(pairs) and all(st for _, st in pairs) and all(D.simple(t) for t, _ in pairs)
    if labelled:
        n_dom, n_exp = numbers(dom), (numbers(exp) if exp else {})
    else:
        n_dom = fill_91(t91[0], dom, True) if t91 else {}
        n_exp = fill_91(t91[1], exp, False) if len(t91) > 1 and exp else {}

    # 9.2 세부표 — 항 아래 표들을 머리행 낱말로 고른다
    def fill_detail(table, first, lots, setter, n_summary=5):
        last = len(table.rows) - 1 - n_summary
        f, l = E.fit_rows(table, first, last, max(1, len(lots)))
        for i, lot in enumerate(lots):
            setter(E.raw_cells(table.rows[f + i]), i + 1, lot)
        return f, l

    def put(table, ri, idx, *vals):
        c = E.raw_cells(table.rows[ri])
        for k, v in zip(idx, vals):
            E.set_cell(c[k], v)

    def fill_92(section, lots, n, is_dom):
        tabs = _tables(document, section)
        if not tabs:
            return
        appearance = app(lots)
        t_bio = _by_header(tabs, "생균수")
        t_fill = _by_header(tabs, "질량", "기밀도")
        t_tube = _by_header(tabs, "튜브인쇄")
        t_metal = _by_header(tabs, "금속성이물")
        t_ident = _by_header(tabs, "확인")
        t_numb = _by_header(tabs, "입자도")
        if t_numb is None and t_ident is not None and "입자도" in _text(t_ident._tbl):
            t_numb = t_ident                      # 확인 블록 + 수치 블록이 한 표 (내수용)
        t_last = _by_header(tabs, "무균")
        if t_bio:
            vals_ = [bio_text(l) for l in lots]
            common_ = max(set(vals_), key=vals_.count)
            fill_detail(t_bio, 1, lots, lambda c, no, lot: [E.set_cell(c[0], str(no)), E.set_cell(c[1], lot), E.set_cell(c[2], appearance),
                                                            E.set_cell(c[3], bio_text(lot) + ("1)" if bio_text(lot) != common_ else ""))])
        if t_fill:
            def s_fill(c, no, lot):
                r = rec(lot, "923")
                E.set_cell(c[0], str(no)); E.set_cell(c[1], lot); E.set_cell(c[2], appearance)
                E.set_cell(c[3], (r.get("mass_avg") or "").replace("g", "")); E.set_cell(c[4], (r.get("mass_each") or "").replace("g", ""))
                E.set_cell(c[5], r.get("leak") or "메틸렌블루시액 침투 없음")
            f, l = fill_detail(t_fill, 2, lots, s_fill)
            if n.get("fa"):
                fmt = "%.1f" if is_dom else "%.2f"
                put(t_fill, l + 1, (2, 3), fmt % max(n["fa"]), "%.2f" % max(n["fhi"]))
                put(t_fill, l + 2, (2, 3), fmt % min(n["fa"]), "%.2f" % min(n["flo"]))
                put(t_fill, l + 3, (2,), "%.2f" % (sum(n["fa"]) / len(n["fa"])))
        if t_tube:
            fill_detail(t_tube, 1, lots, lambda c, no, lot: [E.set_cell(c[0], str(no)), E.set_cell(c[1], lot), E.set_cell(c[2], "인쇄상태가 양호하며 제조번호 및 사용기한의 압인상태가 명확히 식별 가능함")])
        cpk = {}
        if t_metal and n.get("ms"):
            f, l = fill_detail(t_metal, 2, lots, lambda c, no, lot: [E.set_cell(c[0], str(no)), E.set_cell(c[1], lot), E.set_cell(c[2], appearance), E.set_cell(c[3], rec(lot, "924").get("metal_total") or ""), E.set_cell(c[4], rec(lot, "924").get("metal_each") or "")])
            ms, mi = n["ms"], n["mi"]
            put(t_metal, l + 1, (2, 3), "%.0f" % max(ms), "%.0f" % max(mi))
            put(t_metal, l + 2, (2, 3), "%.0f" % min(ms), "%.0f" % min(mi))
            if qc.cpk_applies(len(lots)):
                put(t_metal, l + 3, (2, 3), "%.2f" % (sum(ms) / len(ms)), "%.2f" % (sum(mi) / len(mi)))
                cpk["metal"] = cpk_uni(ms, limits["metal"])
                if cpk["metal"] is not None:
                    put(t_metal, l + 4, (2,), "%.2f" % cpk["metal"]); put(t_metal, l + 5, (2,), "충분" if cpk["metal"] >= 1 else "부족")
            else:
                put(t_metal, l + 3, (2, 3), "%.0f" % (sum(ms) / len(ms)), "%.0f" % (sum(mi) / len(mi)))
        if t_ident is not None and t_ident is t_numb:            # 한 표에 확인 + 수치 두 블록 (내수용)
            t = t_ident
            firsts = [i for i, tr in enumerate(t._tbl.findall(qn("w:tr"))) if _text(tr.findall(qn("w:tc"))[0]).startswith("최댓값")]
            fa_, la_ = E.fit_rows(t, 2, firsts[0] - 1, len(lots))
            for i, lot in enumerate(lots):
                c = E.raw_cells(t.rows[fa_ + i]); E.set_cell(c[0], str(i + 1)); E.set_cell(c[1], lot)
                E.set_cell(c[2], "검액은 표준액과 동일한 주 피크 유지시간을 나타냄"); E.set_cell(c[3], "검액과 표준액의 주 피크 UV spectrum은 동일함")
            base2 = la_ + 6
            fb_, lb_ = E.fit_rows(t, base2 + 2, len(t.rows) - 6, len(lots))
            _fill_numbers(t, fb_, lb_, lots, n, cpk, rec, put, is_dom, limits)
        else:
            if t_ident:
                def s_ident(c, no, lot):
                    E.set_cell(c[0], str(no)); E.set_cell(c[1], lot)
                    if is_dom:
                        E.set_cell(c[2], "검액은 표준액과 동일한 주 피크 유지시간을 나타냄"); E.set_cell(c[3], "검액과 표준액의 주 피크 UV spectrum은 동일함")
                    else:
                        r = rec(lot, "924")
                        E.set_cell(c[2], r.get("ident_color") or "확인 필요"); E.set_cell(c[3], r.get("ident_precip") or "확인 필요")
                        E.set_cell(c[4], r.get("uv_max") or "확인 필요")
                fill_detail(t_ident, 2, lots, s_ident)
            if t_numb and n.get("ct"):
                f, l = E.fit_rows(t_numb, 2, len(t_numb.rows) - 6, len(lots))
                _fill_numbers(t_numb, f, l, lots, n, cpk, rec, put, is_dom, limits)
        if t_last:
            def s_last(c, no, lot):
                r = rec(lot, "924")
                E.set_cell(c[0], str(no)); E.set_cell(c[1], lot)
                E.set_cell(c[2], r.get("sterility") or "음성"); E.set_cell(c[3], r.get("leak") or "메틸렌블루시액 침투 없이 양호")
                E.set_cell(c[4], "각 규격에 적합함")
            fill_detail(t_last, 1, lots, s_last)
        return cpk
    def bio_full(lot):
        m = re.match(r"(\d+)\s*(미만|이하)", bio_text(lot))
        return ("%s CFU/g %s" % (m.group(1), m.group(2))) if m else bio_text(lot)

    def _plain(value, fmt=None):
        """성적서에 적힌 숫자를 적힌 자릿수 그대로 (5.10 을 5.1 로, 5.00 을 5 로 줄이지 않는다).

        담당자 2026-09-07: "소숫점 자리를 왼쪽처럼 통일시켜줘" — 자릿수는 열마다
        detail92.unify_decimals 가 가장 자세한 값에 맞춘다.
        """
        text = str(value if value is not None else "").strip()
        m = re.search(r"-?\d+(?:\.\d+)?", text)
        if m is None:
            return text or None
        return (fmt % float(m.group())) if fmt else m.group()

    makers = []

    def maker_of(process, idx=0):
        return makers[idx](process) if len(makers) > idx else (lambda lab, lot, i: None)

    def fill_92_labelled(pairs, rules, lots, n, prior=None):
        """9.2 표를 머리글 이름으로 채운다 (2026 양식: 조제·충전·포장 공정별 표)."""
        parts = [r["part"] for r in rules if r["part"]]
        found = {}
        prior = prior or {}

        def maker(process):
            def value(lab, lot, i):
                r923, r924 = rec(lot, "923"), rec(lot, "924")
                mine = r923 if process == "충전" else r924
                part = next((p for p in parts if p and p in lab), "")
                if "성상" in lab:
                    return app(lots) or D.criterion_for(rules, process, "성상")
                if "바이오버든" in lab or "생균수" in lab:
                    return bio_full(lot)
                if "함량" in lab:
                    got = [a for a in (r924.get("assays") or [])
                           if not part or D.squeeze(a.get("part") or "") in part or part in D.squeeze(a.get("part") or "")]
                    v = _num(got[0].get("value")) if got else _num(r924.get("assay"))
                    if v is None and part:
                        found.setdefault("assay_miss", set()).add(part)
                    return ("%.1f" % float(v)) if v is not None else None
                if "입자도" in lab:
                    return _plain(r924.get("particle"))
                if "금속성이물" in lab:
                    # 9.1 표는 '합계' 라는 글 없이 줄만 갈라져 있다 — '개개' 가 아니면 합계 줄이다
                    return _plain(r924.get("metal_each") if "개개" in lab else r924.get("metal_total"))
                if "질량" in lab or "용량" in lab:
                    if "평균" in lab:
                        return _plain(mine.get("mass_avg"))
                    if "개개" in lab:
                        if process == "충전":
                            got = re.findall(r"\d+(?:\.\d+)?", str(mine.get("mass_each") or ""))
                            return (" ~ ".join(got[:2])) if len(got) >= 2 else _plain(mine.get("mass_each"))
                        return ("%s 이상" % _plain(mine.get("mass_each_min"))) if mine.get("mass_each_min") else None
                if "무균" in lab:
                    return r924.get("sterility") or D.criterion_for(rules, process, "무균")
                if "기밀도" in lab:
                    return mine.get("leak") or D.criterion_for(rules, process, "기밀도")
                if "튜브개봉" in lab:
                    return mine.get("tube_open")              # 변경관리로 더해진 항목 — 성적서에 없는 Lot 은 빈 칸(사선)
                if "튜브인쇄" in lab:
                    return mine.get("tube_print") or _as_result(D.criterion_for(rules, process, "튜브인쇄", part))
                if "확인" in lab:
                    sub = (re.search(r"\d\)", lab) or [""])[0] if re.search(r"\d\)", lab) else ""
                    crit = D.criterion_for(rules, process, "확인", part, sub)
                    return _ident_result([lot], rec, prior, crit)  # 올해 성적서 '확인시험 적합' + 전년도 문안
                if "포장규격" in lab:
                    return D.criterion_for(rules, process, "포장규격", part)
                # 코드에 박아 두지 않은 항목(pH·삼투압·비중 …)은 성적서의 시험항목 표에서 찾는다
                # (담당자 2026-09-08: "9.2.1 시험 압축 파일을 참고해서 9.2.1 표를 작성하면 돼").
                return _item_value(lot, process, lab)
            return value

        def _item_value(lot, process, lab):
            """성적서 시험항목 표에서 이 열에 맞는 결과 — 없으면 None."""
            keys = ("921", "922", "923") if process != "포장" else ("924",)
            if process == "충전":
                keys = ("923", "921", "922")
            want = D.squeeze(re.sub(r"[(（][^)）]*[)）]", "", lab))
            want = re.sub(r"\d+\)", "", want)
            if not want:
                return None
            for key in keys + ("924", "923", "922", "921"):
                for name, one in (rec(lot, key).get("items") or {}).items():
                    flat = D.squeeze(name)
                    if not flat:
                        continue
                    if flat == want or flat in want or want in flat:
                        value = (one.get("value") or "").strip()
                        if value and value not in ("N/A",):
                            return value
            return None

        def cpk_of(lab, texts):
            if not qc.cpk_applies(len(lots)):
                return None
            vals = [float(x) for x in (_num(t) for t in texts) if x is not None]
            if len(vals) != len(texts):
                return None
            if "함량" in lab:
                # 주성분이 둘이면 성분마다 규격이 다르다(디겐타: 플루오로메톨론 90~110, 겐타마이신
                # 90~120). 열 이름에 든 성분의 규격을 쓴다 — 2026-09 점검에서 둘 다 90~110 으로 계산됐다.
                lo, hi = limits["assay"]
                for part_name, (plo, phi) in (limits.get("assay_parts") or {}).items():
                    if part_name and re.sub(r"\s+", "", part_name) in re.sub(r"\s+", "", lab):
                        lo, hi = plo, phi
                        break
                got = cpk_bi(vals, lo, hi)
            elif "입자도" in lab:
                got = cpk_uni(vals, limits["particle"])
            elif "금속성이물" in lab:
                got = cpk_uni(vals, limits["metal"])
            else:
                return None
            if got is None:
                return None
            found[lab] = got
            return ("%.2f" % got, "충분" if got >= 1 else "부족")

        makers.append(maker)
        for table, process in pairs:
            D.fill(table, lots, maker(process), cpk=cpk_of)
        miss = found.pop("assay_miss", None)
        if miss:
            issues.append(("9.2", ", ".join(sorted(miss)),
                           "이 성분의 함량을 성적서에서 찾지 못해 칸을 비웠습니다 — 원본에서 확인해 적으세요"))
        out = {}
        parts = [n for n in (limits.get("assay_parts") or {}) if n]
        for lab, v in found.items():
            key = "assay" if "함량" in lab else "particle" if "입자도" in lab else "metal"
            if key == "assay" and len(parts) > 1:
                squeezed = re.sub(r"\s+", "", lab)
                part = next((n for n in parts if re.sub(r"\s+", "", n) in squeezed), "")
                if part:
                    key = "assay/%s" % part           # 16.1 표에 '함량(성분)' 으로 적힌다
            out[key] = v
        return out

    NUMERIC = re.compile(r"[\s\d.,~]+(?:\s*(?:이상|이하|미만|초과))?$")
    UNIT = re.compile(r"\d[\d.,]*\s*(㎛|um|μm|%|kg|mg|mL|g|개|매)")

    def unit_of(crit, item, sub):
        if "금속성" in item:
            return "매" if "개개" in sub else "개"
        m = UNIT.search(crit or "")
        return m.group(1) if m else ""

    def summarize(texts, crit, item, sub, unit):
        """9.1 결과 칸 글. 한림 결재본이 쓰는 차림새 그대로.

        자료가 범위면 범위, 허용기준이 한쪽만 정한 값(‘75㎛ 이하’·‘3.60 g 이상’)이면 그 쪽,
        모두 같으면 그 값, 그 밖이면 ‘Av. 평균(최솟값 ~ 최댓값)’.
        """
        texts = [t for t in texts if t]
        if not texts:
            return None
        if not all(NUMERIC.match(t) for t in texts):
            return max(set(texts), key=texts.count)          # 글로 적는 항목 — 가장 많이 나온 글
        tail = (" " + unit) if unit else ""
        if any("~" in t for t in texts):
            top, bottom, _ = D._stats("", texts)
            return "%s ~ %s%s" % (bottom, top, tail)
        if len(set(texts)) == 1:
            return "%s%s" % (texts[0].strip(), tail)
        if "금속성" in item:
            # 결재본은 금속성이물 결과를 평균 하나로 적는다 — 합계 '0 개', 개개 '0 매'
            # (담당자 2026-09-07: "금속성 이물 결과는 왼쪽 내용 참고해서 작성해 줘")
            mean = D._stats("", texts)[2]
            return ("%s%s" % (mean, tail)) if mean else None
        one_low = ("이상" in (crit or "") or any("이상" in t for t in texts)) and "평균" not in sub
        one_high = ("이하" in (crit or "") or any("이하" in t for t in texts)) and "평균" not in sub
        top, bottom, mean = D._stats("이상" if one_low else "이하" if one_high else "", texts)
        if one_low and not one_high:
            return "%s%s 이상" % (bottom, tail)
        if one_high and not one_low:
            return "%s%s 이하" % (top, tail)
        if top == bottom:
            return "%s%s" % (top, tail)
        # 결재본은 평균과 범위를 두 줄로 적는다 — "Av. 3.64 g" / "(3.60 ~ 3.67 g)"
        # (담당자 2026-09-07: "질량 용량도 왼쪽과 같은 형태로 작성해 줘")
        return "Av. %s%s\n(%s ~ %s%s)" % (mean, tail, bottom, top, tail)

    def fill_91_labelled(t91, rules, lots, n, maker, old_index=0):
        """9.1 결과 칸을 9.2 와 같은 판독값으로 채운다 (2026 양식).

        값이 없는 항(확인 시험처럼 '검액은 표준액과 동일한 주피크 유지시간을 나타냄' 같은 글 결과)은 전년도
        결재본의 같은 기준 줄에서 글 결과를 옮긴다 — 숫자가 든 결과는 옮기지 않는다(담당자 2026-09-06: "9.1.1
        시험결과들은 왜 사선인지?"). 결과 칸 왼쪽에 빈 쪽칸이 붙어 있으면(공양식 금속성이물 줄) 합친다.
        """
        prior = _prior_91(data, old_index)
        carried = 0
        done, k, follow = 0, 0, []
        for row in t91.rows[1:]:
            cells = E.raw_cells(row)
            texts = [E.cell_text(c) for c in cells]
            ci = next((i for i, t in enumerate(texts) if D.PREFIX.search(t or "")), None)
            if ci is None:
                # 허용기준 칸이 위 줄과 병합된 줄(금속성이물 ‘개개’) — 위 줄의 기준을 이어 쓴다.
                # 구분 글('개개')이 없는 빈 줄이라도 바로 위가 금속성이물(합계)이면 '개개' 줄로 본다
                # (담당자 PC 공양식 2026-09-06: 금속성이물 결과 칸이 두 줄로 갈라져 아래 줄이 사선으로 남았다)
                sub = "".join(D.squeeze(t) for t in texts[1:-1])
                if not follow or len(cells) < 2:
                    continue
                prev, crit_text = follow[-1]
                if not sub:
                    if "금속성이물" in prev["item"] and "개개" not in prev["sub"]:
                        sub = "개개"
                    else:
                        continue
                hit = dict(prev, sub=sub)
                ci = None
            elif ci >= len(cells) - 1:
                continue
            else:
                crit_text = texts[ci]
                hit = rules[k] if k < len(rules) else None       # criteria() 와 같은 차례로 훑는다
                k += 1
                if hit is None:
                    continue
                follow.append((hit, crit_text))
            lab = hit["item"] + hit["part"] + hit["sub"]
            value = maker(hit["process"])
            got = [value(lab, lot, i) for i, lot in enumerate(lots)]
            crit = D.PREFIX.sub("", crit_text or "")
            out = summarize([g for g in got if g], crit, hit["item"], hit["sub"],
                            unit_of(crit, hit["item"], hit["sub"]))
            if "포장규격" in hit["item"] and out:
                out = "각 규격에 적합함"
            if not out and "확인" in hit["item"]:
                out = _ident_result(lots, rec, prior, crit_text) or ""
                carried += bool(out)
            if not out:
                old = prior.get(_crit_key(crit_text))
                if old and not re.search(r"\d", old):
                    out = old
                    carried += 1
            if out:
                target = cells[-1]
                # 결과 칸 왼쪽에 글 없는 쪽칸이 붙어 있으면(기준 칸이 아니면) 합쳐서 사선이 생기지 않게 한다
                left_ok = (ci is None and len(cells) >= 2) or (ci is not None and len(cells) - 2 > ci)
                if left_ok and not E.cell_text(cells[-2]).strip() and not D.PREFIX.search(E.cell_text(cells[-2])):
                    E.clear_diag(cells[-2])
                    if "금속성이물" in hit["item"]:
                        # 결재본은 이 쪽칸에 '합계'·'개개' 를 적고 결과를 오른쪽 칸에 둔다
                        # (담당자 2026-09-07: "포장에서의 금속성 이물도 왼쪽을 참고해 줘")
                        E.set_cell(cells[-2], "개개" if "개개" in hit["sub"] else "합계")
                        E.set_cell_align(cells[-2], "center")      # 가운데 맞춤 (담당자 2026-09-07)
                        E.set_cell_valign(cells[-2], "center")
                    else:
                        E.merge_right(cells[-2], cells[-1])
                        target = cells[-2]
                E.clear_diag(target)
                E.set_cell(target, *out.split("\n"))
                done += 1
        if carried:
            log("9.1항: 글 결과 %d줄은 전년도 결재본에서 옮김(확인 시험 등)" % carried)
        return done

    if labelled:
        triples = D.tables_92_by_market(document)
        pairs_dom = [(t, st) for t, st, mk in triples if mk != "수출"]
        pairs_exp = [(t, st) for t, st, mk in triples if mk == "수출"]
        cpk_dom = fill_92_labelled(pairs_dom, rules, dom, n_dom, _prior_91(data, 0))
        if t91:
            log("9.1항: 결과 칸 %d줄을 판독값으로 채움" % fill_91_labelled(t91[0], rules, dom, n_dom, lambda pr: maker_of(pr, 0), 0))
        cpk_exp = {}
        if exp and pairs_exp:
            rules_e = D.criteria(t91[1]) if len(t91) > 1 else rules
            cpk_exp = fill_92_labelled(pairs_exp, rules_e, exp, n_exp, _prior_91(data, 1))
            if len(t91) > 1:
                log("9.1항(수출): 결과 칸 %d줄을 판독값으로 채움" % fill_91_labelled(t91[1], rules_e, exp, n_exp, lambda pr: maker_of(pr, 1), 1))
        log("9.2항: 머리글로 짚어 표 %d개 채움 (%s)" % (len(triples), ", ".join((mk + " " if mk else "") + st for _, st, mk in triples)))
        # 공양식의 각주 '1) 모든 시험 결과값이 0매로 동일하여 별도 계산하지 않음.' — 최댓값·최솟값·평균을 적으므로
        # 앞뒤가 맞지 않아 지운다(담당자 2026-09-06: "최댓값, 최솟값 기재가 안 됐네")
        gone = 0
        for para in document.paragraphs:
            if "동일하여" in para.text and ("별도 계산" in para.text or "별도로 작성" in para.text):
                E.set_para_text(para._p, "")
                gone += 1
        if gone:
            log("9.2항: '동일하여 별도 계산하지 않음' 각주 %d줄 지움 — 최댓값·최솟값·평균을 적는다" % gone)
            # 그 각주를 가리키던 머리칸의 윗첨자 '1)' (디겐타 공양식 '금속성이물1)') 도 지운다
            marks = 0
            for t, _ in D.tables_92(document):
                for row in t.rows[:3]:
                    for cell in E.raw_cells(row):
                        if "금속성이물" not in E.cell_text(cell):
                            continue
                        # '1)' 이 '1'·')' 두 윗첨자 조각으로 나뉘어 있기도 하다 — 이어진 윗첨자 조각을 합쳐 본다
                        for p_ in cell._tc.iter(qn("w:p")):
                            runs = list(p_.findall(qn("w:r")))
                            sup = [r_ for r_ in runs if r_.find(qn("w:rPr")) is not None
                                   and r_.find(qn("w:rPr")).find(qn("w:vertAlign")) is not None]
                            txt = "".join(x.text or "" for r_ in sup for x in r_.findall(qn("w:t")))
                            if sup and re.fullmatch(r"\s*\d\)\s*", txt):
                                for r_ in sup:
                                    p_.remove(r_)
                                marks += 1
            if marks:
                log("9.2항: 머리칸의 각주 번호 %d개 지움" % marks)
    else:
        cpk_dom = fill_92("9.2.1", dom, n_dom, True) or {}
        cpk_exp = fill_92("9.2.2", exp, n_exp, False) if exp else {}
    log("9항: Cpk %s" % {k: round(v, 2) for k, v in cpk_dom.items() if v is not None})

    # 각주 — 번호는 위에서 센 것을 쓴다
    notes = []
    if odd_bio:
        notes.append("%d) %s 조제(바이오버든) 공정 시험 성적서의 생균수 기재값은 “%s” 임. 원 기록의 단위 표기 확인 필요."
                     % (note_bio, ", ".join(l for l, _ in odd_bio), odd_bio[0][1]))
        issues.append(("9.2.2", ", ".join(l for l, _ in odd_bio), "생균수 기재값이 다른 Lot 과 다름 — 원본 확인"))
    notes.append("%d) %d년 완제 시험 성적서는 질량·용량 개개를 최솟값(···g 이상)으로만 기재하므로, 각 Lot 의 개개 최솟값으로 기재하였음."
                 % (note_mass, year_from))
    for t in reversed(t91):
        for nt in reversed(notes):
            E.note_after(document, t, nt)
    # 9.2 — 질량 각주는 그것을 설명하는 표(포장 질량·기밀도 표) 바로 아래 둔다. 양식에는 이미
    # '1) 모든 Lot 의 시험결과가 0 개(매)…' 각주가 있으므로 번호는 그다음부터 쓴다
    # (담당자 2026-09 점검: 11쪽에 '1)' 이 둘이고 질량 표보다 앞에 놓여 있었다).
    # 표는 detail92.tables_92 로 찾는다 — 양식의 각주 줄('1) 모든 Lot …')이 제목처럼 보여
    # 제목 기준 찾기(_tables)로는 그 아래 질량 표를 놓친다.
    from . import detail92 as _d92
    tabs92 = [t for t, _ in _d92.tables_92(document)]
    next_no = max(_note_numbers_under(document, "9.2") or [0]) + 1
    tb = _by_header(tabs92, "생균수")
    if tb is not None and odd_bio:
        E.note_after(document, tb, "%d) %s" % (next_no, notes[0].split(") ", 1)[1]))
        _mark_header(tb, "생균수", next_no)
        next_no += 1
    tm = _by_header(tabs92, "질량", "기밀도") or _by_header(tabs92, "질량", "개개")
    if tm is not None:
        E.note_after(document, tm, "%d) %s" % (next_no, notes[-1].split(") ", 1)[1]))
        _mark_header(tm, "개개", next_no)

    # ---------- 10항 ----------
    # 10.1 공정밸리데이션: 평가 년도에 보고서가 난 PV 를 채운다 (마스터파일)
    def fill_pv(table, hint=""):
        """표의 기존 보고서 번호(PV24-2-QUIO3-R …)에서 코드를 알아내 마스터에서 평가 년도 보고서를 찾는다.
        빈 공양식에는 번호가 없다 — 전년도 결재본 같은 자리 표의 번호(hint)로 찾는다(내수 QUIO3·수출 QUIO2)."""
        codes = set(re.findall(r"PV\d{2}-\d-([A-Z0-9]+)-", _text(table._tbl) + " " + hint))
        pv_path = next((p_ for p_ in data.files.get("10.1", []) if p_.lower().endswith(".xlsx")), None)
        if not codes or not pv_path:
            return 0
        from .readers.masters import pv_by_code
        rows = []
        for code in codes:
            for e in pv_by_code(pv_path, code):
                yr = re.match(r"PV(\d{2})", e["report"] or "")
                if yr and 2000 + int(yr.group(1)) == year_from:
                    rows.append(e)
        if not rows:
            return 0
        # 가장 최근 것 하나만 싣는다 (담당자 지시 2026-09) — 마스터파일에는 해묵은 PV 가 함께 있다.
        rows = [max(rows, key=lambda e: (e.get("report_date") or "", e.get("report") or ""))]
        # 열은 머리행 이름으로 짚는다 — 전년도 양식은 'No.' 열이 있고 2026 결재본은 없다.
        # 자리로 쓰면 사유 체크박스가 No. 칸에 들어가는 등 한 칸씩 밀린다.
        head = [re.sub(r"\s+", "", E.cell_text(h)) for h in E.raw_cells(table.rows[0])]

        def col(*words):
            return next((i for i, h in enumerate(head) if any(w in h for w in words)), None)

        i_no, i_why, i_lot = col("No.", "연번"), col("사유"), col("대상Lot", "Lot")
        i_doc, i_done, i_note = col("문서번호"), col("완료일"), col("비고")
        f, l = E.fit_rows(table, 1, len(table.rows) - 1, len(rows))
        for i, e in enumerate(rows):
            c = E.raw_cells(table.rows[f + i])

            def put(index, *lines):
                if index is not None and index < len(c):
                    E.set_cell(c[index], *lines)

            reason = (e.get("reason") or "").strip()
            kind = "변경" if "변경" in reason else ("정기적" if "정기" in reason else "최초")
            put(i_no, str(i + 1))
            put(i_why, *["%s %s" % ("■" if k == kind else "☐", k) for k in ("최초", "변경", "정기적")])
            put(i_lot, *[lot for _, lot, _ in e["lots"] if str(lot).strip()][:3])
            m = re.match(r"(\S+)\s*(\(.*?\))?", e["report"])
            put(i_doc, m.group(1), (m.group(2) or "").lower()) if m else put(i_doc, e["report"])
            put(i_done, e.get("report_date") or "확인 필요")
            put(i_note, reason)                 # 비고에 지난해 글이 남지 않게 늘 덮어쓴다
        return len(rows)
    t101 = _tables(document, "10.1")
    pv_n = 0
    olds = (getattr(data, "prev_sections_all", None) or {}).get("10.1") or []
    carried = 0
    for i, tb in enumerate(t101):
        hint = " ".join(" ".join(c for c in row if c) for row in olds[i]) if i < len(olds) else ""
        n = fill_pv(tb, hint)
        if not n and i < len(olds) and len(t101) > 1:
            n = _carry_rows(tb, olds[i], getattr(data, "previous_name", ""))   # 내수·수출 표마다 전년도 것을
            carried += n
        pv_n += n
    if carried:
        issues.append(("10.1", getattr(data, "previous_name", ""),
                       "평가 연도의 PV 를 마스터파일에서 찾지 못해 전년도 결재본의 10.1 을 옮겼습니다 — 최신 PV 마스터파일로 확인하세요"))
    if not pv_n:
        # PV 마스터에서 평가 연도 보고서를 찾지 못하면 전년도 결재본의 10.1 을 옮긴다 —
        # 빈칸으로 두지 않는다(담당자 2026-09: "10.1 … 전년도 PQR 내용도 참고하고").
        옛 = (getattr(data, "prev_sections", None) or {}).get("10.1") or []
        for tb in t101:
            pv_n += _carry_rows(tb, 옛, getattr(data, "previous_name", ""))
        if pv_n:
            issues.append(("10.1", getattr(data, "previous_name", ""),
                           "평가 연도의 PV 를 마스터파일에서 찾지 못해 전년도 결재본의 10.1 을 "
                           "옮겼습니다 — 최신 PV 마스터파일로 확인하세요"))
    log("10.1 PV 행: %d" % pv_n)

    본문 = _text(document.element.body)          # 보고서에 실제로 실린 설비만 문의 목록에 올린다
    # IQ·OQ·PQ 를 모두 넘긴다 — 빈 서식으로 만들 때 IQ·OQ 칸이 사선인 채로 남지 않게.
    eq_lookup = {k: {kind: [(d, dt) for d, dt in v["docs"] if d.startswith(kind)]
                     for kind in ("IQ", "OQ", "PQ")}
                 for k, v in data.equipment.items()}
    sp_lookup = {}
    빠진것 = {}                                  # 다른 라인·다른 방 공사라 빼 둔 적격성평가
    대신넣음 = {}                                # 뺄 것뿐이라 그대로 넣은 적격성평가
    for k, v in data.support.items():
        pq = [(d.split(" (")[0].strip(), dt.split("/")[0].strip()) for d, dt in v.get("PQ", []) if d and dt]
        if not pq:
            # 설비 행에 PQ 가 없으면 같은 시스템의 PQ 를 쓴다. 시스템 칸이 비었거나 'N/A' 인
            # 줄(클린에어 분배 시스템)은 설비 이름의 낱말로 같은 무리를 찾는다 — 그러지 않으면
            # 바로 위의 남남(질소 분배 시스템) 것을 가져온다.
            system = (v.get("system") or "").strip()
            words = [w for w in re.findall(r"[가-힣A-Za-z]{2,}", v.get("name") or "")
                     if w not in ("시스템", "분배", "라인")][:2]
            for k2, v2 in data.support.items():
                if k2 == k or not v2.get("PQ"):
                    continue
                same = bool(system) and system.upper() != "N/A" and v2.get("system") == system
                like = any(w in (v2.get("name") or "") or w in (v2.get("system") or "") for w in words)
                if same or like:
                    pq = [(d.split(" (")[0].strip(), dt.split("/")[0].strip()) for d, dt in v2["PQ"] if d and dt]
                    break
        # IQ·OQ 는 마스터파일의 '사유' 를 보고 이 라인에 해당하는 것만 싣는다 — 액제 라인
        # 리모델링(23)이나 다른 방 공사는 안연고 보고서에 넣지 않는다(담당자 2026-09).
        # IQ·OQ 를 한 칸에 합쳐 적은 줄(IOQ21-WS-…)은 표를 채울 때 그대로 둔다.
        def 우리것(kind):
            got = v.get(kind) or []
            why = (v.get("why") or {}).get(kind) or [""] * len(got)
            keep = [(d.split(" (")[0].strip(), dt.split("/")[0].strip())
                    for (d, dt), w in zip(got, why)
                    if d and dt and masters_mod.applies(w, LINE)]
            excluded = [(d.split(" (")[0].strip(), dt.split("/")[0].strip())
                        for (d, dt), w in zip(got, why)
                        if d and dt and not masters_mod.applies(w, LINE)]
            if not keep and excluded:
                # 이 설비의 IQ·OQ 가 '다른 라인·다른 방 공사' 문서뿐이면(주사용수 분배시스템 IQ15/OQ15,
                # 질소 분배라인 IOQ23 …) 빈칸에 사선을 긋지 않고 그 문서를 넣되 확인을 남긴다
                # (담당자 2026-09-06: "보고서도 IQ, OQ 가 작성되지 않고").
                대신넣음.setdefault(k, []).extend(d for d, _ in excluded)
                return excluded
            for d, _ in excluded:
                빠진것.setdefault(k, []).append(d)
            return keep

        sp_lookup[k] = {"PQ": pq, "IQ": 우리것("IQ"), "OQ": 우리것("OQ")}
    for k, docs in 빠진것.items():
        보이는 = [d for d in dict.fromkeys(docs)]
        if 보이는 and k in 본문:
            issues.append(("10.3~10.5", k, "다른 라인·다른 방 공사라 빼 둠: %s — 확인 필요"
                           % ", ".join(보이는)))
    for k, docs in 대신넣음.items():
        보이는 = [d for d in dict.fromkeys(docs)]
        if 보이는 and k in 본문:
            issues.append(("10.3~10.5", k, "이 설비의 IQ·OQ 는 다른 라인·다른 방 공사 문서뿐이라 그대로 넣었습니다: %s "
                                          "— 이 라인에 맞는 문서인지 확인" % ", ".join(보이는)))
    upd = 0
    # 빈 공양식은 10.3~10.5(때로 10.2)의 관리번호·설비명이 없다 — 전년도 결재본의 설비 줄을 먼저 세우고
    # 마스터파일로 문서·완료일을 갱신한다(담당자 2026-09-06: "작성할 줄 모르겠으면 16항의 전년도 PQR 결재본을
    # 참고해서 작성하고 … 업로드한 파일로 최신 내용으로 업데이트하면 돼").
    seeds = getattr(data, "prev_equipment", None) or {}
    for prefix in ("10.2", "10.3", "10.4", "10.5"):
        for t in _tables(document, prefix):
            if seeds.get(prefix) and not _has_equipment(t):
                n = _seed_equipment_rows(t, seeds[prefix])
                log("%s: 전년도 결재본에서 설비 %d대 줄을 세움 — 마스터파일로 갱신" % (prefix, n))
    for prefix in ("10.2",):
        for t in _tables(document, prefix):
            gone = drop_equipment(t, DRY_HEAT, keep=("디겐타",), name=name)
            if gone:
                log("10.2: 이 제품에 쓰지 않는 설비 %d대를 뺌" % gone)
            upd += update_qualification(t, eq_lookup)
    for t in _tables(document, "10.4"):        # 제조용수는 층마다 따로다 — 우리 층(1층) 설비로
        for old_mid, new_mid in use_our_floor(t, data.support):
            log("10.4: %s(%s층 설비)를 연고 라인 %s층 설비 %s 로 바꿈"
                % (old_mid, "·".join(sorted(_floors(data.support[old_mid].get("name")))), FLOOR, new_mid))
            issues.append(("10.4", old_mid,
                           "연고 라인(%s층) 설비가 아니라 같은 종류의 %s(%s)로 바꿨습니다 — 확인 필요"
                           % (FLOOR, new_mid, data.support[new_mid].get("name") or "")))
    for prefix in ("10.3", "10.4", "10.5"):
        for t in _tables(document, prefix):
            upd += update_qualification(t, sp_lookup)
    log("10항 IQ·OQ·PQ 갱신: %d" % upd)
    # 검토: 마스터파일에 문서가 있는데 빈 칸이 남았으면 문의 목록 맨 앞에 ★ 로 알린다 —
    # 담당자 PC 에서 10.4·10.5 IOQ 칸이 비어 나간 일(2026-09-06)이 되풀이되지 않게.
    빈칸 = []
    for prefix, lk in (("10.2", eq_lookup), ("10.3", sp_lookup), ("10.4", sp_lookup), ("10.5", sp_lookup)):
        for t in _tables(document, prefix):
            빈칸 += [(prefix, mid, kind) for mid, kind in blank_qualification_cells(t, lk)]
    if 빈칸:
        for prefix, mid, kind in 빈칸:
            issues.insert(0, (prefix, mid, "★ 검토: 마스터파일에 %s 문서가 있는데 칸이 비었습니다 — 프로그램 오류, 제작자에게 알려 주세요" % kind))
        log("★ 10항 검토: 빈 칸 %d — %s" % (len(빈칸), ", ".join("%s %s %s" % x for x in 빈칸)))
    else:
        log("10항 검토: 마스터파일 문서가 있는 IQ·OQ·PQ 칸은 모두 채워짐")
    # 마스터파일에 IQ·OQ 가 아예 없는 설비는 비워 둔다(값을 지어내지 않는다) — 대신 문의
    # 목록에 남겨 담당자가 마스터파일을 보완할지 판단하게 한다.
    for mid, kinds in sorted(eq_lookup.items()):
        빈 = [k for k in ("IQ", "OQ") if not kinds.get(k)]
        if 빈 and re.match(r"^[A-Z]{3}\d{4}", mid) and mid in 본문:
            issues.append(("10.2", mid, "마스터파일에 %s 가 없어 비워 둠 — 확인 필요"
                           % "·".join(빈)))
    if not pv_n:
        issues.append(("10.1", "", "평가 년도의 PV 보고서를 마스터파일에서 찾지 못해 결재본 값을 유지함 — 확인 필요"))

    # ---------- 11항 ----------
    devs = [d for d in data.deviations if (d.get("lot") in dom + exp) or (name and name[:4] in (d.get("title") or ""))]
    t11 = E.join_continuations(_tables(document, "11.1"))   # 쪽 나눔으로 갈라 둔 표는 하나로 잇는다
    if t11:
        tbl = t11[0]
        if devs:
            # 열은 머리행 이름으로 짚는다 — 전년도 양식은 '구분(제품)' 열이 있고 EDMS 서식은 없다.
            head = [re.sub(r"\s+", "", E.cell_text(h)) for h in E.raw_cells(tbl.rows[0])]
            def col(*words):
                return next((i for i, h in enumerate(head) if any(w in h for w in words)), None)
            i_kind, i_lot, i_doc = col("구분"), col("Lot"), col("문서")
            i_det, i_act, i_capa = col("일탈사항"), col("조치사항"), col("CAPA")
            f, l = E.fit_rows(tbl, 1, len(tbl.rows) - 2, len(devs))
            for i, d in enumerate(devs):
                c = E.raw_cells(tbl.rows[f + i])
                E.set_cell(c[0], str(i + 1))
                if i_kind is not None and i_kind < len(c):
                    E.set_cell(c[i_kind], "제품")
                if i_lot is not None and i_lot < len(c):
                    E.set_cell(c[i_lot], d.get("lot") or "")
                if i_doc is not None and i_doc < len(c):
                    E.set_cell(c[i_doc], d.get("doc_no") or "")
                detail = _deviation_lines(d)
                action = [(line, "none") for line in (d.get("correction") or "").split("\n") if line.strip()]
                if d.get("completed"):
                    action.append(("(조치사항 완료일 : %s)" % d["completed"], "none"))
                if i_det is not None and i_det < len(c):
                    E.set_cell_flow(c[i_det], detail)
                    # 제목·'* 일탈 내용'·'* 일탈 원인' 머리만 굵게 (담당자 2026-09-07)
                    E.bold_lines(c[i_det], ("[", "* "))
                if i_act is not None and i_act < len(c):
                    E.set_cell_flow(c[i_act], action)
                if i_capa is not None and i_capa < len(c):
                    E.set_cell(c[i_capa], "☐ Yes", "■ No")
            E.set_cell_plain(E.comment_cell(tbl), "특이사항 (Comment)",
                             "- 평가 년도 내 %d건의 일탈 있었으나, 모두 적합하게 조치되었으며 특이사항 없음."
                             % len(devs))
        else:
            E.set_cell_plain(E.comment_cell(tbl), "특이사항 (Comment)", "평가 년도 내 중요 일탈 및 기준 일탈 이력 없음.")
        for tb in t11[1:]:
            E.set_cell_plain(E.comment_cell(tb), "특이사항 (Comment)", "평가 년도 내 중요 일탈 및 기준 일탈 이력 없음.")
    for tb in E.join_continuations(_tables(document, "11.2")):
        E.set_cell_plain(E.comment_cell(tb), "특이사항 (Comment)", "평가 년도 내 경향 일탈 이력 없음.")

    # ---------- 12항 ----------
    t12 = E.join_continuations(_tables(document, "12."))
    # 주성분 이름 — 주성분을 건드리는 변경은 대상 제품에 이름이 없어도 이 제품에 해당한다
    주성분 = sorted({r["part"] for r in rules if r.get("part") and "함량" in r.get("item", "")})
    if t12:
        tbl = t12[0]
        # 담당자가 12항에 넣어 준 변경요청서는 모두 싣는다. 변경요청서의 '대상 제품' 에 이 제품
        # 이름이 없는 것도 있다(디겐타안연고 2026: 주성분 플루오로메톨론 멸균 온도 변경 건은
        # 대상이 후메론점안액으로만 적혀 있다). 제품 이름으로 걸러 내면 실제로 있었던 변경이
        # 보고서에서 통째로 빠진다 — 싣고, 이름이 없는 건은 확인해 달라고 남긴다.
        ccs = list(data.changes)
        간추림 = False
        if ccs:
            f, l = E.fit_rows(tbl, 1, len(tbl.rows) - 2, len(ccs))
            for i, cc in enumerate(ccs):
                c = E.raw_cells(tbl.rows[f + i])
                E.set_cell(c[0], str(i + 1))
                # 변경사항: 변경 번호와 변경 내용만 (변경 사유는 적지 않는다 — 담당자 지시 2026-09)
                lines = ["[%s] %s" % (cc.get("doc_no"), cc.get("title") or "")]
                lines += brief_change(cc.get("description") or "")
                if cc.get("unread"):               # 글자 없는 스캔本 — 문서번호만 적고 노랑으로 표시
                    lines = ["[%s] 확인 필요 — 변경요청서를 읽지 못했습니다" % cc.get("doc_no")]
                E.set_cell_plain(c[1], *lines)
                if cc.get("unread"):
                    E.highlight_cell(c[1])
                # 첫 줄(문서번호·변경명)만 굵게 — 변경 내용은 보통 글씨 (담당자 2026-09)
                E.bold_first_line_only(c[1])
                # 조치사항: 변경 실행 계획의 부서별 조치사항을 간추린다. 위탁사·위수탁 줄은 뺀다.
                # 변경통보는 위탁사·거래처에 알리는 일이라 자사 제품 PQR 에는 적지 않는다
                # (담당자 지시 2026-09: "변경통보는 삭제해줘 … 자사 제품내용만 PQR 에 작성").
                acts = [a for team, a in (cc.get("actions") or [])
                        if "위수탁" not in team and not re.search(r"위탁사|위수탁|변경\s*통보", a)]
                if acts:
                    E.set_cell_plain(c[2], *["%d. %s" % (k, a) for k, a in enumerate(acts, 1)])
                    간추림 = True
                else:
                    E.set_cell_plain(c[2], "확인 필요")
                    issues.append(("12", cc.get("doc_no") or "",
                                   "변경 실행 계획을 읽지 못했습니다 — 조치사항을 직접 적으세요"))
                E.set_cell(c[3], "확인 필요"); E.set_cell(c[4], "N/A")
                if cc.get("unread"):
                    pass                            # 못 읽은 건은 해당 여부도 알 수 없다 — 위에서 이미 알렸다
                elif not change_covers(cc, name, 주성분):
                    issues.append(("12", cc.get("doc_no") or "",
                                   "이 제품에 해당하는 변경인지 자동으로 확인하지 못했습니다 "
                                   "(관련 제품: %s) — 확인하세요" % (cc.get("products") or "적혀 있지 않음")))
            if 간추림:                       # 같은 안내를 건마다 되풀이하지 않는다
                issues.append(("12", "", "조치사항은 변경 실행 계획을 간추린 것입니다 — "
                                         "이 제품에 해당하지 않는 줄은 지우세요"))
            E.set_cell_plain(E.comment_cell(tbl), "특이사항 (Comment)", "N/A")
        else:
            E.set_cell_plain(E.comment_cell(tbl), "특이사항 (Comment)", "평가 년도 내 변경관리 이력 없음.")

    # ---------- 13항 ----------
    stab = getattr(data, "stability", None)
    logs = getattr(data, "stability_logs", None)
    if logs:
        spec = {r["part"]: r["text"] for r in rules if "함량" in r["item"] and r["part"]}
        if not spec:                                    # 주성분이 하나라 규격 줄에 성분 이름이 없는 제품(퀴노비드)
            spec = {part: "%.1f ~ %.1f%%" % (lo, hi)
                    for part, (lo, hi) in (limits.get("assay_parts") or {}).items()}
        # 실시 사유는 그 Lot 의 공정밸리데이션 사유 — 올해 10.1 을 먼저, 없으면 전년도 것을 쓴다
        why = dict(getattr(data, "pv_reasons", None) or {})
        why.update(CARRY.pv_reasons(document))
        _fill_stability26(document, logs, period, spec, log, issues, why, getattr(data, "prev_packs", None),
                          getattr(data, "prev_entries", None), getattr(data, "previous_name", ""))
    elif stab:
        _fill_stability(document, stab, log, limits["assay"])
    elif _carry_stability(document, getattr(data, "prev_stability", None) or {},
                          dict(getattr(data, "pv_reasons", None) or {},
                               **CARRY.pv_reasons(document)), log, issues,
                          getattr(data, "previous_name", "")):
        pass                    # 전년도 결재본에서 옮겨 왔다 — 빈칸으로 두지 않는다
    else:
        for prefix in ("13.1", "13.2"):
            for tb in _tables(document, prefix):
                if len(tb.rows) < 3:
                    continue
                f, l = E.fit_rows(tb, 1, len(tb.rows) - 2, 1)
                for c in E.raw_cells(tb.rows[f]):
                    E.set_cell(c, ""); E.set_vmerge(c, False)
                cells = E.raw_cells(tb.rows[f])
                E.set_cell(cells[0], "1"); E.set_cell(cells[min(3, len(cells) - 1)], "확인 필요")
                E.set_cell_plain(E.comment_cell(tb), "특이사항 (Comment)", "* 안정성 시험일지 판독 필요 — 담당자 확인 후 기재")
        for tb in _tables(document, "13.3"):
            firsts = [i for i, r in enumerate(tb.rows) if E.cell_text(E.raw_cells(r)[0]).strip().startswith("관리")]
            last = (firsts[0] - 1) if firsts else len(tb.rows) - 6
            f, l = E.fit_rows(tb, 1, last, 1)
            cells = E.raw_cells(tb.rows[f])
            for c in cells:
                E.set_cell(c, ""); E.set_vmerge(c, False)
            E.set_cell(cells[0], "확인 필요")
            E.set_cell_plain(E.comment_cell(tb), "특이사항 (Comment)", "* 안정성 시험일지 판독 필요 — 담당자 확인 후 기재")
        issues.append(("13", "", "안정성 시험(13.1~13.3) 값을 읽지 못해 '확인 필요' 로 두었음 — 시험일지 판독 필요"))

    # 13.2 — 양식은 '* 시판 후 안정성 시험 이력 없음.' 한 줄뿐이라 13.2 항이 비어 보인다
    # (담당자 2026-09: "13.2항이 없어 — 13.2 시판 후 안정성 이력 없음 이런식으로").
    p132 = find_para(document, "시판 후 안정성 시험 이력 없음")
    if p132 is not None and not p132.text.strip().startswith("13.2"):
        E.set_para_text(p132._p, "13.2 시판 후 안정성 시험 이력 없음.")
        log("13.2 줄: 항 번호 붙임")

    # ---------- 내역이 없는 표는 연번을 지우고 한 줄로 (담당자 2026-09-07:
    # "8.2.3, 11.1 및 11.2는 기재할 내용 없으면 연번 삭제하고 사선 처리해") ----------
    빈줄표 = []
    for prefix in ("8.2.3", "11.1", "11.2"):
        for tb in E.join_continuations(_tables(document, prefix)):
            if E.single_blank_row(tb):
                빈줄표.append(tb)
    if 빈줄표:
        log("내역 없는 표를 한 줄로(연번 지움): %d개" % len(빈줄표))

    # ---------- 14·15항 ----------
    us_export = "미국" in (name or "")           # 계획서 비고로 갈라진 '(미국 수출용)' 건
    빈표 = []
    returns = "평가 년도 내 반품 이력 없음." if us_export else "사용기한 경과 외 반품이력 없음"
    for prefix, msg in (("14.1", returns),
                        ("14.2", "평가 년도 내 불만 이력 없음."), ("14.3", "평가 년도 내 회수 이력 없음."),
                        ("15.", "평가 년도 내 시정조치사항 이력 없음.")):
        for tb in E.join_continuations(_tables(document, prefix)):
            if E.single_blank_row(tb):                   # 내역이 없으면 한 줄 (14.2 서식은 빈 줄 셋을 병합해 두었다)
                빈표.append(tb)
            last = E.raw_cells(tb.rows[-1])[0]
            if E.cell_text(last).lstrip().startswith("특이사항"):
                E.set_cell_plain(last, "특이사항 (Comment)", msg)

    # 내역이 없어 한 줄로 줄인 표는 칸마다 사선을 긋는 대신, 서식이 14.1·14.3·15항에 그어 둔
    # '한 줄을 가로지르는 선' 을 그대로 옮겨 온다 (담당자 2026-09-06: "31쪽 불만 사선도 수정이 안 됐어").
    if 빈표 or 빈줄표:
        본보기 = next((t for pre in ("14.1", "14.3", "15.", "14.2")
                     for t in _tables(document, pre) if E.has_drawing(t)), None)
        옮김 = sum(1 for t in 빈표 + 빈줄표 if E.copy_diag_line(본보기, t))
        log("빈 표 사선 한 줄: %d" % 옮김)

    # ---------- 16항 ---------- 배포본 'PQR 작성방법 공유의 건'(2026-09-04) 문안 그대로.
    # 10 Lot 미만이라 Cpk 를 산출하지 않았다는 말은 당연한 것이라 결론에 적지 않는다(담당자 지시).
    from . import conclusion
    plan_year, plan_q = conclusion.plan_quarter(today)      # 계획서 기한 = 작성일의 다음 분기
    # 조치가 끝나지 않은 일탈(완료일이 없는 것)이 있으면 16.2 로 다음 해 확인을 남긴다
    # (한림 결재본 PQR25 퀴노비드안연고 문안; 2026-09 점검).
    dev_open = [d.get("doc_no") or "" for d in devs if not d.get("completed")]
    dev_closed = [d.get("doc_no") or "" for d in devs if d.get("completed")]
    written = conclusion.apply(
        document, full_name, name or full_name, produced=bool(dom or exp), n_lots=len(dom),
        year=year_from, write_year=plan_year, quarter=plan_q, cpk=cpk_dom,
        open_deviation=conclusion.deviation_sentence(dev_closed, dev_open))
    if not written:
        issues.append(("16", "", "'16. 결론' 제목을 찾지 못해 결론을 다시 쓰지 못함 — 확인 필요"))
    log("16항 완료")
    cover = find_para(document, name[:5]) if name else None
    def rename_heading(old_text, new_text):
        """항 제목을 2026 결재본 차림새로 바꾼다. 찾지 못하면 그냥 둔다."""
        para = find_para(document, old_text)
        if para is not None and E.loose(para.text) != E.loose(new_text):
            E.set_para_text(para, new_text)
            return True
        return False

    # ---------- 14 · 17 · 18 항 차림새 (2026 결재본 기준) ----------
    rename_heading("반품 및 불만 회수관련 기록", "14. 반품, 불만 및 회수 현황표")
    # 17 참고 자료에서 첨부 문서(안정성 결과표·경향 분석 결과)를 18 항으로 옮긴다.
    ref = find_para(document, "17. 참고 자료")
    if ref is not None:
        _cpk_references(document, ref, qc.cpk_applies(len(dom)), log)
        moved, tail = [], []
        node = ref._p.getnext()
        while node is not None and node.tag == qn("w:p"):
            text = _text(node).strip()
            if re.match(r"^\s*18\.", text):
                moved = []                         # 이미 18 항이 있으면 손대지 않는다
                break
            if text.startswith("-") and ("HLF-QC-104" in text or "HLF-QC-126-06" in text):
                moved.append(node)
            elif text:
                tail.append(node)
            node = node.getnext()
        if moved:
            head = copy.deepcopy(ref._p)
            for run in head.findall(qn("w:r"))[1:]:
                head.remove(run)
            for t in head.iter(qn("w:t")):
                t.text = "18. 첨부 문서"
                t.set(qn("xml:space"), "preserve")
            anchor = (tail or moved)[-1] if not tail else tail[-1]
            last = moved[-1]
            for el in moved:                       # 옮길 줄들을 문서 끝으로 모은다
                el.getparent().remove(el)
            spacer = copy.deepcopy(moved[0])
            for run in spacer.findall(qn("w:r")):
                spacer.remove(run)
            anchor.addnext(head)
            head.addprevious(spacer)
            after = head
            for el in moved:
                after.addnext(el)
                after = el
            log("17·18항: 첨부 문서 %d줄을 18항으로 나눔" % len(moved))

    return {"issues": issues, "cover_title": (cover.text.strip() if cover is not None else None), "cpk": cpk_dom}


CPK_SHEETS = ("- 제품품질평가 경향분석 Sheet(한쪽 규격 용)(HLF-QC-126-08)",
              "- 제품품질평가 경향분석 Sheet(양쪽 규격 용)(HLF-QC-126-09)")


def _cpk_references(document, ref, applies, log):
    """17 참고 자료의 Cpk 경향분석 Sheet 줄 — 10 Lot 이상이면 있어야 하고, 미만이면 없어야 한다.

    한림 결재본(PQR25 퀴노비드안연고, 18 Lot)은 '별표 17' 줄 다음에 HLF-QC-126-08·-09 두 줄을
    적었다. 전년도가 10 Lot 미만이던 제품이 올해 넘기면 줄이 없고, 반대면 남는다(2026-09 점검).
    """
    lines, node = [], ref._p.getnext()
    while node is not None and node.tag == qn("w:p"):
        text = _text(node).strip()
        if re.match(r"^\s*18\.", text):
            break
        if text.startswith("-"):
            lines.append(node)
        node = node.getnext()
    have = [n for n in lines if "HLF-QC-126-08" in _text(n) or "HLF-QC-126-09" in _text(n)]
    if not applies:
        for n in have:
            n.getparent().remove(n)
        if have:
            log("17항: 10 Lot 미만이라 Cpk 경향분석 Sheet 참고 %d줄 뺌" % len(have))
        return
    if len(have) >= 2 or not lines:
        return
    # '별표 17' 줄 뒤(없으면 첨부 문서 줄 앞·마지막 참고 줄 뒤)에 넣는다
    anchor = next((n for n in lines if "별표" in _text(n)), None)
    if anchor is None:
        plain = [n for n in lines if "HLF-QC-104" not in _text(n) and "HLF-QC-126-06" not in _text(n)]
        anchor = (plain or lines)[-1]
    present = "".join(_text(n) for n in have)
    added = 0
    for text in CPK_SHEETS:
        code = text[text.rindex("(") + 1:-1]
        if code in present:
            continue
        new = copy.deepcopy(anchor)
        E.set_para_text(new, text)
        anchor.addnext(new)
        anchor = new
        added += 1
    if added:
        log("17항: Cpk 경향분석 Sheet 참고 %d줄 넣음" % added)


def _fill_numbers(t, f, l, lots, n, cpk, rec, put, is_dom, limits=None):
    limits = limits or DEFAULT_LIMITS
    for i, lot in enumerate(lots):
        r = rec(lot, "924"); c = E.raw_cells(t.rows[f + i])
        E.set_cell(c[0], str(i + 1)); E.set_cell(c[1], lot)
        E.set_cell(c[2], ("%.2f" % float(r["particle"])) if r.get("particle") else "")
        E.set_cell(c[3], ("%.2f" % float(r["mass_avg"])) if r.get("mass_avg") else ""); E.set_cell(c[4], "%s 이상" % r.get("mass_each_min") if r.get("mass_each_min") else "")
        E.set_cell(c[5], r.get("assay") or "")
    pt, pa, pi, ct = n["pt"], n["pa"], n["pi"], n["ct"]
    put(t, l + 1, (1, 2, 3, 4), "%.2f" % max(pt), "%.2f" % max(pa), "%.2f" % max(pi), "%.1f" % max(ct))
    put(t, l + 2, (1, 2, 3, 4), "%.2f" % min(pt), "%.2f" % min(pa), "%.2f" % min(pi), "%.1f" % min(ct))
    put(t, l + 3, (1, 2, 3, 4), "%.2f" % (sum(pt) / len(pt)), "%.2f" % (sum(pa) / len(pa)), "", "%.1f" % (sum(ct) / len(ct)))
    if qc.cpk_applies(len(lots)):
        cpk["particle"] = cpk_uni(pt, limits["particle"]); cpk["assay"] = cpk_bi(ct, *limits["assay"])
        if cpk["particle"] is not None and cpk["assay"] is not None:
            put(t, l + 4, (1, 4), "%.2f" % cpk["particle"], "%.2f" % cpk["assay"])
            put(t, l + 5, (1, 4), "충분" if cpk["particle"] >= 1 else "부족", "충분" if cpk["assay"] >= 1 else "부족")


# 변경사항 글을 간추린다 — 담당자 지시(2026-09): "변경은 간략히 작성해주면 돼"
ITEM = re.compile(r"^\s*(\d{1,2})[.]\s*(.+)$")            # 1. 큰 항목
SUB = re.compile(r"^\s*(?:\d{1,2}\)|[-–●○*]|[가-힣]\))\s*")  # 1) · - 같은 잔가지
TAIL = re.compile(r"^\s*[:：]\s*(.+)$")                     # ': 내용' — 바로 위 항목에 붙는 설명


def brief_change(text):
    """변경요청서의 변경 내용에서 큰 항목만 남기고 잔가지를 버린다.

    큰 항목('1. 안연고 튜브 자재 시험 검체 수량 확대')과 바로 뒤의 ': …' 설명을 한 줄로 잇고,
    그 아래 '1)'·'-' 로 시작하는 세부 개정 목록은 버린다. 큰 항목이 하나도 없으면 원문을
    그대로 둔다(항목 번호 없이 한두 줄로 적힌 변경건).
    """
    lines = [x.strip() for x in (text or "").split("\n") if x.strip()]
    items = [x for x in lines if ITEM.match(x)]
    if not items:
        return lines
    out, keeping = [], False
    for line in lines:
        if ITEM.match(line):
            out.append(line)
            keeping = True
            continue
        if SUB.match(line):
            keeping = False                      # 잔가지부터는 그 아래 이어진 줄도 버린다
            continue
        tail = TAIL.match(line)
        if tail and out:
            out[-1] = "%s: %s" % (out[-1].rstrip(" :："), tail.group(1))
            keeping = True
        elif keeping and out:
            out[-1] = "%s %s" % (out[-1], line)  # PDF 에서 줄이 접힌 것 — 앞줄에 잇는다
    return out


def _stability_tables(document):
    """13항 표를 제목으로 가른다 — [(kind, market, table)]. kind: '장기'·'시판후'·'경향', market: '내수'·'수출'·''.

    디겐타 서식은 13.1 장기 / 13.3 경향 표 하나씩(market ''). 퀴노비드처럼 내수용·수출용이 따로인 서식은
    13.1.1 내수용 / 13.1.2 수출용 (베트남) / 13.2.x 시판 후 / 13.3.x 경향 으로 갈린다(2026-09).
    """
    out, kind, market = [], None, ""
    for k, value, _ in outline(document):
        if k == "h":
            t = re.sub(r"\s+", "", value)
            m = re.match(r"^(\d+(?:\.\d+)*)", t)
            if not m:
                continue
            num = m.group(1).rstrip(".")
            if not (num == "13" or num.startswith("13.")):
                kind = None
                continue
            if num.count(".") <= 1:                       # 13 · 13.1 · 13.2 · 13.3
                if "시판" in t:
                    kind = "시판후"
                elif "경향" in t:
                    kind = "경향"
                elif "장기" in t:
                    kind = "장기"
                else:
                    kind = None
                market = ""
            else:                                         # 13.1.1 내수용 · 13.1.2 수출용 (베트남)
                market = "수출" if "수출" in t else ("내수" if "내수" in t else market)
        elif kind:
            out.append((kind, market, document.tables[value]))
    return out


_STAND_IN = re.compile(r"동일\s*수탁\s*제품\s*[(（]\s*([^)）]+?)\s*[)）]\s*(?:로|으로)?\s*갈음")


def _stand_in_names(tabs):
    """서식 13항 특이사항이 알려 주는 갈음 제품 — {제품 이름: 그 표의 시장}.

    '동일 수탁 제품(에펙신안연고)로 갈음하였음' 처럼 적혀 있으면, 그 이름이 든 시험일지는
    그 표(수출용)의 것이다. 제품마다 갈음 제품이 다르므로 서식에서 읽어 쓴다.
    """
    out = {}
    for kind, market, table in tabs:
        if not market or not table.rows:
            continue
        글 = E.cell_text(E.raw_cells(table.rows[-1])[0])
        for name in _STAND_IN.findall(글):
            out.setdefault(re.sub(r"\s+", "", name), market)
    return out


def _post_completed(one, taken):
    """시판 후 안정성 — 마지막 시점이 사용기한(제조일자~사용기한)에 닿았으면 '완료', 아니면 '진행중'."""
    months = [int(p["period"][:-1]) for p in one.get("points", []) if (p.get("period") or "").endswith("M")]
    if not months:
        return False
    mfg, exp = one.get("mfg") or "", one.get("expiry") or ""
    if len(mfg) >= 7 and len(exp) >= 7:
        total = (int(exp[:4]) - int(mfg[:4])) * 12 + (int(exp[5:7]) - int(mfg[5:7]))
        return max(months) >= total - 1
    return max(months) >= 36


def _fill_131_table(table, rows, why_of, issues, post=False):
    """13.1(장기)·13.2(시판 후) 실시 내역 — Lot 하나가 한 줄, 시험 기간·완료 일자는 줄바꿈으로 잇는다.
    마지막 열은 장기면 '실시 사유', 시판 후면 '비고'(완료/진행중)."""
    f, _ = E.fit_rows(table, 1, len(table.rows) - 2, len(rows))
    for i, (one, taken) in enumerate(rows):
        cells = E.raw_cells(table.rows[f + i])
        if post:
            last = "완료" if _post_completed(one, taken) else "진행중"
        else:
            last = one.get("why") or (why_of or {}).get(one["lot"]) or ""
        carried = bool(one.get("carried"))                  # 전년도 결재본에서 옮긴 Lot — 올해 시점·완료 일자는 모른다
        periods = ["확인 필요"] if carried else [p["period"] for p in taken]
        dones = ["확인 필요"] if carried else [p["done"] or "확인 필요" for p in taken]
        put = [str(i + 1), one.get("year") or "", periods, one.get("lot_text") or one["lot"],
               one.get("pack") or "", one.get("store") or "", dones, last]
        for k, value in enumerate(put):
            if k >= len(cells):
                break
            E.set_cell(cells[k], *(value if isinstance(value, list) else str(value).split("\n")))
            E.set_vmerge(cells[k], False)
            E.clear_diag(cells[k])
        # 손글씨 판독이 애매한 완료 일자는 노랑 (담당자 2026-09: "애매한 것만 노랑마크로")
        if any("done" in (p.get("unsure") or []) for p in taken) and len(cells) > 6:
            E.highlight_cell(cells[6])                      # 시점마다 한 줄('확인 필요' 도 줄마다) — 시험 기간 줄과 맞춘다
        if carried:
            for k in (2, 6, 7):
                if k < len(cells):
                    E.highlight_cell(cells[k])
    if not post and not any(one.get("why") or (why_of or {}).get(one["lot"]) for one, _ in rows):
        issues.append(("13.1", "", "장기 안정성 시험의 ‘실시 사유’ 는 시험일지에 없습니다 — "
                                   "변경관리·PV 내용을 보고 직접 적으세요"))
    moved = [one["lot"] for one, _ in rows if one.get("carried")]
    if moved:
        cell = E.comment_cell(table)
        kept = [l for l in E.cell_text(cell).split("\n")[1:] if l.strip() and l.strip() != "N/A"]   # 서식의 글은 남긴다
        kept.append("* %s 줄은 전년도 결재본·서식 각주에서 옮긴 것입니다 — 올해 안정성 시험일지로 시험 기간·완료 일자를 채우세요."
                    % ", ".join(moved))
        E.set_cell_plain(cell, "특이사항 (Comment)", *kept)


def _fill_133_table(table, groups, spec, _trim, marks=None):
    """13.3 경향 분석 — groups: [(줄 이름 '시판 후'|'장기', log, 평가 연도까지의 시점들)]. 성분마다 최솟값 ~ 최댓값."""
    labels = D.labels(table)
    parts = []
    for k, name in enumerate(labels):
        got = next((p for p in spec if D.squeeze(p) and D.squeeze(p) in name), None)
        if got:
            parts.append((k, got))
    if not parts and len(spec) == 1:                       # 성분 이름 없이 '함량(%)' 한 열뿐인 표(퀴노비드)
        value_cols = [k for k, name in enumerate(labels) if k >= 2 and name and "시험항목" not in name]
        if value_cols:
            parts.append((value_cols[0], list(spec)[0]))
    if not parts:
        return 0
    width = len(labels)
    heads = [i for i, tr in enumerate(table._tbl.findall(qn("w:tr")))
             if D.squeeze(_text(tr.findall(qn("w:tc"))[0])).startswith(("관리규격", "최소", "최대", "경향"))]
    first = next((i for i, r in enumerate(table.rows)
                  if D.squeeze(E.cell_text(E.raw_cells(r)[0])) in ("장기", "시판후")), 2)
    # 줄 차례는 서식을 따른다 — EDMS 공양식은 '장기' 를 먼저, 2025 결재본은 '시판 후' 를 먼저 적었다
    lead = D.squeeze(E.cell_text(E.raw_cells(table.rows[first])[0])) if first < len(table.rows) else ""
    if lead == "장기":
        groups = [g for g in groups if g[0] == "장기"] + [g for g in groups if g[0] != "장기"]
    last = heads[0] - 1 if heads else len(table.rows) - 1
    f, l = E.fit_rows(table, first, last, len(groups))
    heads = [i for i, tr in enumerate(table._tbl.findall(qn("w:tr")))          # 줄 수를 맞춘 뒤 다시 찾는다 — 번호가 밀린다
             if D.squeeze(_text(tr.findall(qn("w:tc"))[0])).startswith(("관리규격", "최소", "최대", "경향"))]
    values = {k: [] for k, _ in parts}
    guessed = {k: set() for k, _ in parts}                 # 애매하게 읽힌 예상값 — 최소·최대가 여기서 나오면 노랑
    notes, prev_label = [], None
    for i, (label, one, taken) in enumerate(groups):
        cells = _grid_cells_of(table.rows[f + i], width)
        head = label != prev_label
        if cells.get(0) is not None:
            E.set_cell(cells[0], label if head else "")
            E.set_vmerge(cells[0], "restart" if head else None)
        prev_label = label
        same = [g for g in groups if g[0] == label and g[1]["year"] == one["year"]]
        mark = ""
        if marks:                                          # 서식 각주('1) OEX101 …')의 번호를 그대로
            mark = marks.get(one["lot"], "")
        elif len(same) > 1:
            mark = "%d)" % (sum(1 for g in groups[:i] if g[0] == label and g[1]["year"] == one["year"]) + 1)
        if cells.get(1) is not None:
            E.set_cell(cells[1], "%s%s" % (one.get("year") or "", mark))
            E.set_vmerge(cells[1], False)
        notes.append((label, "%s%s" % (one.get("year") or "", mark), one["lot"]))
        for k, part in parts:
            if cells.get(k) is None:
                continue
            shaky = [p for p in taken if part in (p.get("unsure") or [])]
            got = [p["assays"].get(part) for p in taken]
            got = [float(x) for x in got if x is not None]
            E.clear_diag(cells[k])
            if not got and not shaky and one.get("carried"):
                # 전년도 결재본 13.3 의 그 Lot 줄 — 올해 시점 값이 더해져야 하니 노랑으로 남긴다
                prior = one.get("prev_range") or {}
                text = prior.get(part) or (next(iter(prior.values())) if len(prior) == 1 else "")
                nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", text or "")]
                E.set_cell(cells[k], text or "확인 필요")
                E.highlight_cell(cells[k])
                values[k] += nums
                guessed[k].update(nums)
                continue
            if not got and not shaky:
                E.set_cell(cells[k], "")
                continue
            if got:
                values[k] += got
                guessed[k].update(float(p["assays"][part]) for p in shaky if p["assays"].get(part) is not None)
                E.set_cell(cells[k], "%s ~ %s" % (_trim(min(got)), _trim(max(got))) if len(got) > 1 or min(got) != max(got)
                           else _trim(got[0]))
            else:
                E.set_cell(cells[k], "확인 필요")                 # 그해 값을 하나도 못 읽었다
            if shaky:
                E.highlight_cell(cells[k])
    bold_rows = set()
    for ri in heads:
        cells = _grid_cells_of(table.rows[ri], width)
        head = D.squeeze(E.cell_text(cells[0])) if cells.get(0) is not None else ""
        for k, part in parts:
            if cells.get(k) is None or not values[k]:
                continue
            if "관리규격" in head:
                E.set_cell(cells[k], re.sub(r"\s*%$", "", spec.get(part, "")))
                bold_rows.add(ri)
            elif "최소" in head:
                E.set_cell(cells[k], _trim(min(values[k])))
                if min(values[k]) in guessed[k]:
                    E.highlight_cell(cells[k])              # 예상값이 최소가 됐다 — 대조 필요
            elif "최대" in head:
                E.set_cell(cells[k], _trim(max(values[k])))
                if max(values[k]) in guessed[k]:
                    E.highlight_cell(cells[k])
            elif "경향" in head:
                lo_hi = re.findall(r"\d+(?:\.\d+)?", spec.get(part, ""))
                ok = len(lo_hi) < 2 or (float(lo_hi[0]) <= min(values[k]) and max(values[k]) <= float(lo_hi[1]))
                E.set_cell(cells[k], "적합" if ok else "부적합")
    for ri in bold_rows:                          # 관리 규격은 보통 글씨 (담당자 2026-09: "굵게 처리 하지 않음")
        E.unbold_row(table, ri)
    _note_133_lots(table, notes)
    return len(groups)


LOT_NOTE = "* 해당 연도의 제조번호 —"


def _note_133_lots(table, notes):
    """13.3 특이사항에 '어느 해가 어느 Lot 인지' 를 적는다.

    담당자 2026-09-07: "이 정보가 13.3.1 특이사항 칸에 기재되어야지" — 경향표에는 연도만 적혀
    2024¹⁾·2024²⁾ 가 어느 제조번호인지 표에서 알 수 없다. 서식에 담당자가 적어 둔 다른 글
    (갈음 문구 등)은 그대로 두고 이 줄만 새로 쓴다.
    """
    if not notes:
        return
    by = {}
    for label, year, lot in notes:
        by.setdefault(label or "", []).append("%s %s" % (year, lot))
    line = LOT_NOTE + " " + " / ".join("%s: %s" % (label, ", ".join(items)) for label, items in by.items())
    cell = E.comment_cell(table)
    kept = [l for l in E.cell_text(cell).split("\n")[1:]
            if l.strip() and l.strip() != "N/A" and not l.strip().startswith(LOT_NOTE)]
    E.set_cell_plain(cell, "특이사항 (Comment)", *(kept + [line]))


def _heading_material_name(document, table):
    """표 바로 앞 소제목 '8.2.1.1 포비돈 K-25' 의 원료명. 소제목이 원료명을 들고 있지 않으면 빈 값."""
    from .locate import headings_of_tables
    try:
        index = [t._tbl for t in document.tables].index(table._tbl)
    except ValueError:
        return ""
    for head in headings_of_tables(document).get(index, []):
        m = re.match(r"^\s*8\.2\.1\.\d+\.?\s*(.+)$", head)
        if m and m.group(1).strip() and not re.match(r"^(주원료|원료|시험결과)", m.group(1).strip()):
            return m.group(1).strip()
    return ""


def _declared_lots(table):
    """13.3 경향 표에 담당자가 미리 적어 둔 Lot — 특이사항의 '1) OEX101 2) OEX102 …' 각주와
    연도 칸의 '2025 4)' 표시를 맞춘다. {"lots": [(표시, Lot)], "year": {표시: 연도}, "label": {표시: '장기'|'시판후'}}.
    빈 공양식이라도 이 각주는 담당자가 적은 것이라 올해 시험한 Lot 으로 본다(퀴노비드 2026: OEY301·OEY302)."""
    out = {"lots": [], "year": {}, "label": {}}
    if not table.rows:
        return out
    comment = E.cell_text(E.raw_cells(table.rows[-1])[0])
    if not comment.lstrip().startswith("특이사항"):
        return out
    out["lots"] = [(m + ")", lot) for m, lot in re.findall(r"(\d)\)\s*([A-Z]{2}[A-Z0-9]{4})", comment)]
    label = ""
    for row in table.rows[1:-1]:
        cells = E.raw_cells(row)
        head = D.squeeze(E.cell_text(cells[0]))
        if head in ("장기", "시판후"):
            label = head
        if len(cells) < 2 or not label:
            continue
        m = re.match(r"^(\d{4})\s*(\d)\)$", D.squeeze(E.cell_text(cells[1])))
        if m:
            out["year"][m.group(2) + ")"] = m.group(1)
            out["label"][m.group(2) + ")"] = label
    return out


def _fill_stability26(document, logs, period, spec, log, issues, why_of=None, prev_packs=None,
                      prev_entries=None, previous_name=""):
    """2026 양식의 13항 — 장기(13.1)·시판 후(13.2) 실시 내역 · 경향 분석(13.3).

    logs: [{"lot", "year", "pack", "store", "kind", "market", "mfg", "expiry", "why",
            "points": [{"period", "done", "assays": {성분: 값}}, …]}, …]
    평가 기간 안에 끝난 시점만 실시 내역에 적고, 경향(13.3)도 그 시점까지의 값으로 낸다 —
    아직 하지 않은 뒤 시점을 넣으면 그 해의 경향이 아니다(한림 2026 결재본도 그렇게 쓴다).
    내수용·수출용 표가 따로면 시험일지의 시장(제품명 '(수출용)'·파일 이름)으로 갈라 넣는다.
    """
    year_to = lotcode.year_of((period or {}).get("to"))
    _trim = lambda v: ("%%.%df" % _decimals_of(logs)) % v

    def years_of(point):
        got = re.findall(r"\d{4}", point.get("done") or "")
        if got:
            return [int(got[0])]
        if point.get("expected"):
            # 완료 일자를 못 읽은 시점(오프라인 판독) — 제조일자+기간으로 어림한 시기. 연말(9~12월)
            # 예정이면 이듬해에 끝났을 수도 있어 두 해 모두 후보로 둔다. 13.1 에는 '확인 필요' 노랑으로 선다.
            y, mo = int(point["expected"][:4]), int(point["expected"][5:7])
            return [y] + ([y + 1] if mo >= 9 else [])
        return []

    def upto(point):        # 13.3 경향: 평가 연도까지 끝난 시점
        years = years_of(point)
        return (not year_to) or (bool(years) and min(years) <= year_to)

    def during(point):      # 13.1 실시 내역: 평가 연도에 끝난 시점
        years = years_of(point)
        return (not year_to) or (year_to in years)

    def split(some):
        rows, trend = [], []
        for one in some:
            for p in one.get("points", []):
                shaky = list(p.get("unsure") or [])
                # 문의 목록에는 보고서에 실제로 실리는 것만: 완료 일자는 평가 연도 시점(13.1)일 때,
                # 함량은 평가 연도까지의 시점(13.3·경향 엑셀)일 때
                if "done" in shaky and not during(p):
                    shaky.remove("done")
                if not upto(p):
                    shaky = []
                if shaky:
                    what = ", ".join("완료 일자" if u == "done" else "함량(%s)" % u for u in shaky)
                    issues.append(("13", one.get("lot", ""), "%s 시점 손글씨 판독이 애매함 — %s (노랑/주황 표시) 시험일지와 대조하세요"
                                   % (p.get("period"), what)))
            taken = [p for p in one.get("points", []) if during(p)]
            seen = [p for p in one.get("points", []) if upto(p)]
            if taken:
                rows.append((one, taken))
            if seen:
                trend.append((one, seen))
        return rows, trend

    tabs = _stability_tables(document)

    def pick(kind, market):
        same = [t for k, m, t in tabs if k == kind and m == market]
        if same:
            return same[0]
        blank = [t for k, m, t in tabs if k == kind and not m]
        return blank[0] if blank else None

    packs = prev_packs or {}
    # 갈음(대신하는) 제품 — 서식의 특이사항이 스스로 알려 준다:
    # "* 퀴노비드안연고(수출용)의 장기 안정성 시험은 동일 수탁 제품(에펙신안연고)로 갈음하였음."
    # 그 이름이 든 시험일지는 그 표(수출용)의 것이다 (담당자 2026-09-07: "수출용을 에펙신으로
    # 갈음한 거야 — 동일한 충전량이고 에펙신은 퀴노비드와 동일한 제품, 일동제약 수탁품이야").
    갈음이름 = _stand_in_names(tabs)
    갈음 = {}
    for one in logs:
        base = os.path.basename(one.get("source") or "")
        for name, market in 갈음이름.items():
            if name and name in base:
                갈음[base] = market
    if 갈음:
        log("13항: 갈음 제품 시험일지 %d장을 %s 로 봄 — %s"
            % (len(갈음), "·".join(sorted(set(갈음.values()))), ", ".join(sorted(set(갈음이름)))))
    for one in logs:
        # 시장 — 시험일지(파일 이름·'(수출용)')에 없으면 전년도 결재본에서 그 Lot(또는 앞 두 글자)이 어느 표에
        # 있었는지로 가른다. 판독기는 한글 제품명을 못 읽는다.
        if not one.get("market_hint"):
            lot = one.get("lot") or ""
            got = ((갈음.get(os.path.basename(one.get("source") or "")))
                   or (packs.get("market_by_lot") or {}).get(lot)
                   or (packs.get("market_by_prefix") or {}).get(lot[:2]))
            if got:
                one["market"] = got
    # 전년도 결재본에서 올해도 이어지는 시험(장기 36M 전 · 시판 후 '진행중')인데 올해 시험일지를 읽지
    # 못한 Lot 은 옮겨 놓고 '확인 필요' 노랑으로 남긴다 — 빈 표로 두지 않는다 (담당자 2026-09).
    carried = []
    read = {(one.get("kind") or "장기", one.get("market") or "내수", one["lot"]) for one in logs}
    for e in prev_entries or []:
        if not e.get("ongoing") or (e["kind"], e["market"], e["lot"]) in read:
            continue
        carried.append({"lot": e["lot"], "lot_text": e.get("lot_text") or e["lot"], "year": e.get("year") or "",
                        "pack": e.get("pack") or "", "store": e.get("store") or "", "why": e.get("last") or "",
                        "kind": e["kind"], "market": e["market"], "points": [], "carried": True,
                        "prev_range": dict(e.get("range") or {})})
    # 서식 13.3 각주에 적힌 Lot 가운데 시험일지에도 전년도 결재본에도 없는 것(올해 새로 시작한 장기 시험)
    # 도 줄을 세운다 — 값은 '확인 필요' 노랑. 시험일지가 올라오면 그 값으로 바뀐다.
    marks = {}
    for kind_, market_, table_ in tabs:
        if kind_ != "경향":
            continue
        got = _declared_lots(table_)
        marks[market_] = {lot: mark for mark, lot in got["lots"]}
        known = {one["lot"] for one in logs + carried}
        for mark, lot in got["lots"]:
            if lot in known:
                continue
            same = [e for e in prev_entries or [] if e["market"] == (market_ or "내수")]
            carried.append({"lot": lot, "lot_text": lot, "year": got["year"].get(mark, ""),
                            "pack": (packs.get("by_market") or {}).get(market_ or "내수", ""),
                            "store": next((e["store"] for e in same if e.get("store")), ""),
                            "why": "", "kind": got["label"].get(mark, "장기"), "market": market_ or "내수",
                            "points": [], "carried": True, "declared": True, "prev_range": {}})
            known.add(lot)
    markets = []
    for one in logs + carried:
        m = one.get("market") or "내수"
        if m not in markets:
            markets.append(m)
    markets.sort(key=lambda m: m != "내수")
    wrote = []
    for market in markets:
        mlogs = [one for one in logs if (one.get("market") or "내수") == market]
        mcarry = [one for one in carried if one["market"] == market]
        for one in mlogs:
            # 포장 형태 — 전년도 결재본의 표기('5g tube/갑'·'4.0g/Tube')를 따른다: 그 Lot 이 있으면 그 글, 없으면
            # 같은 시장의 글(용량 숫자가 시험일지와 같을 때). 둘 다 없으면 시험일지에서 읽은 것.
            mine = one.get("pack") or ""
            prior = (packs.get("by_lot") or {}).get(one["lot"]) or (packs.get("by_market") or {}).get(market) or ""
            g_mine = re.search(r"(\d+(?:\.\d+)?)\s*(?:g|mL|ml)", mine)
            g_prior = re.search(r"(\d+(?:\.\d+)?)\s*(?:g|mL|ml)", prior)
            same_size = bool(g_mine and g_prior and float(g_mine.group(1)) == float(g_prior.group(1)))
            one["pack"] = prior if (prior and (not mine or same_size)) else mine
        # 용량이 없는 판독값('PE')은 포장 형태가 아니다 — 같은 시장에서 가장 많이 쓰인 표기로 바꾼다
        # (담당자 2026-09-07: "2022 포장형태가 왜 PE야").
        제대로 = {}
        for one in mlogs:
            text = one.get("pack") or ""
            if re.search(r"\d+(?:\.\d+)?\s*(?:g|mL|ml|L)", text):
                제대로[text] = 제대로.get(text, 0) + 1
        흔한 = max(제대로.items(), key=lambda kv: kv[1])[0] if 제대로 else \
               ((packs.get("by_market") or {}).get(market) or "")
        for one in mlogs:
            text = one.get("pack") or ""
            if 흔한 and not re.search(r"\d+(?:\.\d+)?\s*(?:g|mL|ml|L)", text):
                if text:
                    issues.append(("13", one["lot"], "포장 형태를 '%s' 로 읽었는데 용량이 없어 "
                                                     "'%s' 로 적었습니다 — 시험일지와 대조하세요" % (text, 흔한)))
                one["pack"] = 흔한
        long_logs = [one for one in mlogs if (one.get("kind") or "장기") != "시판후"]
        post_logs = [one for one in mlogs if (one.get("kind") or "장기") == "시판후"]
        rows_l, trend_l = split(long_logs)
        rows_p, trend_p = split(post_logs)
        for one in mcarry:
            target = (rows_l, trend_l) if one["kind"] == "장기" else (rows_p, trend_p)
            target[0].append((one, [])); target[1].append((one, []))
        if mcarry:
            moved = [one["lot"] for one in mcarry if not one.get("declared")]
            new = [one["lot"] for one in mcarry if one.get("declared")]
            what = []
            if moved:
                what.append("전년도 결재본%s의 이어지는 Lot(%s)을 옮겼고"
                            % (("(%s)" % os.path.basename(previous_name)) if previous_name else "", ", ".join(moved)))
            if new:
                what.append("서식 13.3 각주의 Lot(%s)은 줄만 세웠습니다" % ", ".join(new))
            issues.insert(0, ("13", ", ".join(one["lot"] for one in mcarry),
                              "★ 13 폴더에 %s 안정성 시험일지가 없습니다 — 그 일지를 올리면 시험 기간(3M·6M…)과 "
                              "완료 일자를 읽어 채웁니다. 지금은 %s(노랑)"
                              % ("·".join(sorted({"%s %s" % (one["kind"], market) for one in mcarry})),
                                 " ".join(what))))
        t = pick("장기", market)
        if t is not None and rows_l:
            _fill_131_table(t, rows_l, why_of, issues, post=False)
            wrote.append("장기·%s %d Lot" % (market, len(rows_l)))
        t = pick("시판후", market)
        if t is not None and rows_p:
            _fill_131_table(t, rows_p, why_of, issues, post=True)
            wrote.append("시판후·%s %d Lot" % (market, len(rows_p)))
        elif rows_p:
            issues.append(("13.2", ", ".join(one["lot"] for one, _ in rows_p),
                           "시판 후 안정성 시험 표가 서식에 없어 넣지 못함 — 서식을 확인하세요"))
        t = pick("경향", market)
        groups = [("시판 후", one, seen) for one, seen in trend_p] + [("장기", one, seen) for one, seen in trend_l]
        if t is not None and groups:
            n = _fill_133_table(t, groups, spec, _trim, marks.get(market) or marks.get(""))
            wrote.append("경향·%s %d줄" % (market, n))
    log("13항: %s" % (", ".join(wrote) if wrote else "평가 기간에 든 시점이 없음"))


def _decimals_of(logs):
    """시험일지에 적힌 자릿수 그대로 쓴다 — 함량 99.0 을 99 로 줄이면 결재본과 달라진다."""
    seen = 1
    for one in logs:
        for point in one.get("points", []):
            for value in (point.get("assays") or {}).values():
                text = ("%s" % value)
                if "." in text:
                    seen = max(seen, len(text.split(".")[1]))
    return seen


def _grid_cells_of(row, width):
    out, col = {}, 0
    for cell in E.raw_cells(row):
        pr = cell._tc.find(qn("w:tcPr"))
        span_el = pr.find(qn("w:gridSpan")) if pr is not None else None
        span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
        if col < width:
            out[col] = cell
        col += span
    return out


def _fill_stability(document, stab, log, assay_limits=None):
    """stab: {"post_dom": [...], "post_exp": [...], "long_dom": [...], "long_exp": [...],
    "trend_dom": [...], "trend_exp": [...]} — 손글씨 판독(비전) 결과. 형식은 kynobuild/data.py 와 같다."""
    def fill_post(table, rows):
        f, l = E.fit_rows(table, 1, _last_data_row(table), max(1, len(rows)))
        for i, (no, yr, per, lot, pack, day, note) in enumerate(rows):
            c = E.raw_cells(table.rows[f + i])
            for k, v in enumerate((no, yr, per, lot, pack)):
                E.set_cell(c[k], *v.split("\n")); E.set_vmerge(c[k], False)
            E.set_cell(c[5], *STORE); E.set_vmerge(c[5], False)
            E.set_cell(c[6], day); E.set_vmerge(c[6], False)
            E.set_cell(c[7], note); E.set_vmerge(c[7], False)

    def fill_long(table, groups):
        total = sum(len(g[5]) for g in groups)
        f, l = E.fit_rows(table, 1, _last_data_row(table), max(1, total))
        ri, prev_year = f, None
        for no, yr, lot, pack, why, points in groups:
            for pi, (period, day) in enumerate(points):
                c = E.raw_cells(table.rows[ri]); head = (pi == 0)
                E.set_cell(c[0], no if head else ""); E.set_vmerge(c[0], "restart" if head else None)
                new_year = head and yr != prev_year
                E.set_cell(c[1], yr if new_year else ""); E.set_vmerge(c[1], "restart" if new_year else None)
                E.set_cell(c[2], period); E.set_vmerge(c[2], False)
                E.set_cell(c[3], *(lot.split("\n") if head else [""])); E.set_vmerge(c[3], "restart" if head else None)
                E.set_cell(c[4], pack if head else ""); E.set_vmerge(c[4], "restart" if head else None)
                E.set_cell(c[5], *(STORE if head else [""])); E.set_vmerge(c[5], "restart" if head else None)
                E.set_cell(c[6], day); E.set_vmerge(c[6], False)
                E.set_cell_plain(c[7], *(why.split("\n") if head else [""])); E.set_vmerge(c[7], "restart" if head else None)
                ri += 1
                if head:
                    prev_year = yr

    # 허용기준을 못 읽었으면 (None, None) 이 온다 — 13.3 관리 규격에는 흔한 값을 적어 둔다
    lo_hi = assay_limits if (assay_limits and assay_limits[0] is not None) else DEFAULT_LIMITS["assay"]

    def fill_trend(table, rows, note, comment):
        f, l = E.fit_rows(table, 1, len(table.rows) - 6, max(1, len(rows)))
        prev = None
        for i, (grp, yr, val) in enumerate(rows):
            c = E.raw_cells(table.rows[f + i]); head = grp != prev
            E.set_cell(c[0], grp if head else ""); E.set_vmerge(c[0], "restart" if head else None)
            E.set_cell(c[1], yr); E.set_cell(c[2], val); prev = grp
        lows = [float(r[2].split("~")[0]) for r in rows if "~" in r[2]]
        highs = [float(r[2].split("~")[1]) for r in rows if "~" in r[2]]
        for ri, v in ((l + 1, "%.1f~%.1f" % tuple(lo_hi)), (l + 2, "%.1f" % min(lows) if lows else ""), (l + 3, "%.1f" % max(highs) if highs else ""), (l + 4, "적합")):
            E.set_cell(E.raw_cells(table.rows[ri])[1], v)
        E.set_cell_plain(E.comment_cell(table), "특이사항 (Comment)", note, comment)

    t131 = _tables(document, "13.1"); t132 = _tables(document, "13.2"); t133 = _tables(document, "13.3")
    if t131:
        fill_post(t131[0], stab.get("post_dom", [])); E.set_cell_plain(E.comment_cell(t131[0]), "특이사항 (Comment)", stab.get("post_dom_note") or "N/A")
        if len(t131) > 1:
            fill_post(t131[1], stab.get("post_exp", [])); E.set_cell_plain(E.comment_cell(t131[1]), "특이사항 (Comment)", stab.get("post_exp_note") or "N/A")
    if t132:
        fill_long(t132[0], stab.get("long_dom", [])); E.set_cell_plain(E.comment_cell(t132[0]), "특이사항 (Comment)", stab.get("long_dom_note") or "N/A")
        if len(t132) > 1:
            fill_long(t132[1], stab.get("long_exp", [])); E.set_cell_plain(E.comment_cell(t132[1]), "특이사항 (Comment)", stab.get("long_exp_note") or "N/A")
    if t133:
        fill_trend(t133[0], stab.get("trend_dom", []), stab.get("trend_dom_lots", ""), stab.get("trend_dom_comment", ""))
        if len(t133) > 1:
            fill_trend(t133[1], stab.get("trend_exp", []), stab.get("trend_exp_lots", ""), stab.get("trend_exp_comment", ""))
    log("13항: 안정성 %d/%d/%d 표" % (len(t131), len(t132), len(t133)))



def _carry_rows(table, grid, source, first_col_number=True):
    """전년도 표의 자료 줄을 열 이름이 같은 칸만 골라 옮긴다.

    서식이 개정되며 열이 하나 늘거나 이름이 바뀌므로, 자리로 옮기면 값이 밀린다.
    옮긴 줄 수를 돌려준다.
    """
    if len(grid) < 2 or not table.rows:
        return 0
    width = E.grid_width(table)
    새머리 = {CARRY.squeeze(E.cell_text(c)): i
              for i, c in E.grid_cells(table.rows[0], width).items() if E.cell_text(c).strip()}
    옛머리 = {CARRY.squeeze(t): j for j, t in enumerate(grid[0]) if t.strip()}
    같은열 = {i: 옛머리[name] for name, i in 새머리.items() if name in 옛머리}
    if not 같은열:
        return 0
    줄 = [row for row in grid[1:]
          if not CARRY.squeeze(row[0] if row else "").startswith("특이사항")
          and any((row[j] or "").strip() for j in 같은열.values())]
    if not 줄:
        return 0
    끝 = len(table.rows) - 1
    if CARRY.squeeze(E.cell_text(E.raw_cells(table.rows[-1])[0])).startswith("특이사항"):
        끝 -= 1
    f, l = E.fit_rows(table, 1, 끝, len(줄))
    for k, row in enumerate(줄):
        cells = E.grid_cells(table.rows[f + k], width)
        for i, cell in cells.items():
            j = 같은열.get(i)
            글 = (row[j] if j is not None else "") or ""
            if j is None and i == 0 and first_col_number:
                글 = str(k + 1)               # 연번은 새로 매긴다
            elif j is None:
                continue
            E.set_vmerge(cell, False)
            if "\n" in 글:
                E.set_cell_plain(cell, *글.split("\n"))
            else:
                E.set_cell(cell, 글.strip())
            E.clear_diag(cell)
    return len(줄)


def _carry_stability(document, prev, why_of, log, issues, source=""):
    """전년도 결재본의 13.1 · 13.3 을 그대로 옮겨 놓는다 (담당자 2026-09).

    올해 안정성 시험일지를 읽지 못했을 때 13항을 빈칸으로 두지 않기 위한 것이다.
    옮긴 값은 작년 것이므로 두 표의 '특이사항' 에 갱신하라는 줄을 덧붙이고 문의 목록에도 남긴다.
    값을 새로 지어내지는 않는다 — 전년도 결재본에 적힌 글자를 그대로 옮길 뿐이다.
    """
    쓴표 = 0
    쓴표 += _carry_131(document, prev.get("13.1") or [], why_of or {}, source)
    쓴표 += _carry_133(document, prev.get("13.3") or [], source)
    if 쓴표:
        log("13항: 전년도 결재본에서 옮겨 옴 (표 %d개)" % 쓴표)
        issues.append(("13", os.path.basename(source or ""),
                       "올해 안정성 시험일지를 읽지 못해 전년도 결재본의 13항을 그대로 옮겼습니다 "
                       "— 13항 최신 안정성 시험 자료를 올려 다시 만드세요"))
    return 쓴표


def _carry_131(document, grid, why_of, source):
    """13.1 장기 안정성 실시 내역 — 열 이름이 같은 칸만 옮긴다.

    서식이 바뀌며 마지막 열이 '비고' 에서 '실시 사유' 가 되었다. 이름이 다르면 옮기지 않고,
    실시 사유는 10.1 밸리데이션 사유에서 채운다.
    """
    if len(grid) < 2:
        return 0
    tables = _tables(document, "13.1")
    if not tables:
        return 0
    table = tables[0]
    width = E.grid_width(table)
    새머리 = {CARRY.squeeze(t): i for i, t in
              enumerate(E.cell_text(c) for c in E.grid_cells(table.rows[0], width).values())}
    옛머리 = {CARRY.squeeze(t): j for j, t in enumerate(grid[0])}
    같은열 = {i: 옛머리[name] for name, i in 새머리.items() if name in 옛머리}
    lot_col = 옛머리.get("제조번호")
    줄 = [row for row in grid[1:]
          if lot_col is not None and CARRY.LOT.match((row[lot_col] or "").strip())]
    if not 줄:
        return 0
    f, l = E.fit_rows(table, 1, len(table.rows) - 2, len(줄))
    for k, row in enumerate(줄):
        cells = E.grid_cells(table.rows[f + k], width)
        for i, cell in cells.items():
            if i == 0:
                E.set_cell(cell, str(k + 1)); E.set_vmerge(cell, False)
                E.clear_diag(cell)
                continue
            j = 같은열.get(i)
            글 = (row[j] if j is not None else "") or ""
            if not 글 and i == max(새머리.values() or [0]):
                글 = why_of.get((row[lot_col] or "").strip(), "")
            E.set_vmerge(cell, False)
            if "\n" in 글:
                E.set_cell_plain(cell, *글.split("\n"))
            else:
                E.set_cell(cell, 글)
            E.clear_diag(cell)
    _note_carried(table, source, _old_comment(grid))
    return 1


def _carry_133(document, grid, source):
    """13.3 경향 분석 결과 — 성분 이름으로 열을 맞춰 옮긴다."""
    if len(grid) < 3:
        return 0
    tables = _tables(document, "13.3")
    if not tables:
        return 0
    table = tables[0]
    width = E.grid_width(table)

    # 열은 그리드 번호로 짚는다 — '시험항목' 칸이 두 열을 덮어 자리로 세면 한 칸씩 밀린다.
    새성분 = {CARRY.squeeze(E.cell_text(c)): i
              for i, c in E.grid_cells(table.rows[1], width).items()
              if CARRY.squeeze(E.cell_text(c))}
    옛성분 = {CARRY.squeeze(t): j for j, t in enumerate(grid[1]) if CARRY.squeeze(t)}
    if not 새성분 or not 옛성분:
        return 0
    라벨 = {}                                    # 옛 표의 '관리규격·최소·최대·경향분석결과' 줄
    장기 = []                                    # 연도별 값 줄
    for row in grid[2:]:
        머리 = CARRY.squeeze(row[0]) if row else ""
        if 머리.startswith("특이사항"):
            continue
        if 머리 in ("관리규격", "최소", "최대", "경향분석결과"):
            라벨[머리] = row
        elif any((row[k] or "").strip() for k in 옛성분.values()):
            장기.append(row)
    if not 장기 and not 라벨:
        return 0
    옮김 = 0
    첫줄 = next((i for i, r in enumerate(table.rows)
                 if CARRY.squeeze(E.cell_text(E.raw_cells(r)[0])) == "장기"), None)
    끝줄 = next((i for i, r in enumerate(table.rows)
                 if CARRY.squeeze(E.cell_text(E.raw_cells(r)[0])) == "관리규격"), None)
    if 첫줄 is not None and 끝줄 is not None and 장기:
        f, l = E.fit_rows(table, 첫줄, 끝줄 - 1, len(장기))
        for k, row in enumerate(장기):
            cells = E.grid_cells(table.rows[f + k], width)
            for name, i in 새성분.items():
                j = 옛성분.get(name)
                if j is not None and i in cells:
                    E.set_vmerge(cells[i], False)
                    E.set_cell(cells[i], (row[j] or "").strip())
                    E.clear_diag(cells[i])
            # 연도 칸(성분 열이 아닌 칸)도 그리드 번호 그대로 옮긴다 — 표 너비가 같을 때만.
            if len(grid[1]) == width:
                for i, cell in cells.items():
                    if i == 0 or i in 새성분.values():
                        continue
                    E.set_vmerge(cell, False)
                    E.set_cell(cell, (row[i] or "").strip())
                    E.clear_diag(cell)
        옮김 += len(장기)
    for ri, row in enumerate(table.rows):
        머리 = CARRY.squeeze(E.cell_text(E.raw_cells(row)[0]))
        옛 = 라벨.get(머리)
        if 옛 is None:
            continue
        if 머리 == "관리규격":
            E.unbold_row(table, ri)               # 담당자 2026-09: 관리 규격은 굵게 하지 않는다
        cells = E.grid_cells(row, width)
        for name, i in 새성분.items():
            j = 옛성분.get(name)
            if j is not None and i in cells:
                E.set_vmerge(cells[i], False)
                E.set_cell(cells[i], (옛[j] or "").strip())
                E.clear_diag(cells[i])
        옮김 += 1
    if not 옮김:
        return 0
    _note_carried(table, source, _old_comment(grid))
    return 1


def _note_carried(table, source, 옛글=None):
    """표 맨 아래 '특이사항' 칸을 전년도 문안으로 바꾸고, 어디서 옮겼는지 한 줄 덧붙인다.

    옮겨 온 값이 작년 것이므로 문안도 작년 것이어야 앞뒤가 맞는다 — 빈 서식의 예시
    문구를 그대로 두면 표에 없는 해를 가리키게 된다.
    """
    last = E.raw_cells(table.rows[-1])[0]
    글 = E.cell_text(last)
    if not 글.lstrip().startswith("특이사항"):
        return
    바탕 = 옛글 if 옛글 is not None else 글.split("\n")[1:]
    줄 = [l for l in 바탕 if l.strip() and l.strip() != "N/A"]
    줄.append("* 전년도 결재본%s의 내용을 옮긴 것입니다 — 올해 안정성 시험 자료로 갱신하세요."
              % (("(%s)" % os.path.basename(source)) if source else ""))
    E.set_cell_plain(last, "특이사항 (Comment)", *줄)


def _old_comment(grid):
    """전년도 표 맨 아래 '특이사항' 칸의 글줄 (머리말 '특이사항 (Comment)' 은 뺀다)."""
    if not grid:
        return None
    글 = (grid[-1] or [""])[0] or ""
    return 글.split("\n")[1:] if 글.lstrip().startswith("특이사항") else None
