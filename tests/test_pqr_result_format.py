"""9.1·9.2 결과 차림새와 11항 굵은 글씨, 10.4 층 맞추기 (담당자 지시 2026-09-07)."""
import os
import sys
import unittest

import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import detail92 as D, docedit as E                      # noqa: E402
from pqr.engine.recipe_ointment import use_our_floor                    # noqa: E402


def _table(rows):
    d = docx.Document()
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            E.set_cell(t.rows[i].cells[j], v)
    return t


class 소수_자릿수(unittest.TestCase):
    def test_한_열을_가장_자세한_값에_맞춘다(self):
        self.assertEqual(D.unify_decimals(["5.11", "5.1", "5", "5.07"]),
                         ["5.11", "5.10", "5.00", "5.07"])

    def test_이상_꼬리는_그대로_두고_숫자만_맞춘다(self):
        self.assertEqual(D.unify_decimals(["5.08 이상", "5 이상"]), ["5.08 이상", "5.00 이상"])

    def test_글로_적는_결과는_건드리지_않는다(self):
        self.assertEqual(D.unify_decimals(["음성", "음성"]), ["음성", "음성"])
        self.assertEqual(D.unify_decimals(["10 CFU/g 미만"]), ["10 CFU/g 미만"])


class 요약_줄(unittest.TestCase):
    def test_생균수처럼_글로_적는_결과는_최댓값도_평균도_내지_않는다(self):
        self.assertEqual(D._stats("생균수시험(CFU/g)", ["10 CFU/g 미만"] * 2), (None, None, None))

    def test_확인시험_문장도_숫자를_뽑아_적지_않는다(self):
        texts = ["224-228nm 및 292~296nm에서 흡수극대를 나타냄"] * 2
        self.assertEqual(D._stats("확인3)", texts), (None, None, None))

    def test_숫자_열은_자릿수를_지켜_적는다(self):
        self.assertEqual(D._stats("질량용량(g)평균", ["5.11", "5.10", "5.00"]), ("5.11", "5.00", "5.07"))


class 굵은_글씨(unittest.TestCase):
    def test_일탈_제목과_머리만_굵게(self):
        t = _table([["x"]])
        cell = t.rows[0].cells[0]
        E.set_cell_plain(cell, "[충전 수율 일탈 건]", "* 일탈 내용", "충전 수율이 관리 기준을 벗어남.")
        E.bold_lines(cell, ("[", "* "))
        got = []
        for para in E.cell_paras(cell._tc):
            text = "".join(t_.text or "" for t_ in para.iter(qn("w:t")))
            bold = [r.find(qn("w:rPr")).find(qn("w:b")).get(qn("w:val"))
                    for r in para.findall(qn("w:r"))]
            got.append((text, bold))
        self.assertEqual([b for _, b in got], [["1"], ["1"], ["0"]])
        self.assertEqual(got[0][0], "[충전 수율 일탈 건]")


class 제조용수_층(unittest.TestCase):
    support = {
        "HJA5037": {"name": "주사용수 제조 시스템 (2t - 2층)", "system": "주사용수 시스템"},
        "HJA5002": {"name": "주사용수 제조 시스템 (1.5t - 1층, 냉주사용수)", "system": "주사용수 시스템"},
        "WFIG2301": {"name": "주사용수 제조 시스템 (1층 안구이식제)", "system": "안구이식제 주사용수 시스템"},
        "HBA5089": {"name": "A.H.U 26호", "system": "공기조화장치"},
    }

    def _table(self):
        return _table([["No.", "관리번호", "설비명"],
                       ["1", "HJA5037", "주사용수 제조시스템"],
                       ["", "", ""]])

    def test_2층_설비를_같은_시스템의_1층_설비로_바꾼다(self):
        t = self._table()
        self.assertEqual(use_our_floor(t, self.support), [("HJA5037", "HJA5002")])
        self.assertEqual(E.cell_text(E.raw_cells(t.rows[1])[1]).strip(), "HJA5002")

    def test_층이_안_적힌_설비는_그대로_둔다(self):
        t = _table([["No.", "관리번호", "설비명"], ["1", "HBA5089", "A.H.U 26호"]])
        self.assertEqual(use_our_floor(t, self.support), [])


class 결론_표_폭(unittest.TestCase):
    def test_16_1_아래_표는_본문_들여쓰기만큼_들어가고_그만큼_좁아진다(self):
        d = docx.Document()
        p = d.add_paragraph("16.1 본 제품 품질 평가를 통해")
        pr = p._p.get_or_add_pPr()
        ind = pr.makeelement(qn("w:ind"), {}); pr.append(ind)
        ind.set(qn("w:left"), "420")
        t = d.add_table(rows=1, cols=3)
        for j, w in enumerate((3300, 3300, 3300)):
            t._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))[j].set(qn("w:w"), str(w))
        self.assertEqual(E.para_left_before(t), 420)
        self.assertTrue(E.inset_table(t, 420, 9978 - 420))
        pr = t._tbl.find(qn("w:tblPr"))
        self.assertEqual(pr.find(qn("w:tblInd")).get(qn("w:w")), "420")
        self.assertEqual(pr.find(qn("w:tblW")).get(qn("w:w")), "9558")
        self.assertEqual(pr.find(qn("w:tblW")).get(qn("w:type")), "dxa")
        cols = [int(g.get(qn("w:w"))) for g in t._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
        self.assertEqual(sum(cols), 9558)


if __name__ == "__main__":
    unittest.main()
