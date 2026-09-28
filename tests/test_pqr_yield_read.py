# -*- coding: utf-8 -*-
"""7항 수율 판독 — 서식 프로필, 판독값 검산, 수율현황표 초안, 화면에서 부르는 길.

담당자 2026-09-27: "우리 팀은 직접 작성하는 것이 편하다는 의견이고 다른 팀은 기록서 스캔본에서
수율이 기재된 페이지를 직접 AI 가 읽어서 정리해줬으면 좋겠다는 의견이야." → 두 길을 다 둔다.
쪽수는 제품마다가 아니라 "수탁사 × 기록서 서식마다 한 번 … 버전마다 관리".
"""
import json
import os
import shutil
import tempfile
import threading
import unittest
import urllib.error
import urllib.request

os.environ["PQR_REVIEW"] = "0"

from pqr import server as server_module                     # noqa: E402
from pqr.engine import claude_cli, yield_profile as P, yield_read as Y   # noqa: E402
from pqr.engine import collect as C                          # noqa: E402
from pqr.engine.readers import yield_sheet                  # noqa: E402
from pqr.sample import write_samples                        # noqa: E402

TODAY = "2026-08-27"

READING = {
    "lots": [
        {"lot": "GVY601", "values": {"조제": 99.85, "충전": 95.31, "포장": 99.99},
         "production": {}, "page": {"조제": 7, "충전": 12, "포장": 18}, "unsure": []},
        {"lot": "GVYD01", "values": {"조제": 99.85, "충전": 98.13},
         "production": {}, "page": {"조제": 7, "충전": 12}, "unsure": ["포장"]},
    ],
    "specs": {"조제": "99.5±0.45%", "충전": "98.0±3.0", "포장": "98.0±2.0%"},
}


class 쪽수_적기(unittest.TestCase):
    def test_공정_이름과_쪽을_함께_적는다(self):
        self.assertEqual(P.parse_pages("조제 7, 충전 12, 포장 18-19"),
                         {"조제": [7], "충전": [12], "포장": [18, 19]})

    def test_콜론도_슬래시도_받는다(self):
        self.assertEqual(P.parse_pages("조제: 7 / 충전: 12"), {"조제": [7], "충전": [12]})

    def test_이름_없이_쪽만_적어도_된다(self):
        self.assertEqual(P.parse_pages("7,12,18-19"), {"": [7, 12, 18, 19]})

    def test_쪽이_겹쳐도_한_번만(self):
        self.assertEqual(P.page_list({"pages": {"조제": [7], "충전": [12, 7]}}), [7, 12])


class 서식_프로필(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)

    def test_판이_다르면_다른_프로필이다(self):
        a = P.make_id("코오롱제약", "정산기록서", "MBRM-00004", "4.0")
        b = P.make_id("코오롱제약", "정산기록서", "MBRM-00004", "5.0")
        self.assertNotEqual(a, b)
        self.assertIn("4.0", a)

    def test_저장하고_다시_읽는다(self):
        data = P.empty()
        one = P.put(data, {"maker": "코오롱제약", "record": "정산기록서",
                           "form_no": "MBRM-00004", "rev": "4.0",
                           "pages": P.parse_pages("조제 1, 충전 2")})
        P.assign(data, "QC1-2080", one["id"])
        P.save(self.dir, data)
        again = P.load(self.dir)
        self.assertEqual(len(again["profiles"]), 1)
        self.assertEqual(P.for_product(again, "QC1-2080")["id"], one["id"])
        self.assertEqual(P.page_list(P.for_product(again, "qc1-2080")), [1, 2])

    def test_공통_폴더에_둔다(self):
        P.save(self.dir, P.empty())
        self.assertTrue(os.path.isfile(os.path.join(self.dir, "공통", P.FILE_NAME)))

    def test_서식_판이_바뀌면_알아챈다(self):
        one = {"form_no": "MBRM-00004", "rev": "4.0"}
        self.assertTrue(P.rev_changed(one, "MBRM-00004", "5.0"))
        self.assertFalse(P.rev_changed(one, "MBRM 00004", "4.0"))
        self.assertFalse(P.rev_changed(one, "", ""))       # 못 읽었으면 바뀌었다고 하지 않는다

    def test_판을_올리면_옛_프로필은_남는다(self):
        data = P.empty()
        old = P.put(data, {"maker": "동화약품", "record": "제조총괄표", "rev": "1.0",
                           "pages": {"조제": [3]}})
        P.put(data, P.bump(old, "2.0", {"조제": [4]}))
        self.assertEqual(len(data["profiles"]), 2)
        self.assertEqual(P.find(data, old["id"])["pages"], {"조제": [3]})


class 판독값_검산(unittest.TestCase):
    def test_숫자가_아니면_버리고_확인_필요로_남긴다(self):
        got = claude_cli._yield_clean({"lots": [{"lot": "gvy601",
                                                 "values": {"조제": "99.85%", "충전": "흐림"}}]})
        self.assertEqual(got["lots"][0]["values"], {"조제": 99.85})
        self.assertIn("충전", got["lots"][0]["unsure"])
        self.assertEqual(got["lots"][0]["lot"], "GVY601")

    def test_수율일_수_없는_값은_버린다(self):
        got = claude_cli._yield_clean({"lots": [{"lot": "A1234", "values": {"조제": "150"}}]})
        self.assertEqual(got["lots"][0]["values"], {})

    def test_기준을_벗어나면_짚는다(self):
        notes = Y.check({"lots": [{"lot": "A1234", "values": {"조제": 80.0}, "unsure": []}],
                         "specs": {"조제": "99.5±0.45%"}})
        self.assertTrue(any("기준" in n for n in notes), notes)

    def test_제조내역에_없는_제조번호를_짚는다(self):
        notes = Y.check(READING, known_lots=["GVY601"])
        self.assertTrue(any("GVYD01" in n and "제조내역" in n for n in notes), notes)

    def test_흐린_칸을_짚는다(self):
        notes = Y.check(READING)
        self.assertTrue(any("GVYD01 포장" in n for n in notes), notes)

    def test_기준_읽기(self):
        self.assertEqual(Y._spec_range("99.5±0.45%"), (99.05, 99.95))
        self.assertEqual(Y._spec_range("98.0~102.0"), (98.0, 102.0))
        self.assertEqual(Y._spec_range("98.0 이상"), (98.0, None))
        self.assertEqual(Y._spec_range("적합"), (None, None))


class 수율현황표_초안(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.path = Y.write_sheet(self.dir, READING, Y.check(READING))

    def test_보고서가_읽는_모양으로_쓴다(self):
        """이 파일은 '보고서 작성' 이 그대로 읽어야 한다 — 읽는 쪽으로 되읽어 확인한다."""
        got = dict(yield_sheet.read_yields(self.path))
        self.assertEqual(got["GVY601"]["조제"], "99.85")
        self.assertEqual(got["GVY601"]["충전"], "95.31")
        self.assertEqual(yield_sheet.read_specs(self.path)["조제"], "99.5±0.45%")

    def test_못_읽은_칸은_확인_필요로_둔다(self):
        got = dict(yield_sheet.read_yields(self.path))
        self.assertEqual(got["GVYD01"]["포장"], "확인 필요")

    def test_근거_쪽수를_비고에_남긴다(self):
        import openpyxl
        ws = openpyxl.load_workbook(self.path).active
        text = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
        self.assertIn("7, 12, 18쪽", text)

    def test_초안이라고_적어_둔다(self):
        import openpyxl
        ws = openpyxl.load_workbook(self.path).active
        text = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
        self.assertIn("초안", text)

    def test_판독_결과를_제품_폴더에_남긴다(self):
        Y.save_reading(self.dir, READING, ["확인 필요"], {"profile": "x"})
        got = Y.load_reading(self.dir)
        self.assertEqual(len(got["lots"]), 2)
        self.assertEqual(got["profile"], "x")


class 공정_이름_대응(unittest.TestCase):
    def test_수탁사_말을_우리_말로_바꾼다(self):
        got = Y.apply_names({"캡슐충전": 99.1, "조제": 98.0},
                            {"names": {"캡슐충전": "충전"}})
        self.assertEqual(got, {"충전": 99.1, "조제": 98.0})


class 담당자가_만든_표가_먼저다(unittest.TestCase):
    def test_직접_만든_수율현황표를_알아본다(self):
        folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, folder, True)
        open(os.path.join(folder, "7. 수율현황표 - 개인.xlsx"), "wb").close()
        self.assertTrue(Y.own_sheet(folder))

    def test_판독본은_담당자_표로_치지_않는다(self):
        folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, folder, True)
        open(os.path.join(folder, Y.SHEET_NAME), "wb").close()
        self.assertIsNone(Y.own_sheet(folder))


class 파일_이름의_쪽(unittest.TestCase):
    """담당자 2026-09-28: "제조기록서 스캔 파일 이름명 - 수율 페이지를 기재하면 알아서 수율 엑셀파일로
    만들어주는게 좋겠어" — 파일 이름에 적은 쪽이 프로필보다 먼저다."""

    def test_괄호_안_수율_n페이지(self):
        self.assertEqual(P.pages_from_name("한림 다파로엠서방정 9VY301 기록서 (수율 6페이지).pdf"), [6])

    def test_여러_쪽과_범위(self):
        self.assertEqual(P.pages_from_name("기록서 (수율 6, 8 페이지).pdf"), [6, 8])
        self.assertEqual(P.pages_from_name("기록서 수율 6-7쪽.pdf"), [6, 7])
        self.assertEqual(P.pages_from_name("기록서 (yield p6).pdf"), [6])

    def test_쪽이_없으면_빈_목록(self):
        self.assertEqual(P.pages_from_name("9VY301 기록서.pdf"), [])
        self.assertEqual(P.pages_from_name("7. 수율현황표 - 2025.xlsx"), [])

    def test_읽을_쪽을_정하는_차례(self):
        folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, folder, True)
        prof = {"id": "x", "pages": {"": [20]}}
        named = os.path.join(folder, "기록서 (수율 6페이지).pdf")
        plain = os.path.join(folder, "기록서.pdf")
        # ① 화면에서 적은 쪽
        self.assertEqual(Y.pages_for(named, prof, "9", folder)[0], [9])
        # ② 파일 이름
        self.assertEqual(Y.pages_for(named, prof, "", folder), ([6], "파일 이름의 쪽"))
        # ③ 지난번 메모
        Y.save_pages_note(folder, "11-12")
        self.assertEqual(Y.pages_for(plain, prof, "", folder)[0], [11, 12])
        Y.save_pages_note(folder, "")
        # ④ 프로필
        self.assertEqual(Y.pages_for(plain, prof, "", folder)[0], [20])
        # 아무것도 없으면 빈 목록
        self.assertEqual(Y.pages_for(plain, None, "", folder), ([], ""))

    def test_스캔_여럿을_하나로_합친다(self):
        a = {"lots": [{"lot": "9VY301", "values": {"타정": 95.4}}], "specs": {"타정": "75.0~100.0"}}
        b = {"lots": [{"lot": "9VY302", "values": {"타정": 96.0}}, {"lot": "9VY301", "values": {"타정": 1}}],
             "specs": {"코팅": "96.0~100.0"}}
        got = Y._merge_reading(dict(a), b)
        self.assertEqual([one["lot"] for one in got["lots"]], ["9VY301", "9VY302"])
        self.assertEqual(got["lots"][0]["values"]["타정"], 95.4)     # 먼저 읽은 것이 남는다
        self.assertEqual(set(got["specs"]), {"타정", "코팅"})


class 보고서가_바로_쓴다(unittest.TestCase):
    """담당자 2026-09-28: 판독본을 PQR 자동 작성 때 바로 쓴다 — 담당자 표가 있으면 그것만."""

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.folder, True)

    def _sheet(self, name, lot, value):
        # 판독본을 쓰는 그 함수로 만든다 — 보고서가 읽는 모양 그대로
        reading = {"lots": [{"lot": lot, "values": {"타정": value}, "production": {}, "page": {}}],
                   "specs": {"타정": "75.0~100.0"}}
        Y.write_sheet(self.folder, reading, (), path=os.path.join(self.folder, name))

    def test_판독본만_있으면_그것을_읽는다(self):
        self._sheet(Y.SHEET_NAME, "9VY301", 95.4)
        data = C.collect(self.folder)
        self.assertIn("9VY301", data.yields)

    def test_담당자_표가_있으면_판독본은_읽지_않는다(self):
        self._sheet(Y.SHEET_NAME, "9VY301", 1.0)
        self._sheet("7. 수율현황표 - 2026.xlsx", "9VY301", 95.4)
        data = C.collect(self.folder)
        self.assertEqual(str(data.yields["9VY301"]["타정"]), "95.40")
        self.assertTrue(any("판독본은 쓰지 않음" in str(row) for row in data.ledger))

    def test_공정이_하나뿐인_표도_이름을_잃지_않는다(self):
        """값 열이 하나면 머리 줄이 모두 '같은 글' 이 되어 공정 이름이 제목으로 버려졌다."""
        from pqr.engine.readers import yield_sheet
        self._sheet("7. 수율현황표 - 2026.xlsx", "9VY301", 95.4)
        got = dict(yield_sheet.read_yields(os.path.join(self.folder, "7. 수율현황표 - 2026.xlsx")))
        self.assertEqual(got["9VY301"]["타정"], "95.40")

    def test_ensure_sheet_는_표가_있으면_아무것도_하지_않는다(self):
        self._sheet("7. 수율현황표 - 2026.xlsx", "9VY301", 95.4)
        self.assertIsNone(Y.ensure_sheet(self.folder, "X"))

    def test_ensure_sheet_는_쪽을_모르면_까닭만_남긴다(self):
        open(os.path.join(self.folder, "7. 기록서.pdf"), "wb").write(b"%PDF-1.4\n")
        lines = []
        self.assertIsNone(Y.ensure_sheet(self.folder, "X", lines.append))
        self.assertTrue(any("건너뜁니다" in line for line in lines), lines)


class 화면에서_부르는_길(unittest.TestCase):
    """단추를 누르는 길은 HTTP 로 끝까지 밟아 본다 (규정집: 안쪽 함수만 시험하면 놓친다)."""

    def setUp(self):
        self.dir = tempfile.mkdtemp()
        write_samples(self.dir, layout="tree")
        self.folder = next(n for n in os.listdir(self.dir) if n.startswith("HP-110"))
        self.out = tempfile.mkdtemp()
        self.httpd = server_module.serve(self.dir, host="127.0.0.1", port=0, out_dir=self.out,
                                         today=TODAY, log=lambda *a: None)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        self.base = "http://127.0.0.1:%d" % self.httpd.server_port

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        shutil.rmtree(self.dir, ignore_errors=True)
        shutil.rmtree(self.out, ignore_errors=True)

    def call(self, path, body):
        request = urllib.request.Request(
            self.base + path, data=json.dumps(body).encode("utf-8"),
            headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as error:
            return json.loads(error.read().decode("utf-8"))

    def test_프로필을_만들고_제품에_배정한다(self):
        got = self.call("/api/yield-profile", {
            "profile": {"maker": "코오롱제약", "record": "정산기록서",
                        "form_no": "MBRM-00004", "rev": "4.0",
                        "pages": "조제 1, 충전 2", "rules": {"note": "공정수율 행만"}}})
        self.assertTrue(got.get("ok"), got.get("error"))
        pid = got["saved"]
        got = self.call("/api/yield-profile", {"product": "HP-110", "profile_id": pid})
        self.assertEqual(got["assigned"]["HP-110"], pid)
        again = self.call("/api/yield-profiles", {})
        self.assertEqual([p["id"] for p in again["profiles"]], [pid])
        self.assertTrue(os.path.isfile(os.path.join(self.dir, "공통", P.FILE_NAME)))

    def test_쪽수가_없으면_거절한다(self):
        got = self.call("/api/yield-profile", {
            "profile": {"maker": "동화약품", "record": "제조총괄표", "pages": ""}})
        self.assertFalse(got.get("ok"))
        self.assertIn("쪽", got.get("error", ""))

    def test_쪽을_모르면_무엇이_필요한지_알려_준다(self):
        """프로필이 없어도 파일 이름이나 수율 쪽 칸에 쪽이 있으면 읽는다 — 둘 다 없을 때만 묻는다."""
        open(os.path.join(self.dir, self.folder, "7. 기록서.pdf"), "wb").write(b"%PDF-1.4\n")
        got = self.call("/api/read-yield", {"product": "HP-110"})
        self.assertFalse(got.get("ok"))
        self.assertEqual(got.get("need"), "pages")
        self.assertIn("(수율 6페이지)", got.get("why", ""))

    def test_스캔이_없으면_먼저_올리라고_한다(self):
        got = self.call("/api/read-yield", {"product": "HP-110"})
        self.assertFalse(got.get("ok"))
        self.assertIn("스캔", got.get("why", ""))

    def test_화면에서_적은_쪽은_제품_폴더에_남는다(self):
        open(os.path.join(self.dir, self.folder, "7. 기록서.pdf"), "wb").write(b"%PDF-1.4\n")
        self.call("/api/read-yield", {"product": "HP-110", "pages": "6"})
        self.assertEqual(Y.load_pages_note(os.path.join(self.dir, self.folder)), "6")
        rows = self.call("/api/yield-profiles", {})["rows"]
        mine = next(r for r in rows if r["code"] == "HP-110")
        self.assertEqual(mine["pages_note"], "6")
        self.assertTrue(mine["ready"])

    def test_파일_이름의_쪽이_목록에_실린다(self):
        open(os.path.join(self.dir, self.folder, "7. 기록서 (수율 6페이지).pdf"), "wb").write(b"%PDF-1.4\n")
        rows = self.call("/api/yield-profiles", {})["rows"]
        mine = next(r for r in rows if r["code"] == "HP-110")
        self.assertEqual(mine["scan_pages"][0]["pages"], [6])
        self.assertEqual(mine["scan_pages"][0]["source"], "파일 이름의 쪽")

    def test_판독_기록을_제품_폴더에_남긴다(self):
        self.call("/api/read-yield", {"product": "HP-110"})
        self.assertTrue(os.path.isfile(os.path.join(self.dir, self.folder, Y.LOG_NAME)))

    def test_제품_목록을_내려보낸다(self):
        got = self.call("/api/yield-profiles", {})
        self.assertTrue(got.get("ok"))
        self.assertIsInstance(got.get("rows"), list)


if __name__ == "__main__":
    unittest.main()
