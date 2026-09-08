# -*- coding: utf-8 -*-
"""보고서를 만든 뒤 Claude Code 가 첨부와 대조해 문의 목록에 보태는 검토 단계
(담당자 2026-09-08: "단추 한 번에 만들고 → Claude 가 검토해 문의 목록에 적기")."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx                                                            # noqa: E402

from pqr.engine import review as R                                     # noqa: E402


def _report(path):
    d = docx.Document()
    d.add_paragraph("7. 수율현황표")
    t = d.add_table(rows=2, cols=3)
    for j, v in enumerate(["Lot No.", "조제", "충전"]):
        t.rows[0].cells[j].text = v
    for j, v in enumerate(["LWY201", "99.97", ""]):
        t.rows[1].cells[j].text = v
    d.add_paragraph("8. 원료")
    d.save(path)
    return path


class 보고서를_글로(unittest.TestCase):
    def test_항_차례대로_표를_줄로_푼다(self):
        text = R.dump_report(_report(os.path.join(tempfile.mkdtemp(), "r.docx")))
        self.assertIn("7. 수율현황표", text)
        self.assertIn("| Lot No. | 조제 | 충전 |", text)
        self.assertIn("| LWY201 | 99.97 | - |", text)             # 빈 칸은 '-'
        self.assertLess(text.index("7. 수율현황표"), text.index("| LWY201"))
        self.assertLess(text.index("| LWY201"), text.index("8. 원료"))


class 검토_결과_읽기(unittest.TestCase):
    def test_JSON을_문의_줄로(self):
        got = R._findings('설명\n```json\n{"findings": [{"item": "9.2.4", "lot": "LWY201", '
                          '"why": "제제균일성 칸이 비었는데 성적서에는 2.7% 가 있음", "cause": "코드"},'
                          '{"item": "13", "lot": "", "why": "LWV301 시험일지 없음", "cause": "자료"},'
                          '{"item": "", "why": "항 없음"}]}\n```')
        self.assertEqual(len(got), 2)
        self.assertEqual(got[0][0], "9.2.4"); self.assertEqual(got[0][1], "LWY201")
        self.assertIn("제작자에게", got[0][2])                     # 코드 원인은 제작자에게 보내라고
        self.assertNotIn("제작자에게", got[1][2])


class 검토_단계(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-rv-")
        self.folder = os.path.join(self.dir, "QC1-5087 아이퓨어점안액"); os.makedirs(self.folder)
        made = os.path.join(self.folder, "PQR 작성본"); os.makedirs(made)
        self.report = _report(os.path.join(made, "r.docx"))
        with open(os.path.join(self.folder, "7. 수율현황표.txt"), "w", encoding="utf-8") as h:
            h.write("LWY201 99.97 93.47")

    def test_Claude_Code가_없으면_빈_목록(self):
        from pqr.engine import claude_cli
        old = claude_cli._exe
        claude_cli._exe = lambda: None
        try:
            log = []
            self.assertEqual(R.review(self.folder, {"code": "QC1-5087"}, self.report, [], log.append), [])
            self.assertTrue(any("Claude Code 가 없어" in l for l in log))
        finally:
            claude_cli._exe = old

    def test_있으면_본문을_글로_놓고_물어_문의_줄을_돌려준다(self):
        from pqr.engine import claude_cli
        old_exe, old_ask = claude_cli._exe, claude_cli._ask
        asked = {}

        def fake_ask(exe, prompt, folder, log=None, paths=()):
            asked["prompt"], asked["paths"] = prompt, list(paths)
            return '{"findings": [{"item": "7", "lot": "LWY201", "why": "충전 칸이 비었는데 수율현황표에 93.47 있음", "cause": "코드"}]}'
        claude_cli._exe, claude_cli._ask = (lambda: "claude"), fake_ask
        try:
            got = R.review(self.folder, {"code": "QC1-5087"}, self.report,
                           [("13", "", "이미 있는 문의")], None)
        finally:
            claude_cli._exe, claude_cli._ask = old_exe, old_ask
        self.assertEqual(got[0][:2], ("7", "LWY201"))
        self.assertIn("93.47", got[0][2])
        text_path = os.path.join(self.folder, "PQR 작성본", "PQR 검토용 본문 - QC1-5087.txt")
        self.assertTrue(os.path.isfile(text_path))                     # 보고서를 글로 풀어 두었다
        self.assertIn(text_path, asked["paths"])
        self.assertTrue(any(p.endswith("7. 수율현황표.txt") for p in asked["paths"]))
        self.assertIn("이미 있는 문의", asked["prompt"])               # 있는 문의는 다시 적지 말라고 준다


if __name__ == "__main__":
    unittest.main()
