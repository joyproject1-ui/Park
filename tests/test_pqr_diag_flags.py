# -*- coding: utf-8 -*-
"""사선 칸 규칙 — 모든 PQR 에 함께 걸리는 것.

담당자 2026-09-15:
  · "사선 완료한 칸에 하이폰이나 해당없음 기재는 하지 마 — 모든 PQR 작성 시 공통 사항이야."
  · "1~3행 제조일자 칸이 사선입니다 … 이런식으로 일반적이지 않으면 노랑마크를 표시해줘야
     추가 메모와 함께 확인할 수 있어."
"""
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx                                              # noqa: E402
from docx.oxml.ns import qn                              # noqa: E402

from pqr.engine import docedit as E                      # noqa: E402


def 표(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, line in enumerate(rows):
        for c, text in enumerate(line):
            E.set_cell(E.raw_cells(table.rows[r])[c], text)
    return table


class 사선_칸에는_글자를_적지_않는다(unittest.TestCase):
    def test_하이픈도_지운다(self):
        document = docx.Document()
        table = 표(document, [["연번", "비고"], ["1", "-"], ["2", "해당 없음"], ["3", "N/A"]])
        for r in (1, 2, 3):
            E.add_diag(E.raw_cells(table.rows[r])[1])
        self.assertEqual(E.drop_na_in_diag_cells(document), 3)
        for r in (1, 2, 3):
            self.assertEqual(E.cell_text(E.raw_cells(table.rows[r])[1]).strip(), "")

    def test_사선이_없는_칸은_그대로_둔다(self):
        document = docx.Document()
        table = 표(document, [["연번", "비고"], ["1", "-"]])
        self.assertEqual(E.drop_na_in_diag_cells(document), 0)
        self.assertEqual(E.cell_text(E.raw_cells(table.rows[1])[1]).strip(), "-")

    def test_값이_있는_칸은_지우지_않는다(self):
        document = docx.Document()
        table = 표(document, [["연번", "비고"], ["1", "재발 방지 조치 완료"]])
        E.add_diag(E.raw_cells(table.rows[1])[1])
        self.assertEqual(E.drop_na_in_diag_cells(document), 0)


class 해당_없을_수_없는_칸의_사선(unittest.TestCase):
    def 만들기(self):
        document = docx.Document()
        table = 표(document, [["연번", "제조번호", "제조일자", "비고"],
                              ["1", "XMY401", "", ""],
                              ["2", "XMY402", "2025-03-04", ""]])
        E.add_diag(E.raw_cells(table.rows[1])[2])          # 제조일자 — 해당 없을 수 없다
        E.add_diag(E.raw_cells(table.rows[1])[3])          # 비고 — 사선이 곧 '해당 없음'
        E.add_diag(E.raw_cells(table.rows[2])[3])
        return document, table

    def test_제조일자의_사선만_짚는다(self):
        document, _ = self.만들기()
        found = E.flag_empty_diag_cells(document)
        self.assertEqual([name for name, _row in found], ["제조일자"])

    def test_노랑으로_칠한다(self):
        document, table = self.만들기()
        E.flag_empty_diag_cells(document)
        pr = E.raw_cells(table.rows[1])[2]._tc.find(qn("w:tcPr"))
        shd = pr.find(qn("w:shd"))
        self.assertIsNotNone(shd, "제조일자 칸이 칠해지지 않았습니다")
        self.assertEqual(shd.get(qn("w:fill")), "FFFF00")
        비고 = E.raw_cells(table.rows[1])[3]._tc.find(qn("w:tcPr"))
        self.assertIsNone(비고.find(qn("w:shd")) if 비고 is not None else None)

    def test_값이_있으면_짚지_않는다(self):
        document = docx.Document()
        table = 표(document, [["연번", "제조일자"], ["1", "2025-03-04"]])
        E.add_diag(E.raw_cells(table.rows[1])[1])
        self.assertEqual(E.flag_empty_diag_cells(document), [])

    def test_줄_이름으로도_알아본다(self):
        """세로로 선 표(왼쪽 칸이 항목 이름)도 같은 규칙."""
        document = docx.Document()
        table = 표(document, [["항목", "내용"], ["제조일자", ""], ["비고", ""]])
        for r in (1, 2):
            E.add_diag(E.raw_cells(table.rows[r])[1])
        self.assertEqual([name for name, _row in E.flag_empty_diag_cells(document)], ["제조일자"])


if __name__ == "__main__":
    unittest.main()
