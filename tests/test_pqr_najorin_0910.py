# -*- coding: utf-8 -*-
"""2026-09-10 나조린점안액 문의 목록 답변으로 굳힌 규칙들."""
from __future__ import unicode_literals

import re
import unittest

from pqr.engine.readers import erp
from pqr.engine import recipe_ointment as R
from pqr.engine import excel_attach as XA


class 제조내역ERP(unittest.TestCase):
    def test_ERP_제조번호를_Lot_으로(self):
        text = ("제조번호 (제조) 제조일자 사용기한\n"
                "NJS2-2025-L0K2Y-2000101 2025.02.13 2027.02.12\n"
                "NJS2-2025-L1K1Y-N000011 2025.11.13 2027.11.12\n"
                "NJS2-2025-L1K2Y-D000033 2025.12.08 2027.12.07\n")
        got = [erp.lot_from_erp(m) for m in erp.ERP_LINE.finditer(text)]
        self.assertEqual(got, ["LKY201", "LKYN01", "LKYD03"])

    def test_달이_맞지_않으면_버린다(self):
        m = erp.ERP_LINE.search("NJS2-2025-L0K2Y-3000101 2025.02.13 2027.02.12\n")
        self.assertIsNone(erp.lot_from_erp(m))


class 변경관리해당(unittest.TestCase):
    def test_등_N품목은_해당한다(self):
        self.assertTrue(R.change_covers({"products": "브리딘티 점안액 0.15% 등 29품목", "title": "…"}, "나조린점안액"))

    def test_이름이_없고_무리도_아니면_아니다(self):
        self.assertFalse(R.change_covers({"products": "후메론점안액, 오클점안액", "title": "…"}, "나조린점안액"))


class 무균되찾기(unittest.TestCase):
    def test_포장_판독의_bioburden_음성은_무균이다(self):
        self.assertEqual(R._sterility_from_bioburden({"bioburden": "음성"}), "음성")
        self.assertIsNone(R._sterility_from_bioburden({"bioburden": "0"}))
        self.assertIsNone(R._sterility_from_bioburden({}))


class 소수자릿수(unittest.TestCase):
    def test_성적서_글의_자릿수(self):
        self.assertEqual(XA._decimals("6.40"), 2)
        self.assertEqual(XA._decimals("7.0"), 1)
        self.assertEqual(XA._decimals("7"), 0)
        self.assertIsNone(XA._decimals("음성"))

    def test_표시_형식(self):
        self.assertEqual(XA._formats({"N6": 6.2, "P6": 7.0}, 2, 1), {"N6": "0.0", "P6": "0.0", "B": "0.00"})
        self.assertEqual(XA._formats({"P6": 15.0}, 0, 0), {})


class 기준변경각주(unittest.TestCase):
    def test_범위_꼴만_본다(self):
        self.assertTrue(R.RANGE_SPEC.match("6.2 ~ 7.0"))
        self.assertTrue(R.RANGE_SPEC.match("0.906~1.108"))
        self.assertFalse(R.RANGE_SPEC.match("10CFU/100mL 이하"))
        self.assertFalse(R.RANGE_SPEC.match("무색 투명한 액"))



class 판독이상치(unittest.TestCase):
    def test_다른_Lot_과_크게_다르면_애매함(self):
        logs = [{"lot": "A", "kind": "장기", "market": "내수", "points": [{"period": "6M", "assays": {"X": 107.0, "pH": 6.4}}]},
                {"lot": "B", "kind": "장기", "market": "내수", "points": [{"period": "6M", "assays": {"X": 103.4, "pH": 6.4}}]},
                {"lot": "C", "kind": "장기", "market": "내수", "points": [{"period": "6M", "assays": {"X": 102.9, "pH": 7.5}}]}]
        issues = []
        self.assertEqual(R._flag_outliers(logs, issues, lambda *a: None), 1)
        self.assertEqual(logs[0]["points"][0]["unsure"], ["X"])
        self.assertEqual(len(issues), 1)
        self.assertNotIn("pH", str(issues))

    def test_비슷하면_그대로(self):
        logs = [{"lot": "A", "kind": "장기", "market": "내수", "points": [{"period": "6M", "assays": {"X": 103.0}}]},
                {"lot": "B", "kind": "장기", "market": "내수", "points": [{"period": "6M", "assays": {"X": 103.4}}]}]
        self.assertEqual(R._flag_outliers(logs, [], lambda *a: None), 0)


class 유효기간(unittest.TestCase):
    def test_허가증_유효기간(self):
        class D: license = {"shelf_life": "제조일로부터 36개월"}
        self.assertEqual(R._shelf_months(D()), 36)
        class E: license = None
        self.assertIsNone(R._shelf_months(E()))



class 퀴노비드_0910(unittest.TestCase):
    def test_항보다_깊은_번호는_위_항으로(self):
        from pqr import build as B
        m = B.item_matcher(B.load_config()["items"])
        self.assertEqual(m("8.2.2.1 P12060 -1차 자재 ERP.xls"), "8.2.2")
        self.assertEqual(m("9.2.1.1 x.pdf"), "9.2.1")
        self.assertIsNone(m("3M필름.pdf"))

    def test_판독_시험항목_이름_맞춤(self):
        from pqr.engine import collect as C

        class D:
            coa = {"L1": {"924": {"assays": [{"part": "오플록사신"}, {"part": "벤잘코늄염화물"}]}}}
            stability_logs = [{"lot": "A", "points": [{"period": "12M", "assays": {"오플록사신": 100.1, "벤잘코늄염화물": 92.8}, "unsure": ["오플록사신"]}]},
                              {"lot": "B", "points": [{"period": "12M", "assays": {"함량": 100.6, "보존제": 89.5, "pH": 6.4}}]}]
        d = D()
        self.assertEqual(C.normalize_stability_keys(d), 2)
        self.assertEqual(d.stability_logs[0]["points"][0]["assays"], {"함량": 100.1, "보존제": 92.8})
        self.assertEqual(d.stability_logs[0]["points"][0]["unsure"], ["함량"])
        self.assertEqual(d.stability_logs[1]["points"][0]["assays"], {"함량": 100.6, "보존제": 89.5, "pH": 6.4})

    def test_형제_Lot_끼리만_이상치(self):
        logs = [{"lot": "EHV101", "kind": "시판후", "market": "내수", "year": "2022", "points": [{"period": "36M", "assays": {"보존제": 96.0}}]},
                {"lot": "EHW101", "kind": "시판후", "market": "내수", "year": "2023", "points": [{"period": "36M", "assays": {"보존제": 83.7}}]}]
        self.assertEqual(R._flag_outliers(logs, [], lambda *a: None, {"보존제": "80.0 ~ 120.0%"}), 0)

    def test_단위_뒤_이상(self):
        self.assertEqual(R._drop_unit("5.1mL이상", "질량∙용량(mL)개개"), "5.1 이상")
        self.assertEqual(R._drop_unit("5.3mL", "질량∙용량(mL)평균"), "5.3")

    def test_첨부_문서_줄_맞춤(self):
        import docx
        from pqr.engine import docedit as E
        d = docx.Document()
        d.add_paragraph("18. 첨부 문서")
        d.add_paragraph("- 안정성 시험 결과표(HLF-QC-104-01)")
        d.add_paragraph("- 안정성 시험 결과표(HLF-QC-104-22)")
        d.add_paragraph("- 안정성 시험 경향 분석 결과(HLF-QC-126-06)")
        r, a = E.sync_attachment_list(d, ["HLF-QC-126-06 안정성 시험 경향 분석 결과 - 퀴노비드점안액.xlsx",
                                         "HLF-QC-126-09 제품품질평가 경향분석 Sheet(양쪽 규격 용)_조제 pH.xls"])
        self.assertEqual((r, a), (2, 1))
        texts = [p.text for p in d.paragraphs]
        self.assertEqual(texts, ["18. 첨부 문서", "- 안정성 시험 경향 분석 결과(HLF-QC-126-06)",
                                 "- 제품품질평가 경향분석 Sheet(양쪽 규격 용)(HLF-QC-126-09)"])


if __name__ == "__main__":
    unittest.main()
