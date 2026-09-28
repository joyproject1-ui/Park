# -*- coding: utf-8 -*-
"""공급업체 목록 읽기 — 실시간 시트 · 정식 업체명 · 목록에 없는 관리번호 (담당자 2026-09-28).

"실시간 시트 또는 가장 최종 작성본으로 읽어야돼" / "제조원도 Supriya Lifescience Ltd 인데 Supriya Lifescience 로
기재되어 있어" / "실시간에서 RSN101은 확인되지 않아 그럼 노랑마크로 표시하고 설명해줘".
"""
import datetime
import os
import shutil
import tempfile
import unittest

os.environ["PQR_REVIEW"] = "0"

from openpyxl import Workbook                          # noqa: E402

from pqr.engine import recipe_ointment as R            # noqa: E402
from pqr.engine.readers import suppliers               # noqa: E402

HEAD = ["No.", "공급업체명\nSupplier name", "제조소 국가\nCountry", "제조소 주소\nManufacturer Address",
        "공급되는 품목\nItem supplied", "원료코드\nMaterial Code", "평가방법(현지/서면)\nEvaluation Type\n(On-site/Written)",
        "문서번호\nDoument No.", "개정번호\nRevision No.", "최초승인일\nInitial Approval Date",
        "평가승인일\nApproval date", "평가등급\nEvaluation grade", "차기 평가 예정 년도\nNext evaluation year", "비고\nRemark"]


def _sheet(ws, rows, head=HEAD):
    """스크린샷 모양 — 1줄 제목(병합), 2줄 개정번호, 3줄 머리행(줄바꿈·영문 섞임), 4줄부터 자료."""
    ws["E1"] = "원료약품 공급업체 목록(List of Raw Material Suppliers)"
    ws.merge_cells("E1:H1")
    ws["M2"], ws["N2"] = "개정번호(Revision No.) :", "Rev.26"
    ws.append(head)
    for row in rows:
        ws.append(row)


def _row(no, name, code, doc, first, day, grade="A", item="말레인산페니라민\nPheniramine Maleate"):
    return [no, name, "India", "A-5/2, Lote …", item, code, "서면", doc, "Rev.5", first, day, grade, "2026년", "나조린점안액"]


def make_book(path, live=True):
    wb = Workbook()
    old = wb.active
    old.title = "202503(Rev.21)"
    _sheet(old, [_row(1, "Supriya Lifescience", "RSP107", "VAR-R-Supriya", "2016. 05. 16", "2020. 10. 10")])
    ws = wb.create_sheet("202603(Rev.25)")
    _sheet(ws, [_row(1, "Supriya Lifescience Ltd", "RSP107", "VAR-R-Supriya", "2016. 05. 16", "2022. 10. 10")])
    if live:
        ws = wb.create_sheet("실시간")
        _sheet(ws, [
            _row(591, "Curia Spain S.A.U.", "RSM105", "VAR-R-Curia Spain", "2024. 10. 31", "박탈됨"),
            _row(592, "Micro Labs Limited", "RSN103", "VAR-R-Micro Labs", "2021. 06. 07", datetime.datetime(2025, 6, 18)),
            _row(595, "Supriya Lifescience Ltd", "RSP107", "VAR-R-Supriya", datetime.datetime(2016, 5, 16),
                 datetime.datetime(2024, 10, 10)),
            [None] * 14,
        ])
    wb.create_sheet("202210~개정이력").append(["개정 이력"])
    wb.save(path)


class 실시간_시트(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-sup-")
        self.path = os.path.join(self.dir, "8.1.1 원료 공급업체 List_(Rev.26).xlsx")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_실시간_시트가_있으면_그것을_읽는다(self):
        make_book(self.path)
        rows = suppliers.read_supplier_list(self.path)
        self.assertEqual(suppliers.sheet_of(rows), "실시간")
        self.assertEqual([r["원료코드"] for r in rows], ["RSM105", "RSN103", "RSP107"])   # 빈 줄은 뺀다
        got = suppliers.find_supplier(rows, code="RSP107")
        self.assertEqual(got["공급업체명"], "Supriya Lifescience Ltd")
        self.assertEqual(got["문서번호"], "VAR-R-Supriya")
        self.assertEqual(R.approval_date(got), "2024.10.10")        # 최종 평가 승인일(날짜 칸이라도)
        # '박탈됨' 은 날짜가 아니라 최초승인일로 물러선다 — 박탈된 줄이 코드로 골라지는 일은 없다(살아 있는 줄이 이긴다)
        self.assertEqual(R.approval_date(suppliers.find_supplier(rows, code="RSM105")), "2024.10.31")

    def test_실시간이_없으면_가장_높은_Rev_시트(self):
        make_book(self.path, live=False)
        rows = suppliers.read_supplier_list(self.path)
        self.assertEqual(suppliers.sheet_of(rows), "202603(Rev.25)")
        self.assertEqual(R.approval_date(rows[0]), "2022.10.10")

    def test_머리행_첫_칸이_달라도_열_이름으로_찾는다(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "실시간"
        head = list(HEAD)
        head[0] = "연번"
        _sheet(ws, [_row(1, "Merck KGaA", "RSP105", "VAR-R-Merck", "2016. 04. 16", "2025. 11. 28")], head)
        wb.save(self.path)
        rows = suppliers.read_supplier_list(self.path)
        self.assertEqual(rows[0]["공급업체명"], "Merck KGaA")
        head[0] = "품목"                      # 'No.' 도 '연번' 도 아니어도 열 이름 둘 이상이면 머리행
        wb = Workbook()
        _sheet(wb.active, [_row(1, "Merck KGaA", "RSP105", "VAR-R-Merck", "2016. 04. 16", "2025. 11. 28")], head)
        wb.save(self.path)
        self.assertEqual(suppliers.read_supplier_list(self.path)[0]["원료코드"], "RSP105")

    def test_머리행이_없으면_시트_이름을_말하며_멈춘다(self):
        wb = Workbook()
        wb.active.title = "실시간"
        wb.active.append(["아무것도", "없음"])
        wb.save(self.path)
        with self.assertRaises(ValueError) as caught:
            suppliers.read_supplier_list(self.path)
        self.assertIn("'실시간' 시트", str(caught.exception))


class 정식_업체명(unittest.TestCase):
    RECS = [{"공급업체명": "Supriya Lifescience Ltd"}, {"공급업체명": "Merck KGaA"},
            {"공급업체명": "Alps Pharmaceutical Ind. Co., Ltd."}, {"공급업체명": "주식회사 영일"}]

    def test_접미어가_빠진_이름은_목록의_정식_이름으로(self):
        self.assertEqual(R.official_company("Supriya Lifescience", self.RECS), "Supriya Lifescience Ltd")
        self.assertEqual(R.official_company("Merck", self.RECS), "Merck KGaA")
        self.assertEqual(R.official_company("Alps Pharmaceutical Ind.", self.RECS), "Alps Pharmaceutical Ind. Co., Ltd.")

    def test_같거나_아예_다른_이름은_그대로(self):
        self.assertEqual(R.official_company("Supriya Lifescience Ltd", self.RECS), "Supriya Lifescience Ltd")
        self.assertEqual(R.official_company("린하르트 GmbH (Pausa)", self.RECS), "린하르트 GmbH (Pausa)")
        self.assertEqual(R.official_company("", self.RECS), "")
        self.assertEqual(R.official_company("주식회사 영일", []), "주식회사 영일")

    def test_나라_이름이_붙은_목록_값도(self):
        self.assertEqual(R.official_company("Supriya Lifescience", [{"공급업체명": "Supriya Lifescience Ltd / India"}]),
                         "Supriya Lifescience Ltd")


class 목록에_없는_관리번호(unittest.TestCase):
    def test_설명에_시트와_줄_수와_열쇠가_들어간다(self):
        recs = [{"_sheet": "실시간", "원료코드": "RSN103"}, {"_sheet": "실시간", "원료코드": "RSP107"}]
        why = R.unlisted_note("RSN101", "VAR-R-Micro Labs", "Micro Labs", recs)
        self.assertIn("'실시간' 시트 2줄", why)
        self.assertIn("관리번호 RSN101", why)
        self.assertIn("평가문서번호 'VAR-R-Micro Labs'", why)
        self.assertIn("제조원 'Micro Labs'", why)
        self.assertIn("노랑", why)


class 주성분_공급망_마스터파일(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-chain-")
        self.path = os.path.join(self.dir, "8.1.2 HLF-GR-15-25(Rev.000) 주성분 공급망 마스터파일.xlsx")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_실시간_시트를_읽고_열_이름_띄어쓰기가_달라도(self):
        wb = Workbook()
        old = wb.active
        old.title = "Rev.000"
        old.append(["No.", "원료코드", "주성분 명", "제조소/국가", "1차 납품처", "2차 납품처"])
        old.append([1, "RSP107", "말레인산페니라민", "Supriya Lifescience / India", "제조소 납품", "N/A"])
        ws = wb.create_sheet("실시간")
        ws["A1"] = "주성분 공급망 마스터파일"
        ws.append([])
        ws.append(["No.", "원료코드\nMaterial Code", "주성분명\nAPI", "제조소 (국가)\nManufacturer", "1차\n납품처", "2차\n납품처"])
        ws.append([1, "RSP107", "말레인산페니라민", "Supriya Lifescience / India", "한국 공급사", "N/A"])
        ws.append([2, "RSN103", "나파졸린염산염", "Micro Labs Limited / India", "제조소 납품", None])
        wb.save(self.path)
        got = suppliers.read_api_chain(self.path)
        self.assertEqual(sorted(got), ["RSN103", "RSP107"])
        self.assertEqual(got["RSP107"]["_sheet"], "실시간")
        self.assertEqual(got["RSP107"]["api"], "말레인산페니라민")
        self.assertEqual(got["RSP107"]["manufacturer"], "Supriya Lifescience / India")
        self.assertEqual(got["RSP107"]["chain"], ["한국 공급사"])
        self.assertEqual(got["RSN103"]["chain"], ["제조소 납품"])


if __name__ == "__main__":
    unittest.main()
