# -*- coding: utf-8 -*-
"""10.2 제조설비 IQ·OQ — 마스터파일에서 읽어 채운다 (담당자 2026-09).

빈 공양식의 IQ·OQ 칸은 사선만 그어져 있고 비어 있다. 값은 '10.2 생산장비 적격성
평가 마스터파일(Rev.32)_OLD(IQ, OQ 정보 확인).xlsx' 에만 있고, 최신본(rev.033)에는
PQ 만 있다 — 두 파일을 합쳐 읽어야 IQ·OQ 가 채워진다.
"""
import os
import shutil
import tempfile
import unittest

from openpyxl import Workbook

import docx
from docx.oxml.ns import qn

from pqr.engine import collect, docedit as E
from pqr.engine.readers import masters
from pqr.engine.recipe_ointment import update_qualification


def _master(path, rows):
    """마스터파일 한 장 — 머리행(관리번호·보고서 문서번호·승인일자) 뒤에 자료 줄."""
    wb = Workbook()
    ws = wb.active
    ws.title = "생산장비"
    ws.append(["라인", "장비명", "관리번호", "보고서 문서번호", "승인일자"])
    for line in rows:
        ws.append(line)
    wb.save(path)


class 마스터파일_합치기(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="pqr-iqoq-")
        _master(os.path.join(self.root, "10.2 Qualification Master File-제조설비(rev.033).xlsx"),
                [["안연고제", "Main mixer", "DAA5114", "PQ26-2-DAA5114-R", "2026.05.14"]])
        _master(os.path.join(self.root,
                             "10.2 생산장비 적격성 평가 마스터파일(Rev.32)_OLD(IQ, OQ 정보 확인).xlsx"),
                [["안연고제", "Main Mixer", "DAA5114", "IQ18-2-DAA5114-R", "2018.10.16"],
                 ["", "", "", "OQ18-2-DAA5114-R", "2018.10.26"],
                 ["", "", "", "PQ22-2-DAA5114-R", "2022.10.28"]])

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_옛_파일의_IQ_OQ_를_이어받는다(self):
        data = collect.collect(self.root)
        docs = dict(data.equipment["DAA5114"]["docs"])
        self.assertEqual(docs.get("IQ18-2-DAA5114-R"), "2018.10.16")
        self.assertEqual(docs.get("OQ18-2-DAA5114-R"), "2018.10.26")

    def test_최신_파일의_PQ_가_옛것에_덮이지_않는다(self):
        data = collect.collect(self.root)
        docs = dict(data.equipment["DAA5114"]["docs"])
        self.assertIn("PQ26-2-DAA5114-R", docs)
        self.assertEqual(docs["PQ26-2-DAA5114-R"], "2026.05.14")
        self.assertEqual(data.equipment["DAA5114"]["name"], "Main mixer")   # 최신본 이름


def _span(cell, n):
    """가로 병합 — IQ·OQ 를 한 칸에 합쳐 적은 줄(IOQ…)을 흉내 낸다."""
    pr = cell._tc.get_or_add_tcPr()
    el = pr.makeelement(qn("w:gridSpan"), {qn("w:val"): str(n)})
    pr.append(el)


def _표(합침=False):
    """10.2 꼴의 표 — 머리 4줄 + 설비 한 대(문서번호 줄 · 완료일 줄)."""
    doc = docx.Document()
    t = doc.add_table(rows=6, cols=7)
    머리 = ["No.", "관리번호", "설비명", "완료일", "", "", "비고"]
    for j, v in enumerate(머리):
        E.set_cell(t.rows[0].cells[j], v)
    for j, v in enumerate(["", "", "", "IQ", "OQ", "PQ", ""]):
        E.set_cell(t.rows[1].cells[j], v)
    for ri, 이름 in ((2, "문서번호"), (3, "완료일")):
        for j, v in enumerate(["", "", "", 이름, 이름, 이름, ""]):
            E.set_cell(t.rows[ri].cells[j], v)
    for j, v in enumerate(["1", "DAA5114", "MAINMIXER", "", "", "", ""]):
        E.set_cell(t.rows[4].cells[j], v)
    if 합침:
        row = t.rows[4]
        E.set_cell(row.cells[3], "IOQ21-2-DAA5114-R")
        _span(row.cells[3], 2)          # IQ·OQ 를 한 칸으로
        E.set_cell(row.cells[5], "PQ22-2-DAA5114-R")
    return t


LOOKUP = {"DAA5114": {"IQ": [("IQ18-2-DAA5114-R", "2018.10.16")],
                      "OQ": [("OQ18-2-DAA5114-R", "2018.10.26")],
                      "PQ": [("PQ26-2-DAA5114-R", "2026.05.14")]}}


class 표_채우기(unittest.TestCase):
    def test_빈_IQ_OQ_칸을_마스터_값으로_채운다(self):
        t = _표()
        self.assertEqual(update_qualification(t, LOOKUP), 3)
        문서 = [E.cell_text(c).strip() for c in E.raw_cells(t.rows[4])]
        완료 = [E.cell_text(c).strip() for c in E.raw_cells(t.rows[5])]
        self.assertEqual(문서[3:6], ["IQ18-2-DAA5114-R", "OQ18-2-DAA5114-R", "PQ26-2-DAA5114-R"])
        self.assertEqual(완료[3:6], ["2018.10.16", "2018.10.26", "2026.05.14"])

    def test_사선을_지운다(self):
        t = _표()
        for ri in (4, 5):
            for j in (3, 4):
                E.add_diag(t.rows[ri].cells[j])
        update_qualification(t, LOOKUP)
        for ri in (4, 5):
            for j in (3, 4):
                self.assertFalse(E.has_diag(t.rows[ri].cells[j]))

    def test_IQ_OQ_가_한_칸이면_옆의_PQ_를_덮지_않는다(self):
        """제조지원 설비 표에는 IQ·OQ 를 합쳐 적은 줄이 있다 — 자리로 세면 PQ 가 날아간다."""
        t = _표(합침=True)
        update_qualification(t, LOOKUP)
        글 = [E.cell_text(c).strip() for c in E.raw_cells(t.rows[4])]
        self.assertEqual(글[3], "IOQ21-2-DAA5114-R")     # 합쳐 적은 칸은 그대로
        self.assertEqual(글[4], "PQ26-2-DAA5114-R")      # 그 옆은 PQ 자리 — PQ 로 갱신


def _지원마스터(path, rows):
    """제조지원 설비 꼴 — 관리번호 · IQ/OQ 문서번호 · 승인일 · 사유."""
    wb = Workbook()
    ws = wb.active
    ws.title = "제조지원 설비 & IT 시스템"
    for _ in range(3):
        ws.append([])
    ws.append(["시스템", "구분", "설비명", "관리번호",
               "IQ 문서번호", "IQ 승인일", "사유",
               "OQ 문서번호", "OQ 승인일", "사유",
               "PQ 문서번호", "PQ 승인일"])
    for line in rows:
        ws.append(line)
    wb.save(path)


class 지원설비_파일_나누기(unittest.TestCase):
    """IQ·OQ 는 'IQ, OQ 정보확인' 파일에서, PQ 는 'PQ …' 파일에서 읽는다 (담당자 2026-09)."""

    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="pqr-sp-")
        _지원마스터(os.path.join(self.root, "10.3, 10.4, 10.5 제조지원 설비 마스터파일(IQ, OQ 정보확인).xlsx"),
                    [["질소 분배 시스템", "질소 분배 시스템", "질소 분배 시스템", "HEA5029",
                      "IOQ20-UT-HEA5029-R\nIOQ23-TS-HEA5029-R", "2020.05.05\n2023.08.18",
                      "최초 제정 (20)\n액제 라인 리모델링에만 적용 (23)",
                      "IOQ20-UT-HEA5029-R\nIOQ23-TS-HEA5029-R", "2020.05.05\n2023.08.18",
                      "최초 제정 (20)\n액제 라인 리모델링에만 적용 (23)",
                      "QM24-F1,2,3-N2-R", "2024.05.23"]])
        _master(os.path.join(self.root,
                             "10.5 Qualification Master File(양식버전004)-제조지원설비(rev.002) "
                             "(PQ, 측정위치 타당성 정보 확인).xlsx"),
                [["N/A", "질소 분배 시스템", "HEA5029", "QM26-F1,2,3-N2-R", "2026.05.23"]])

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_PQ_는_PQ_마스터에서_읽는다(self):
        data = collect.collect(self.root)
        docs = dict(data.support["HEA5029"]["PQ"])
        self.assertIn("QM26-F1,2,3-N2-R", docs)         # QM… 으로 적힌 성능적격성평가도 PQ 다
        self.assertNotIn("QM24-F1,2,3-N2-R", docs)      # IQ·OQ 파일의 묵은 PQ 가 이기지 않는다

    def test_IQ_OQ_는_사유와_함께_읽힌다(self):
        data = collect.collect(self.root)
        v = data.support["HEA5029"]
        self.assertEqual([d for d, _ in v["IQ"]], ["IOQ20-UT-HEA5029-R", "IOQ23-TS-HEA5029-R"])
        self.assertEqual(v["why"]["IQ"][1], "액제 라인 리모델링에만 적용 (23)")


class 라인_가리기(unittest.TestCase):
    """마스터파일의 '사유' 는 그 적격성평가가 어디에 적용되는지를 적어 둔다.

    담당자 2026-09: "IQ, OQ에서 액제라인만 23년도에 공사한 경우는 표시했으니 잘 참고해서 작성해".
    """

    def test_우리_라인_것과_사유가_없는_것은_싣는다(self):
        for 사유 in ("최초 제정 (20)", "최초 제정 (20, 연고 리모델링)", "갭 보완 문서 (19)", ""):
            self.assertTrue(masters.applies(사유, "연고"), 사유)

    def test_다른_라인_공사는_빼_둔다(self):
        self.assertFalse(masters.applies("액제 라인 리모델링에만 적용 (23)", "연고"))
        self.assertTrue(masters.applies("액제 라인 리모델링에만 적용 (23)", "액제"))

    def test_다른_방_공사도_빼_둔다(self):
        for 사유 in ("3층 예비실, 입고대기실 리모델링 (21)",
                     "병충전 3실 POU 추가 - 분말스틱포장기 관련 (25)",
                     "병충전 4실 & 중앙포장실 pou 증설(26)"):
            self.assertFalse(masters.applies(사유, "연고"), 사유)


if __name__ == "__main__":
    unittest.main()
