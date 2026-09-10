# -*- coding: utf-8 -*-
"""서식에 남은 '이어짐'(vMerge) 흔적 — 위 칸이 병합이 아닌데 아래 빈 칸만 이어짐이면 낱칸으로 되돌려
빈 칸 사선이 그어지게 한다 (담당자 2026-09-10 나조린 9.2.2: "pH cpk 가장 아래칸이 공란인데 사선처리되지 않음")."""
from __future__ import unicode_literals

import unittest

import docx
from docx.oxml.ns import qn

from pqr.engine import docedit as E


def _table(rows, cols):
    doc = docx.Document()
    return doc, doc.add_table(rows=rows, cols=cols)


def _tc(t, r, c):
    """r행의 c번째 <w:tc> — python-docx 의 cell() 은 병합된 칸을 위 칸으로 바꿔 주므로 XML 로 본다."""
    from docx.table import _Cell
    return _Cell(t.rows[r]._tr.findall(qn("w:tc"))[c], t)


class 떠있는세로병합(unittest.TestCase):
    def test_위_칸이_병합이_아니면_이어짐을_뗀다(self):
        doc, t = _table(3, 2)
        t.cell(0, 1).text = "pH"
        E.set_vmerge(_tc(t, 2, 1), None)             # 위(1,1)는 병합이 아닌데 아래만 '이어짐'
        self.assertEqual(E.unmerge_stray_vmerge(doc), 1)
        self.assertIsNone(E._vmerge_state(_tc(t, 2, 1)._tc))
        # 되돌린 칸은 빈 칸 사선의 대상이 된다
        E.diag_all_empty(doc)
        self.assertTrue(E.has_diag(_tc(t, 2, 1)))

    def test_제대로_된_병합은_두지_않는다(self):
        doc, t = _table(3, 2)
        E.set_vmerge(_tc(t, 1, 1), "restart")
        E.set_vmerge(_tc(t, 2, 1), None)
        self.assertEqual(E.unmerge_stray_vmerge(doc), 0)
        self.assertEqual(E._vmerge_state(_tc(t, 2, 1)._tc), "continue")

    def test_이어짐이_사슬로_이어지면_모두_뗀다(self):
        doc, t = _table(4, 1)
        E.set_vmerge(_tc(t, 2, 0), None)
        E.set_vmerge(_tc(t, 3, 0), None)             # (2,0) 이 낱칸이 되면 (3,0) 도 떠 있다
        self.assertEqual(E.unmerge_stray_vmerge(doc), 2)

    def test_글이_든_칸은_두다(self):
        doc, t = _table(2, 1)
        _tc(t, 1, 0).text = "값"
        E.set_vmerge(_tc(t, 1, 0), None)
        self.assertEqual(E.unmerge_stray_vmerge(doc), 0)

    def test_gridSpan_을_세어_같은_열을_본다(self):
        doc, t = _table(3, 3)
        # 1행: 첫 칸이 두 열을 덮는다 → 2행 셋째 칸의 위는 1행 둘째 칸(열 2)
        first, second = t.rows[1]._tr.findall(qn("w:tc"))[:2]
        E._set_span(first, 2)
        second.getparent().remove(second)
        self.assertEqual(len(t.rows[1]._tr.findall(qn("w:tc"))), 2)
        E.set_vmerge(_tc(t, 1, 1), "restart")           # 1행 둘째 칸(열 2)이 병합 시작
        E.set_vmerge(_tc(t, 2, 2), None)                # 2행 셋째 칸(열 2)은 제대로 이어짐
        E.set_vmerge(_tc(t, 2, 1), None)                # 2행 둘째 칸(열 1)의 위는 두 열 덮는 낱칸 → 떠 있음
        self.assertEqual(E.unmerge_stray_vmerge(doc), 1)
        self.assertIsNone(E._vmerge_state(_tc(t, 2, 1)._tc))
        self.assertEqual(E._vmerge_state(_tc(t, 2, 2)._tc), "continue")


if __name__ == "__main__":
    unittest.main()
