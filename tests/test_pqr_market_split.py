"""내수용·수출용이 따로인 서식(퀴노비드) — 전년도 결재본 읽기, 13항 표 가르기, 시판 후 완료 판정, 시험일지 여러 쪽 (2026-09)."""
import os
import sys
import unittest

import docx

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import carry, docedit as E, handwriting     # noqa: E402
from pqr.engine.recipe_ointment import _stability_tables, _post_completed, _has_equipment, _seed_equipment_rows   # noqa: E402


def _doc(headings_tables):
    d = docx.Document()
    for heading, table in headings_tables:
        d.add_paragraph(heading)
        if table:
            t = d.add_table(rows=len(table), cols=len(table[0]))
            for i, row in enumerate(table):
                for j, v in enumerate(row):
                    E.set_cell(t.rows[i].cells[j], v)
    return d


class 표_가르기(unittest.TestCase):
    def test_내수용_수출용_시판후_장기_경향을_제목으로_가른다(self):
        d = _doc([("13. 안정성시험", None), ("13.1 장기 안정성 시험", None),
                  ("13.1.1 내수용", [["연번", "제조 번호"]]), ("13.1.2 수출용 (베트남)", [["연번", "제조 번호"]]),
                  ("13.2 시판 후 안정성 시험", None), ("13.2.1 내수용", [["연번", "비고"]]),
                  ("13.3 안정성 시험 경향 분석 결과", None), ("13.3.1 내수용", [["시험항목", "함량(%)"]]),
                  ("14. 반품", [["a"]])])
        got = [(k, m) for k, m, _ in _stability_tables(d)]
        self.assertEqual(got, [("장기", "내수"), ("장기", "수출"), ("시판후", "내수"), ("경향", "내수")])

    def test_디겐타_서식은_시장_없이_장기와_경향뿐(self):
        d = _doc([("13.1 장기 안정성 시험", [["연번"]]), ("13.2 시판 후 안정성 시험 이력 없음.", None),
                  ("13.3 안정성 시험 경향 분석 결과", [["시험항목"]])])
        self.assertEqual([(k, m) for k, m, _ in _stability_tables(d)], [("장기", ""), ("경향", "")])


class 시판후_완료(unittest.TestCase):
    def test_사용기한까지_시험했으면_완료(self):
        one = {"mfg": "2022.03.18", "expiry": "2025.03.17", "points": [{"period": "36M"}]}
        self.assertTrue(_post_completed(one, []))
        one = {"mfg": "2023.01.06", "expiry": "2026.01.05", "points": [{"period": "24M"}]}
        self.assertFalse(_post_completed(one, []))


class 전년도_결재본(unittest.TestCase):
    def test_포장과_설비줄과_10_1_표를_읽는다(self):
        d = _doc([("10.1 공정밸리데이션", None), ("10.1.1 내수용", [["사유", "대상Lot", "문서번호"], ["변경", "OEX101", "PV24-2-QUIO3-R"]]),
                  ("10.1.2 수출용", [["사유", "대상Lot", "문서번호"], ["변경", "OAX101", "PV24-2-QUIO2-R"]]),
                  ("10.3 공기조화장치", [["No.", "관리번호", "설비명", "IQ", "OQ", "PQ"], ["", "", "", "문서번호", "문서번호", "문서번호"],
                                     ["", "", "", "완료일", "완료일", "완료일"], ["", "", "", "", "", ""],
                                     ["1", "HBA5089", "A.H.U 26호", "IQ18-HA-HBA5089-R", "OQ18-HA-HBA5089-R", "PQ18-HA-HBA5089-FR"],
                                     ["1", "", "", "2018.09.03", "2018.11.01", "2019.12.04"]]),
                  ("13.1 시판 후 안정성 시험", None),
                  ("13.1.1 내수용", [["연번", "제조 번호", "포장 형태"], ["1", "OEV301", "5g tube/갑"]]),
                  ("13.1.2 수출용 (베트남)", [["연번", "제조 번호", "포장 형태"], ["1", "OZW101", "3.5g tube/갑"]])])
        packs = carry.stability_packs(d)
        self.assertEqual(packs["by_lot"], {"OEV301": "5g tube/갑", "OZW101": "3.5g tube/갑"})
        self.assertEqual(packs["by_market"], {"내수": "5g tube/갑", "수출": "3.5g tube/갑"})
        self.assertEqual(packs["market_by_lot"], {"OEV301": "내수", "OZW101": "수출"})
        self.assertEqual(packs["market_by_prefix"], {"OE": "내수", "OZ": "수출"})
        eq = carry.equipment_rows(d)
        self.assertEqual(eq["10.3"][0]["mid"], "HBA5089")
        self.assertEqual(eq["10.3"][0]["docs"]["OQ"], ("OQ18-HA-HBA5089-R", "2018.11.01"))
        grids = carry.section_grids_all(d, "10.1")
        self.assertEqual(len(grids), 2)
        self.assertIn("PV24-2-QUIO2-R", grids[1][1][2])

    def test_빈_설비_표에_전년도_줄을_세운다(self):
        d = _doc([("10.3 공기조화장치", [["No.", "관리번호", "설비명", "완료일", "", "", "비고"], ["No.", "관리번호", "설비명", "IQ", "OQ", "PQ", "비고"],
                                     ["", "", "", "문서번호", "문서번호", "문서번호", ""], ["", "", "", "완료일", "완료일", "완료일", ""],
                                     ["1", "", "", "", "", "", ""], ["1", "", "", "", "", "", ""]])])
        t = d.tables[0]
        self.assertFalse(_has_equipment(t))
        n = _seed_equipment_rows(t, [{"mid": "HBA5089", "name": "A.H.U 26호", "docs": {"IQ": ("IQ18-HA-HBA5089-R", "2018. 09. 03")}},
                                     {"mid": "HBA5033", "name": "A.H.U 4호", "docs": {}}])
        self.assertEqual(n, 2)
        self.assertTrue(_has_equipment(t))
        rows = [[E.cell_text(c).strip() for c in E.raw_cells(r)] for r in t.rows[4:]]
        self.assertEqual(rows[0][:4], ["1", "HBA5089", "A.H.U 26호", "IQ18-HA-HBA5089-R"])
        self.assertEqual(rows[1][3], "2018.09.03")
        self.assertEqual(rows[2][:3], ["2", "HBA5033", "A.H.U 4호"])
        self.assertEqual(len(rows), 4)


class 시험일지_여러_쪽(unittest.TestCase):
    def test_read_folder_는_쪽마다_한_Lot(self):
        import types
        fake = [{"lot": "OEV301", "year": "2022", "page": 1, "points": []}, {"lot": "OEW101", "year": "2023", "page": 2, "points": []}]
        old = handwriting.read_pages
        handwriting.read_pages = lambda p, specs, log: list(fake)
        try:
            out = handwriting.read_folder(["/x/a.pdf"], None, None)
        finally:
            handwriting.read_pages = old
        self.assertEqual([o["lot"] for o in out], ["OEV301", "OEW101"])


class 삼항_비고_NA(unittest.TestCase):
    def test_NA_칸은_지우고_빈_줄과_합쳐_사선_하나(self):
        d = docx.Document()
        t = d.add_table(rows=6, cols=3)
        for i, row in enumerate([["No.", "점검 항목", "비 고"], ["1", "제형", "N/A"], ["2", "제품분류", ""],
                                 ["3", "제품명", "<주성분>"], ["4", "허가번호", "N/A"], ["5", "허가일자", ""]]):
            for j, v in enumerate(row):
                E.set_cell(t.rows[i].cells[j], v)
        groups, rows = E.merge_empty_runs(t, "비고")
        self.assertEqual((groups, rows), (2, 4))
        self.assertEqual(E.cell_text(t.rows[1].cells[2]).strip(), "")
        self.assertEqual(E.cell_text(t.rows[4].cells[2]).strip(), "")
        self.assertTrue(E.has_diag(t.rows[1].cells[2]) and E.has_diag(t.rows[4].cells[2]))
        self.assertFalse(E.has_diag(E.raw_cells(t.rows[2])[2]))
        self.assertEqual(E.cell_text(t.rows[3].cells[2]).strip(), "<주성분>")


class 구이_요약줄(unittest.TestCase):
    def test_값이_모두_같아도_최댓값_최솟값_평균을_적는다(self):
        from pqr.engine import detail92 as D
        d = docx.Document()
        t = d.add_table(rows=8, cols=4)
        heads = [["연번", "Lot No.", "금속성이물 합계", "금속성이물 개개"], ["1", "", "", ""], ["2", "", "", ""],
                 ["최댓값", "", "", ""], ["최솟값", "", "", ""], ["평균", "", "", ""], ["공정능력지수(Cpk)", "", "", ""], ["Cpk 판정 결과", "", "", ""]]
        for i, row in enumerate(heads):
            for j, v in enumerate(row):
                E.set_cell(t.rows[i].cells[j], v)
        D.fill(t, ["A1", "A2"], lambda lab, lot, i: "0" if "금속" in lab else None)
        got = {E.cell_text(r.cells[0]).strip(): E.cell_text(r.cells[2]).strip() for r in t.rows[3:6]}
        self.assertEqual(got, {"최댓값": "0", "최솟값": "0", "평균": "0"})


class Cpk_사선(unittest.TestCase):
    def test_사선이_그어진_Cpk_칸에는_값을_적지_않는다(self):
        from pqr.engine import detail92 as D
        d = docx.Document()
        t = d.add_table(rows=7, cols=4)
        heads = [["연번", "Lot No.", "함량(%)", "입자도(㎛이하)"], ["1", "", "", ""], ["2", "", "", ""],
                 ["최댓값", "", "", ""], ["최솟값", "", "", ""], ["공정능력지수(Cpk)", "", "", ""], ["Cpk 판정 결과", "", "", ""]]
        for i, row in enumerate(heads):
            for j, v in enumerate(row):
                E.set_cell(t.rows[i].cells[j], v)
        E.add_diag(t.rows[5].cells[3]); E.add_diag(t.rows[6].cells[3])       # 입자도 열의 Cpk 칸은 사선
        D.fill(t, ["A1", "A2"], lambda lab, lot, i: "100" if "함량" in lab else "60",
               cpk=lambda lab, texts: ("1.23", "충분"))
        self.assertEqual(E.cell_text(t.rows[5].cells[2]).strip(), "1.23")
        self.assertEqual(E.cell_text(t.rows[5].cells[3]).strip(), "")
        self.assertEqual(E.cell_text(t.rows[6].cells[3]).strip(), "")
        self.assertTrue(E.has_diag(t.rows[5].cells[3]))


class 갈음_제품(unittest.TestCase):
    """수출용을 동일 수탁 제품(에펙신안연고)으로 갈음한 건 — 그 이름이 든 시험일지는 수출용이다
    (담당자 2026-09-07: "에펙신은 퀴노비드와 동일한 제품, 일동제약 수탁품")."""

    HEAD = ["연번", "해당 연도", "시험 기간", "제조 번호", "포장 형태", "보관 조건", "완료 일자", "실시 사유"]

    def _doc_with_note(self):
        빈줄 = [""] * 8
        note = ["특이사항 (Comment)\n* 퀴노비드안연고(수출용)의 장기 안정성 시험은 "
                "동일 수탁 제품(에펙신안연고)로 갈음하였음."] + [""] * 7
        return _doc([("13. 안정성시험", None), ("13.1 장기 안정성 시험", None),
                     ("13.1.1 내수용", [self.HEAD, 빈줄, ["특이사항 (Comment)\nN/A"] + [""] * 7]),
                     ("13.1.2 수출용 (베트남)", [self.HEAD, 빈줄, note])])

    def test_서식_특이사항에서_갈음_제품_이름을_읽는다(self):
        from pqr.engine.recipe_ointment import _stand_in_names
        got = _stand_in_names(_stability_tables(self._doc_with_note()))
        self.assertEqual(got, {"에펙신안연고": "수출"})

    def test_갈음_문구가_없으면_빈_값(self):
        d = _doc([("13.1 장기 안정성 시험", None),
                  ("13.1.2 수출용 (베트남)", [["연번"], ["1"], ["특이사항 (Comment)\nN/A"]])])
        from pqr.engine.recipe_ointment import _stand_in_names
        self.assertEqual(_stand_in_names(_stability_tables(d)), {})

    def test_갈음_제품_일지는_수출용_표로_간다(self):
        from pqr.engine.recipe_ointment import _fill_stability26
        d = self._doc_with_note()
        logs = [{"lot": "OAX101", "year": "2024", "kind": "장기", "market": "내수", "market_hint": False,
                 "pack": "3.5g tube/갑", "store": "", "why": "",
                 "source": "[LT-24-2]에펙신안연고_OAX101(MTC001)_24M.pdf",
                 "points": [{"period": "24M", "done": "2025.03.10",
                             "assays": {"오플록사신": 101.2}, "unsure": []}]}]
        _fill_stability26(d, logs, {"from": "2025-01-01", "to": "2025-12-31"},
                          {"오플록사신": "90.0 ~ 110.0%"}, lambda *a: None, [], {}, {})
        self.assertEqual(logs[0]["market"], "수출")
        수출표, 내수표 = _stability_tables(d)[1][2], _stability_tables(d)[0][2]
        줄 = lambda t: " ".join(E.cell_text(c) for r in t.rows for c in E.raw_cells(r))
        self.assertIn("OAX101", 줄(수출표))
        self.assertNotIn("OAX101", 줄(내수표))
