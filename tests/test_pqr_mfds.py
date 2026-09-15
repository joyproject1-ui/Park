# -*- coding: utf-8 -*-
"""3항 '대상 제품' 을 식약처 허가정보와 대조한다 (담당자 2026-09-12: "자동으로 해").

회사 밖으로 나가지 않는다 — 응답을 흉내 낸 글로만 시험한다.
"""
from __future__ import unicode_literals

import json
import os
import tempfile
import unittest

from pqr.engine.readers import mfds


올로원스 = {
    "ITEM_NAME": "올로원스점안액(올로파타딘염산염)",
    "ENTP_NAME": "한림제약(주)",
    "ITEM_PERMIT_DATE": "20111130",
    "ITEM_SEQ": "201110503",
    "ETC_OTC_CODE": "전문의약품",
    "STORAGE_METHOD": "기밀용기, 2~25℃보관",
    "VALID_TERM": "제조일로부터 24 개월",
    "REEXAM_TARGET": "",
    "RMP_TARGET": "",
    "PACK_UNIT": "[다회용]3mL/병",
}

표3 = {"제품명": "올로원스점안액(올로파타딘염산염)", "허가일자": "2011년 11월 30일",
       "보관조건": "기밀용기, 2~25℃보관", "사용기한": "제조일로부터 24개월",
       "제품분류": "전문의약품"}


def _응답(items):
    return json.dumps({"body": {"items": [{"item": one} for one in items]}}).encode("utf-8")


class 받아오기(unittest.TestCase):
    def test_응답에서_품목을_꺼낸다(self):
        got = mfds.fetch("올로원스점안액", "열쇠",
                         opener=lambda url, timeout: _응답([올로원스]))
        self.assertEqual(len(got), 1)
        self.assertEqual(got[0]["허가일자"], "20111130")
        self.assertEqual(got[0]["업체명"], "한림제약(주)")

    def test_열쇠가_없으면_부르지_않는다(self):
        def 부르면안됨(url, timeout):
            raise AssertionError("열쇠 없이 불렀습니다")
        self.assertEqual(mfds.fetch("올로원스점안액", "", opener=부르면안됨), [])

    def test_이름이_가장_잘_맞는_것을_고른다(self):
        일회용 = dict(올로원스, ITEM_NAME="올로원스점안액(올로파타딘염산염)(1회용)")
        rows = mfds.fetch("x", "열쇠", opener=lambda u, t: _응답([일회용, 올로원스]))
        got = mfds.pick(rows, "올로원스점안액(올로파타딘염산염)")
        self.assertEqual(got["제품명"], "올로원스점안액(올로파타딘염산염)")


class 대조(unittest.TestCase):
    def _info(self, **바꿀것):
        rows = mfds.fetch("x", "열쇠", opener=lambda u, t: _응답([dict(올로원스, **바꿀것)]))
        return rows[0]

    def test_모두_같으면_어긋남_없음(self):
        self.assertEqual(mfds.compare(표3, self._info()), [])

    def test_허가일자가_다르면_올린다(self):
        got = mfds.compare(표3, self._info(ITEM_PERMIT_DATE="20111129"))
        self.assertEqual([one[0] for one in got], ["허가일자"])

    def test_보관조건_글이_다르면_올린다(self):
        got = mfds.compare(표3, self._info(STORAGE_METHOD="차광기밀용기, 실온보관"))
        self.assertEqual([one[0] for one in got], ["보관조건"])

    def test_허가정보_칸이_비면_견주지_않는다(self):
        """서비스가 그 항목을 주지 않는 것과 '해당 없음' 은 다르다."""
        self.assertEqual(mfds.compare(표3, self._info(STORAGE_METHOD="")), [])

    def test_취소_이력은_크게_알린다(self):
        got = mfds.compare(표3, self._info(CANCEL_DATE="20250401", CANCEL_NAME="자진취하"))
        self.assertIn("허가 상태", [one[0] for one in got])

    def test_재심사_RMP_가_비면_해당_없음으로_적는다(self):
        self.assertIn("재심사대상: 해당 없음", mfds.notes(self._info()))
        self.assertIn("RMP대상: 대상", mfds.notes(self._info(RMP_TARGET="대상")))


class 열쇠모양(unittest.TestCase):
    """마이페이지의 두 인증키 어느 쪽을 넣어도 되게 한다 (담당자 2026-09-14)."""

    def test_Decoding_키는_그대로_쓴다(self):
        self.assertEqual(mfds.plain_key("abc+de/fg=="), "abc+de/fg==")

    def test_Encoding_키는_되돌려_쓴다(self):
        self.assertEqual(mfds.plain_key("abc%2Bde%2Ffg%3D%3D"), "abc+de/fg==")

    def test_앞뒤_빈칸과_줄바꿈은_지운다(self):
        self.assertEqual(mfds.plain_key("  키값\n"), "키값")

    def test_빈_열쇠는_빈_글(self):
        self.assertEqual(mfds.plain_key(None), "")


class 오류알림(unittest.TestCase):
    """인증이 안 되면 포털은 HTTP 200 에 XML 오류를 준다 — 그 까닭을 그대로 올린다."""

    KEY_ERROR = ('<OpenAPI_ServiceResponse><cmmMsgHeader>'
                 '<returnAuthMsg>SERVICE_KEY_IS_NOT_REGISTERED_ERROR</returnAuthMsg>'
                 '<returnReasonCode>30</returnReasonCode></cmmMsgHeader></OpenAPI_ServiceResponse>')

    def test_아는_오류는_우리말로_알린다(self):
        got = mfds.service_error(self.KEY_ERROR)
        self.assertIn("등록되지 않은 서비스 키", got)
        self.assertIn("SERVICE_KEY_IS_NOT_REGISTERED_ERROR", got)

    def test_JSON_이_아니면_그_까닭을_올린다(self):
        with self.assertRaises(ValueError) as caught:
            mfds.fetch("올로원스점안액", "열쇠", opener=lambda u, t: self.KEY_ERROR.encode())
        self.assertIn("등록되지 않은 서비스 키", str(caught.exception))

    def test_오류_글이_없으면_빈_글(self):
        self.assertEqual(mfds.service_error("<x><y>1</y></x>"), "")


처분 = {"PRDUCT": "올로원스점안액", "ENTP_NAME": "한림제약(주)",
        "DISPOS_DATE": "20250401", "DISPOS_CONT": "판매업무정지 1개월",
        "VIOLATION_CONT": "표시 기재 위반"}


def _처분응답(items):
    return json.dumps({"body": {"items": [{"item": one} for one in items]}}).encode("utf-8")


class 행정처분(unittest.TestCase):
    """3항 6번 '행정 처분 이력 없음' 을 식약처 행정처분 정보로 확인한다 (담당자 2026-09-14)."""

    def test_이력을_읽는다(self):
        rows, base = mfds.fetch_penalties("올로원스점안액", "열쇠",
                                          opener=lambda u, t: _처분응답([처분]))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["처분일자"], "20250401")
        self.assertEqual(rows[0]["처분내용"], "판매업무정지 1개월")
        self.assertTrue(base)

    def test_우리_제품_것만_고른다(self):
        남의것 = dict(처분, PRDUCT="다른회사점안액", ENTP_NAME="다른제약")
        rows, _ = mfds.fetch_penalties("x", "열쇠", opener=lambda u, t: _처분응답([처분, 남의것]))
        got = mfds.penalties_for(rows, "올로원스점안액(올로파타딘염산염)")
        self.assertEqual([one["제품명"] for one in got], ["올로원스점안액"])

    def test_주소_후보를_차례로_두드린다(self):
        mfds.forget_base()
        불린곳 = []
        def opener(url, timeout):
            불린곳.append(url.split("?")[0])
            if len(불린곳) == 1:
                raise OSError("없는 주소")
            return _처분응답([처분])
        rows, base = mfds.fetch_penalties("올로원스점안액", "열쇠", opener=opener)
        self.assertEqual(len(rows), 1)
        self.assertEqual(base, mfds.PENALTY_BASES[1])
        self.assertEqual(len(불린곳), 2)

    def test_담당자가_넣은_주소를_먼저_쓴다(self):
        root = tempfile.mkdtemp()
        os.makedirs(os.path.join(root, "공통"))
        with open(os.path.join(root, "공통", mfds.PENALTY_URL_FILE), "w", encoding="utf-8") as h:
            h.write("http://example.test/행정처분?serviceKey=xxx\n")
        old = os.environ.pop("MFDS_PENALTY_URL", None)
        try:
            self.assertEqual(mfds.penalty_base(root), ["http://example.test/행정처분"])
        finally:
            if old is not None:
                os.environ["MFDS_PENALTY_URL"] = old

    def test_이력이_없으면_빈_목록(self):
        rows, _ = mfds.fetch_penalties("올로원스점안액", "열쇠", opener=lambda u, t: _처분응답([]))
        self.assertEqual(mfds.penalties_for(rows, "올로원스점안액"), [])


class 주소와_열쇠_가리기(unittest.TestCase):
    """파일 두 개가 나란히 있어 서로 바꿔 넣기 쉽다 (담당자 2026-09-14: "둘다 키주소는 동일하네")."""

    def _folder(self, key_text, url_text):
        root = tempfile.mkdtemp()
        os.makedirs(os.path.join(root, "공통"))
        for name, text in ((mfds.KEY_FILE, key_text), (mfds.PENALTY_URL_FILE, url_text)):
            with open(os.path.join(root, "공통", name), "w", encoding="utf-8") as handle:
                handle.write(text)
        return root

    def setUp(self):
        self.old = (os.environ.pop("MFDS_API_KEY", None),
                    os.environ.pop("MFDS_PENALTY_URL", None))

    def tearDown(self):
        for name, value in zip(("MFDS_API_KEY", "MFDS_PENALTY_URL"), self.old):
            if value is not None:
                os.environ[name] = value

    def test_제대로_넣으면_둘_다_읽는다(self):
        root = self._folder("abc123+key==", "http://apis.data.go.kr/1471000/Adm/getAdm")
        self.assertEqual(mfds.api_key(root), "abc123+key==")
        self.assertEqual(mfds.penalty_base(root), ["http://apis.data.go.kr/1471000/Adm/getAdm"])

    def test_주소_자리에_열쇠를_넣으면_쓰지_않는다(self):
        root = self._folder("abc123+key==", "abc123+key==")
        self.assertEqual(mfds.api_key(root), "abc123+key==")
        self.assertEqual(mfds.penalty_base(root), [])       # 후보로 물러서지도 않는다

    def test_열쇠_자리에_주소를_넣으면_쓰지_않는다(self):
        root = self._folder("http://apis.data.go.kr/1471000/x", "http://apis.data.go.kr/1471000/y")
        self.assertIsNone(mfds.api_key(root))

    def test_주소의_물음표_뒤는_떼어낸다(self):
        root = self._folder("k", "http://apis.data.go.kr/1471000/Adm/getAdm?serviceKey=zz&type=json")
        self.assertEqual(mfds.penalty_base(root), ["http://apis.data.go.kr/1471000/Adm/getAdm"])


class 부르는_횟수(unittest.TestCase):
    """느려지지 않게 — 주소를 모르면 부르지 않고, 한 번 통한 주소는 기억한다
    (담당자 2026-09-14: "너무 오래 걸리면 생략할까 고민하고 있어")."""

    def setUp(self):
        mfds.forget_base()
        self.old = os.environ.pop("MFDS_PENALTY_URL", None)

    def tearDown(self):
        mfds.forget_base()
        if self.old is not None:
            os.environ["MFDS_PENALTY_URL"] = self.old

    def _folder(self, url_text=None):
        root = tempfile.mkdtemp()
        os.makedirs(os.path.join(root, "공통"))
        if url_text is not None:
            with open(os.path.join(root, "공통", mfds.PENALTY_URL_FILE), "w", encoding="utf-8") as h:
                h.write(url_text)
        return root

    def test_주소_파일이_없으면_아예_부르지_않는다(self):
        self.assertEqual(mfds.penalty_base(self._folder()), [])

    def test_자동이라고_적으면_후보를_두드린다(self):
        self.assertEqual(mfds.penalty_base(self._folder("자동")), list(mfds.PENALTY_BASES))

    def test_한_번_통한_주소는_다음_제품부터_한_번만_부른다(self):
        불린것 = []
        body = _처분응답([])
        def opener(url, timeout):
            base = url.split("?")[0]
            불린것.append(base)
            if base != mfds.PENALTY_BASES[2]:
                raise OSError("없는 주소")
            return body
        for 제품 in ("가", "나", "다"):
            mfds.fetch_penalties(제품, "열쇠", bases=list(mfds.PENALTY_BASES), opener=opener)
        self.assertEqual(len(불린것), 5)                  # 첫 제품 3번, 그 뒤 1번씩
        self.assertEqual(불린것[3:], [mfds.PENALTY_BASES[2]] * 2)

    def test_후보를_두드릴_때는_짧게_기다린다(self):
        기다린것 = []
        def opener(url, timeout):
            기다린것.append(timeout)
            raise OSError("없는 주소")
        mfds.fetch_penalties("가", "열쇠", bases=list(mfds.PENALTY_BASES), opener=opener)
        self.assertEqual(기다린것, [mfds.PROBE_TIMEOUT] * 3)

    def test_주소가_하나면_제대로_기다린다(self):
        기다린것 = []
        def opener(url, timeout):
            기다린것.append(timeout)
            return _처분응답([])
        mfds.fetch_penalties("가", "열쇠", bases=["http://example.test/a"], opener=opener)
        self.assertEqual(기다린것, [mfds.TIMEOUT])


class 행정처분_노랑표시(unittest.TestCase):
    """이력이 있으면 3항 6번 칸을 노랑으로 칠하고 문의에 올린다."""

    def _표(self):
        import docx
        from pqr.engine import docedit as E
        doc = docx.Document()
        t = doc.add_table(rows=4, cols=4)
        rows = (("제품명", "올로원스점안액(올로파타딘염산염)"),
                ("허가 및 시판 후 준수 사항의 이행 여부 검토",
                 "1. 허가상 제조방법 준수하여 생산함을 확인함."),
                ("보관조건", "기밀용기, 2~25℃보관"))
        for i, (이름, 값) in enumerate(rows):
            cells = E.raw_cells(t.rows[i + 1])
            E.set_cell(cells[1], 이름)
            E.set_cell(cells[2], 값)
        return doc, t

    def _돌린다(self, items):
        from pqr.engine import recipe_ointment as R, docedit as E
        doc, t = self._표()
        칸3 = {}
        for row in t.rows[1:]:
            cells = E.raw_cells(row)
            이름 = E.cell_text(cells[1]).strip()
            if 이름:
                칸3[이름] = cells[2]
        old = (mfds.api_key, mfds.fetch_penalties, mfds.penalty_base)
        mfds.api_key = lambda folder=None: "열쇠"
        mfds.penalty_base = lambda folder=None: ["http://example.test"]
        mfds.fetch_penalties = lambda name, key, bases=None, **kw: (
            [{k: mfds._penalty_value(one, k) for k in mfds.PENALTY_FIELDS} for one in items],
            "http://example.test")
        issues = []
        try:
            n = R._check_penalty("올로원스점안액(올로파타딘염산염)", "한림제약(주)", "",
                                 issues, lambda *a: None, 칸3)
        finally:
            mfds.api_key, mfds.fetch_penalties, mfds.penalty_base = old
        칠한칸 = {이름 for 이름, 칸 in 칸3.items() if "highlight" in 칸._tc.xml}
        return n, issues, 칠한칸

    def test_이력이_있으면_6번_칸을_칠한다(self):
        n, issues, 칠한칸 = self._돌린다([처분])
        self.assertEqual(n, 1)
        self.assertEqual(len(issues), 1)
        self.assertIn("행정처분 이력", issues[0][2])
        self.assertIn("판매업무정지 1개월", issues[0][2])
        self.assertEqual(칠한칸, {"허가 및 시판 후 준수 사항의 이행 여부 검토"})

    def test_이력이_없으면_아무것도_하지_않는다(self):
        self.assertEqual(self._돌린다([]), (0, [], set()))


class 열쇠찾기(unittest.TestCase):
    def test_제품_폴더의_파일에서_읽는다(self):
        root = tempfile.mkdtemp()
        folder = os.path.join(root, "QC1-5059 올로원스점안액")
        os.makedirs(os.path.join(folder, "공통"))
        with open(os.path.join(folder, "공통", mfds.KEY_FILE), "w", encoding="utf-8") as handle:
            handle.write("  열쇠값  \n")
        old = os.environ.pop("MFDS_API_KEY", None)
        try:
            self.assertEqual(mfds.api_key(folder), "열쇠값")
        finally:
            if old is not None:
                os.environ["MFDS_API_KEY"] = old

    def test_없으면_None(self):
        old = os.environ.pop("MFDS_API_KEY", None)
        try:
            self.assertIsNone(mfds.api_key(tempfile.mkdtemp()))
        finally:
            if old is not None:
                os.environ["MFDS_API_KEY"] = old


class 노랑표시(unittest.TestCase):
    """어긋난 칸만 노랑으로 칠하고 문의에 올린다 (담당자 2026-09-12)."""

    def _표(self):
        import docx
        from pqr.engine import docedit as E
        doc = docx.Document()
        t = doc.add_table(rows=4, cols=4)
        for i, (이름, 값) in enumerate((("제품명", "올로원스점안액(올로파타딘염산염)"),
                                        ("허가일자", "2011년 11월 30일"),
                                        ("보관조건", "기밀용기, 2~25℃보관"))):
            cells = E.raw_cells(t.rows[i + 1])
            E.set_cell(cells[1], 이름)
            E.set_cell(cells[2], 값)
        return doc, t

    def _돌린다(self, **바꿀것):
        from pqr.engine import recipe_ointment as R, docedit as E
        doc, t = self._표()
        표3, 칸3 = {}, {}
        for row in t.rows[1:]:
            cells = E.raw_cells(row)
            이름 = E.cell_text(cells[1]).strip()
            if 이름:
                표3[이름] = E.cell_text(cells[2]).strip()
                칸3[이름] = cells[2]
        old_key, old_fetch = mfds.api_key, mfds.fetch
        mfds.api_key = lambda folder=None: "열쇠"
        mfds.fetch = lambda name, key, **kw: [dict(
            {k: mfds._value(dict(올로원스, **바꿀것), k) for k in mfds.FIELDS})]
        issues = []
        try:
            n = R._check_license(표3, "올로원스점안액(올로파타딘염산염)", "", issues, lambda *a: None, 칸3)
        finally:
            mfds.api_key, mfds.fetch = old_key, old_fetch
        칠한칸 = {이름 for 이름, 칸 in 칸3.items() if "highlight" in 칸._tc.xml}
        return n, issues, 칠한칸

    def test_어긋난_칸만_칠한다(self):
        n, issues, 칠한칸 = self._돌린다(ITEM_PERMIT_DATE="20111129")
        self.assertEqual(n, 1)
        self.assertEqual(칠한칸, {"허가일자"})

    def test_칸_바탕도_노랑이고_문의에_두_값이_다_적힌다(self):
        """담당자 2026-09-16: "보고서에 노랑으로 표시해주고 문의목록에 내용을 기재해줘"."""
        from pqr.engine import recipe_ointment as R, docedit as E
        doc, t = self._표()
        표3, 칸3 = {}, {}
        for row in t.rows[1:]:
            cells = E.raw_cells(row)
            이름 = E.cell_text(cells[1]).strip()
            if 이름:
                표3[이름] = E.cell_text(cells[2]).strip(); 칸3[이름] = cells[2]
        old_key, old_fetch = mfds.api_key, mfds.fetch
        mfds.api_key = lambda folder=None: "열쇠"
        mfds.fetch = lambda name, key, **kw: [dict(
            {k: mfds._value(dict(올로원스, ITEM_PERMIT_DATE="20111129"), k) for k in mfds.FIELDS})]
        issues = []
        try:
            R._check_license(표3, "올로원스점안액(올로파타딘염산염)", "", issues, lambda *a: None, 칸3)
        finally:
            mfds.api_key, mfds.fetch = old_key, old_fetch
        xml = 칸3["허가일자"]._tc.xml
        self.assertIn("w:shd", xml)                       # 칸 바탕
        self.assertIn("highlight", xml)                   # 글자 형광펜
        self.assertEqual(len(issues), 1)
        item, where, text = issues[0]
        self.assertEqual((item, where), ("3", "3항 허가일자"))
        self.assertIn("2011년 11월 30일", text)           # 보고서 값
        self.assertIn("2011", text) and self.assertIn("식약처", text)
        self.assertIn("노랑", text)
        self.assertEqual(len(issues), 1)
        self.assertIn("허가일자", issues[0][1])

    def test_모두_같으면_아무것도_칠하지_않는다(self):
        n, issues, 칠한칸 = self._돌린다()
        self.assertEqual((n, issues, 칠한칸), (0, [], set()))


if __name__ == "__main__":
    unittest.main()


class 판독_대장에_남김(unittest.TestCase):
    """식약처 확인 결과가 자료 판독 대장에도 한 줄 남는다 (담당자 2026-09-16: 대장만 보고
    "허가 식약처 사이트에서 확인한 것인지?")."""

    def _run(self, key, fetch=None, pick=None):
        from pqr.engine import recipe_ointment as R
        from pqr.engine.readers import mfds
        saved = (mfds.api_key, mfds.fetch, mfds.pick, mfds.compare)
        mfds.api_key = lambda folder: key
        mfds.fetch = fetch or (lambda name, k: [{"제품명": name, "허가일자": "20111130"}])
        mfds.pick = pick or (lambda rows, name: rows[0] if rows else None)
        mfds.compare = lambda s3, info: []
        ledger = []
        try:
            R._check_license({"제품명": "올로원스점안액"}, "올로원스점안액", "", [], lambda *a: None, {}, ledger=ledger)
        finally:
            mfds.api_key, mfds.fetch, mfds.pick, mfds.compare = saved
        return ledger

    def test_대조하면_읽음으로_남는다(self):
        got = self._run("열쇠")
        self.assertEqual(len(got), 1)
        item, name, status, detail = got[0]
        self.assertEqual((item, status), ("3", "읽음"))
        self.assertIn("식약처 허가정보", name)
        self.assertIn("어긋난 칸 없음", detail)

    def test_열쇠가_없으면_못_읽음(self):
        got = self._run("")
        self.assertEqual(got[0][2], "못 읽음")
        self.assertIn("서비스 키", got[0][3])

    def test_받지_못하면_안_읽음(self):
        def 터짐(name, k): raise RuntimeError("연결 실패")
        got = self._run("열쇠", fetch=터짐)
        self.assertEqual(got[0][2], "안 읽음")
        self.assertIn("연결 실패", got[0][3])

    def test_대장이_없어도_돈다(self):
        from pqr.engine import recipe_ointment as R
        from pqr.engine.readers import mfds
        saved = mfds.api_key; mfds.api_key = lambda folder: ""
        try:
            self.assertEqual(R._check_license({}, "x", "", [], lambda *a: None, {}), 0)
        finally:
            mfds.api_key = saved


class 삼항_제품명_확인(unittest.TestCase):
    """3항 표의 제품명이 이 제품이 아니면 잡아낸다 (담당자 PC 2026-09-16: 서식에 남은 한림포비돈 3항)."""

    def test_다른_제품이면_그_이름을_돌려준다(self):
        from pqr.engine.recipe_ointment import section3_mismatch
        self.assertEqual(section3_mismatch({"제품명": "한림포비돈점안액"}, "올로원스점안액"), "한림포비돈점안액")

    def test_같은_제품이면_빈_글(self):
        from pqr.engine.recipe_ointment import section3_mismatch
        self.assertEqual(section3_mismatch({"제품명": "올로원스점안액(올로파타딘염산염)"}, "올로원스점안액"), "")
        self.assertEqual(section3_mismatch({"제품명": "올로원스 점안액"}, "올로원스점안액(다회용)"), "")

    def test_제품명_줄이_없으면_빈_글(self):
        from pqr.engine.recipe_ointment import section3_mismatch
        self.assertEqual(section3_mismatch({"허가번호": "제 99 호"}, "올로원스점안액"), "")
