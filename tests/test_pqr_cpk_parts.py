# -*- coding: utf-8 -*-
"""10 Lot 이상 — 주성분이 둘인 제품의 Cpk (2026-09 점검, 디겐타안연고 합성 12 Lot).

- 함량 Cpk 는 성분마다 그 성분의 규격으로 (플루오로메톨론 90~110, 겐타마이신 90~120)
- Cpk 계산 파일은 성분마다 하나, 시트 머리(제품명·시험항목)와 규격은 올해 값
- 양식이 Cpk 칸과 판정 칸을 세로 병합해 두었으면 풀어서 '충분' 이 보이게
- 16.1 표에는 '함량(성분)' 으로
"""
from __future__ import unicode_literals

import unittest

import docx

from pqr.engine import conclusion, detail92, excel_attach
from pqr.engine import docedit as E


class Data(object):
    def __init__(self, coa):
        self.coa = coa
        self.issues = []
        self.lots = [(l, "", False) for l in coa]

    @property
    def domestic(self):
        return list(self.coa)


def two_part_coa(n=12):
    coa = {}
    for i in range(n):
        lot = "OGY9%02d" % i
        coa[lot] = {"924": {"assay": "%.1f" % (98 + i * 0.1), "particle": "15", "particle_spec": "75",
                            "metal_total": "0", "metal_each": "0",
                            "assays": [{"part": "플루오로메톨론", "value": "%.1f" % (98 + i * 0.1), "lo": "90.0", "hi": "110.0"},
                                       {"part": "겐타마이신황산염", "value": "%.1f" % (100 + i * 0.2), "lo": "90.0", "hi": "120.0"}]}}
    return coa


class 성분별_Cpk_파일(unittest.TestCase):
    def test_함량은_성분마다_한_파일(self):
        data = Data(two_part_coa())
        jobs = excel_attach.cpk_jobs(data, data.domestic, "디겐타안연고")
        labels = [j["label"] for j in jobs]
        self.assertIn("함량(플루오로메톨론)", labels)
        self.assertIn("함량(겐타마이신황산염)", labels)
        self.assertNotIn("함량", labels)
        genta = next(j for j in jobs if j["label"] == "함량(겐타마이신황산염)")
        self.assertEqual((genta["cells"]["N6"], genta["cells"]["P6"]), (90.0, 120.0))
        self.assertEqual(genta["cells"]["C4"], "디겐타안연고 (내수용)")
        self.assertEqual(genta["cells"]["K5"], "함량(%) - 겐타마이신황산염")
        self.assertEqual(len(genta["values"]), 12)
        self.assertAlmostEqual(genta["values"][1], 100.2)

    def test_입자도_규격은_성적서의_상한(self):
        data = Data(two_part_coa())
        jobs = excel_attach.cpk_jobs(data, data.domestic, "디겐타안연고")
        particle = next(j for j in jobs if j["word"] == "입자도")
        self.assertEqual(particle["cells"]["P6"], 75.0)
        self.assertIsNone(particle["cells"]["O6"])

    def test_주성분_하나면_파일_하나(self):
        coa = {l: {"924": {"assay": "101.0", "assays": [{"part": "오플록사신", "value": "101.0", "lo": "90.0", "hi": "110.0"}]}}
               for l in ("A%02d" % i for i in range(10))}
        jobs = excel_attach.cpk_jobs(Data(coa), list(coa), "퀴노비드안연고")
        assay = [j for j in jobs if j["word"] == "함량"]
        self.assertEqual(len(assay), 1)
        self.assertEqual(assay[0]["label"], "함량")
        self.assertEqual((assay[0]["cells"]["N6"], assay[0]["cells"]["P6"]), (90.0, 110.0))


class 결론_이름(unittest.TestCase):
    def test_성분별_함량은_괄호로(self):
        got = conclusion.low_cpk_items({"assay/겐타마이신황산염": 0.8, "assay/플루오로메톨론": 1.5, "particle": 0.9})
        self.assertEqual(got, [("함량(겐타마이신황산염)", 0.8), ("입자도", 0.9)])

    def test_예전_모양도_그대로(self):
        self.assertEqual(conclusion.low_cpk_items({"assay": 0.93, "particle": 1.2, "metal": 24.1}), [("함량", 0.93)])


class 병합된_Cpk_칸(unittest.TestCase):
    def test_판정_칸이_병합돼_있으면_풀고_쓴다(self):
        d = docx.Document()
        t = d.add_table(rows=6, cols=3)
        for c, h in enumerate(("연번", "Lot No.", "함량(%)")):
            t.cell(0, c).text = h
        t.cell(1, 0).text = "1"; t.cell(2, 0).text = "2"
        for r, h in ((3, "평균"), (4, "공정능력지수(Cpk)"), (5, "Cpk 판정 결과")):
            t.cell(r, 0).text = h
        E.set_vmerge(t.cell(4, 2), "restart"); E.set_vmerge(t.cell(5, 2), None)
        lots = ["L%02d" % i for i in range(12)]
        detail92.fill(t, lots, lambda lab, lot, i: "100.0" if "함량" in lab else None,
                      cpk=lambda lab, texts: ("1.50", "충분") if "함량" in lab else None)
        rows = [r for r in t.rows]
        cpk_row = next(r for r in rows if "공정능력" in r.cells[0].text)
        judge_row = next(r for r in rows if "판정" in r.cells[0].text)
        self.assertEqual(cpk_row.cells[2].text, "1.50")
        self.assertEqual(judge_row.cells[2].text, "충분")
        self.assertIsNone(judge_row.cells[2]._tc.tcPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vMerge"))


if __name__ == "__main__":
    unittest.main()
