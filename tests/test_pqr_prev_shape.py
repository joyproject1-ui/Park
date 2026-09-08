"""작성 전에 전년도 결재본을 분석하고, 다 채운 뒤 견주어 빈 항을 잡아낸다
(담당자 2026-09-08: "PQR 작성 전에 16. 전년도 PQR 분석을 한 뒤에 작성해 줘야 돼")."""
import os
import sys
import unittest

import docx

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import docedit as E, writer                          # noqa: E402


def _doc(heading, rows):
    d = docx.Document()
    d.add_paragraph(heading)
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            E.set_cell(t.rows[i].cells[j], v)
    return d


HEAD = ["연번", "Lot No.", "제조일자"]


class 전년도_분석(unittest.TestCase):
    def test_자료_줄만_센다(self):
        d = _doc("6. 제조내역 확인", [HEAD, ["1", "OEX101", "2025.01.02"],
                                  ["2", "OEX102", "2025.02.03"],
                                  ["특이사항 (Comment)", "", ""]])
        self.assertEqual(writer._rows_of(d, "6"), 2)

    def test_요약_줄과_빈_줄은_세지_않는다(self):
        d = _doc("7. 수율현황표", [HEAD, ["", "", ""], ["최댓값", "99.9", ""],
                                ["평균", "99.5", ""]])
        self.assertEqual(writer._rows_of(d, "7"), 0)

    def test_NA_만_있는_줄도_세지_않는다(self):
        d = _doc("6. 제조내역 확인", [HEAD, ["1", "N/A", "-"]])
        self.assertEqual(writer._rows_of(d, "6"), 0)

    def test_분석_결과를_기록에_남긴다(self):
        said = []
        writer._analyse_previous(_doc("6. 제조내역 확인", [HEAD, ["1", "OEX101", "2025.01.02"]]),
                                 said.append)
        self.assertTrue(any("전년도 결재본 분석" in line and "6 1줄" in line for line in said))


class 올해와_견주기(unittest.TestCase):
    def _shape(self, rows):
        return writer._analyse_previous(_doc("6. 제조내역 확인", rows), lambda *a: None)

    def test_전년도에는_있는데_올해_비면_알린다(self):
        shape = self._shape([HEAD, ["1", "OEX101", "2025.01.02"]])
        now = _doc("6. 제조내역 확인", [HEAD, ["", "", ""]])
        got = writer._compare_with_previous(now, shape, lambda *a: None)
        self.assertEqual([item for item, _ in got], ["6"])
        self.assertIn("전년도 결재본에는 6항에 1줄", got[0][1])
        self.assertTrue(got[0][1].startswith("★"))

    def test_올해도_채워졌으면_알리지_않는다(self):
        shape = self._shape([HEAD, ["1", "OEX101", "2025.01.02"]])
        now = _doc("6. 제조내역 확인", [HEAD, ["1", "OEY201", "2026.01.05"]])
        self.assertEqual(writer._compare_with_previous(now, shape, lambda *a: None), [])

    def test_전년도에도_없었으면_알리지_않는다(self):
        shape = self._shape([HEAD, ["", "", ""]])
        now = _doc("6. 제조내역 확인", [HEAD, ["", "", ""]])
        self.assertEqual(writer._compare_with_previous(now, shape, lambda *a: None), [])

    def test_전년도를_읽지_못했으면_아무것도_하지_않는다(self):
        now = _doc("6. 제조내역 확인", [HEAD, ["", "", ""]])
        self.assertEqual(writer._compare_with_previous(now, None, lambda *a: None), [])


if __name__ == "__main__":
    unittest.main()
