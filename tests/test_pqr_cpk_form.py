"""Cpk 계산 파일 — 전년도 서식(.xlsx)의 칸 값만 갈아 끼운다 (담당자 2026-09-06:
"Cpk 는 전년도 양식으로 작성하되 2026년 PQR 작성본 내용을 참고해서 업데이트하면 돼")."""
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import cpk_form, excel_attach                      # noqa: E402

DATA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pqr", "data")
BIL = os.path.join(DATA, excel_attach.FORMS["Bilateral"])
UNI = os.path.join(DATA, excel_attach.FORMS["Unilateral"])


class 서식_채우기(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-cpkform-")

    def _fill(self, form, cells, values):
        dst = os.path.join(self.dir, "out.xlsx")
        kind = cpk_form.fill(form, dst, cells, values, "2026.09.06")
        return kind, openpyxl.load_workbook(dst), openpyxl.load_workbook(dst, data_only=True)

    def test_양쪽_규격은_머리와_값과_계산값을_모두_채운다(self):
        kind, wb, wv = self._fill(BIL, {"C4": "퀴노비드안연고 (내수용)", "C5": "포장 완료 후", "K5": "함량(%)",
                                        "N6": 90, "O6": "N/A", "P6": 110}, [101.0, 99.0, 103.0, 100.0])
        ws, vs = wb.worksheets[0], wv.worksheets[0]
        self.assertEqual(kind, "Bilateral")
        self.assertEqual((ws["C4"].value, ws["K5"].value, ws["K4"].value),
                         ("퀴노비드안연고 (내수용)", "함량(%)", "2026.09.06"))
        self.assertEqual((ws["N6"].value, ws["O6"].value, ws["P6"].value), (90, "N/A", 110))
        self.assertEqual([ws.cell(row=r, column=2).value for r in (10, 13, 14)], [101.0, 100.0, None])
        self.assertTrue(str(ws["G9"].value).startswith("="))            # 수식은 그대로
        self.assertAlmostEqual(vs["G9"].value, 100.75, places=2)        # 계산값도 함께 들어간다
        self.assertAlmostEqual(vs["P9"].value, 1.81, places=2)
        self.assertIn("공정 능력 충분", vs["I11"].value)
        self.assertEqual(len(ws._charts), 1)                            # 그래프·로고는 그대로
        self.assertEqual(len(ws._images), 1)

    def test_한쪽_규격은_서식에_적힌_규격을_그대로_쓴다(self):
        kind, wb, wv = self._fill(UNI, {"C4": "퀴노비드안연고 (내수용)", "K5": "입자도(㎛)", "P6": 75},
                                  [60, 50, 55, 58])
        ws, vs = wb.worksheets[0], wv.worksheets[0]
        self.assertEqual(kind, "Unilateral")
        self.assertEqual(ws["P6"].value, 75)
        self.assertIsNone(ws["O6"].value)
        self.assertAlmostEqual(vs["P9"].value, 1.475, places=3)
        self.assertEqual(vs["S9"].value, "UCL")
        self.assertEqual(vs["T9"].value, "USL")

    def test_규격을_주지_않으면_서식에_적힌_값으로_계산한다(self):
        # 전년도 파일에는 규격이 적혀 있다 — 올해 값이 없으면 그것을 그대로 쓴다(금속성이물)
        작년 = os.path.join(self.dir, "작년.xlsx")
        cpk_form.fill(UNI, 작년, {"K5": "금속성이물(합계)", "P6": 50}, [2, 1, 3, 1], "2025.02.11")
        kind, wb, wv = self._fill(작년, {"K5": "금속성이물(합계)"}, [1, 0, 2, 1])
        ws, vs = wb.worksheets[0], wv.worksheets[0]
        self.assertEqual(ws["P6"].value, 50)                            # 서식에 남아 있던 규격
        self.assertIsNotNone(vs["P9"].value)

    def test_모든_값이_같으면_계산값을_비운다(self):
        kind, wb, wv = self._fill(UNI, {"K5": "금속성이물(개개)", "P6": 1}, [0, 0, 0, 0])
        vs = wv.worksheets[0]
        self.assertIsNone(vs["P9"].value)                               # σ=0 — 지난해 값이 남지 않는다
        self.assertIsNone(vs["I11"].value)


if __name__ == "__main__":
    unittest.main()
