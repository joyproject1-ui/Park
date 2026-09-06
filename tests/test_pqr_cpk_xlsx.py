"""Cpk 계산 파일 — Excel·LibreOffice 없이도 서식·수식·그래프를 그려 반드시 만든다 (담당자 2026-09-06)."""
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import cpk_xlsx, excel_attach                      # noqa: E402


class 직접_그리기(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-cpk-")

    def test_양쪽_규격(self):
        dst = os.path.join(self.dir, "a. 함량 Cpk 계산 파일.xlsx")
        kind = cpk_xlsx.build(dst, [101.0, 99.5, 100.2], {"C4": "퀴노비드안연고 (내수용)", "C5": "포장 완료 후",
                                                         "K5": "함량(%)", "N6": 90, "P6": 110}, "2026.09.06")
        self.assertEqual(kind, "Bilateral")
        ws = openpyxl.load_workbook(dst).worksheets[0]
        self.assertEqual(ws.title, "Bilateral")
        self.assertEqual(ws["K4"].value, "2026.09.06")
        self.assertEqual((ws["N6"].value, ws["P6"].value), (90, 110))
        self.assertEqual([ws["B%d" % r].value for r in (10, 11, 12, 13)], [101.0, 99.5, 100.2, None])
        self.assertEqual(ws["P9"].value, "=MIN((P6-G9)/(3*G8),(G9-N6)/(3*G8))")
        self.assertEqual(ws["V10"].value, "=N$6")
        self.assertEqual(len(ws._charts), 1)
        self.assertEqual(len(ws._charts[0].series), 6)                  # 결과·평균·UCL·LCL·USL·LSL
        self.assertTrue(ws["B13"].border.diagonalUp)                    # 빈 칸 사선
        self.assertIn("HLF-QC-126-09", ws["A45"].value)

    def test_한쪽_규격(self):
        dst = os.path.join(self.dir, "a. 입자도 Cpk 계산 파일.xlsx")
        kind = cpk_xlsx.build(dst, [61, 48], {"K5": "입자도(㎛)", "O6": None, "P6": 75}, "2026.09.06")
        self.assertEqual(kind, "Unilateral")
        ws = openpyxl.load_workbook(dst).worksheets[0]
        self.assertIsNone(ws["O6"].value); self.assertEqual(ws["P6"].value, 75)
        self.assertEqual(ws["P9"].value, "=P8")
        self.assertEqual(ws["T10"].value, "=MAX(O$6:P$6)")
        self.assertEqual(len(ws._charts[0].series), 4)
        self.assertIn("HLF-QC-126-08", ws["A45"].value)


class _Data(object):
    def __init__(self, lots):
        self.domestic = lots
        self.issues = []
        self.coa = {lot: {"924": {"assay": "%.1f" % (100 + i * 0.3), "metal_each": "0", "metal_total": "1",
                                  "particle": "%d" % (50 + i), "particle_spec": "75",
                                  "assays": [{"part": "오플록사신", "value": "%.1f" % (100 + i * 0.3), "lo": "90.0", "hi": "110.0"}]}}
                    for i, lot in enumerate(lots)}


class 전년도_파일이_없어도(unittest.TestCase):
    def test_네_파일을_직접_그려_만든다(self):
        folder = tempfile.mkdtemp(prefix="pqr-cpk-out-")
        data = _Data(["OEY%03d" % i for i in range(1, 13)])
        made = excel_attach.write_cpk_files(folder, data, None, "2026.09.06", product_name="퀴노비드안연고")
        self.assertEqual(sorted(n for n, _ in made),
                         ["a. 금속성이물(개개) Cpk 계산 파일.xlsx", "a. 금속성이물(합계) Cpk 계산 파일.xlsx",
                          "a. 입자도 Cpk 계산 파일.xlsx", "a. 함량 Cpk 계산 파일.xlsx"])
        for _, path in made:
            self.assertTrue(os.path.isfile(path))
        ws = openpyxl.load_workbook(dict(made)["a. 함량 Cpk 계산 파일.xlsx"]).worksheets[0]
        self.assertEqual(ws["C4"].value, "퀴노비드안연고 (내수용)")
        self.assertEqual(ws["B10"].value, 100.0)
        self.assertTrue(any("직접 그렸습니다" in i[2] for i in data.issues))

    def test_10_lot_미만이면_만들지_않는다(self):
        folder = tempfile.mkdtemp(prefix="pqr-cpk-out-")
        data = _Data(["OEY%03d" % i for i in range(1, 5)])
        self.assertEqual(excel_attach.write_cpk_files(folder, data, None, "2026.09.06"), [])


if __name__ == "__main__":
    unittest.main()
