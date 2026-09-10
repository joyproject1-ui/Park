"""수율현황표의 포장이 시장별로 갈린 표 (담당자 2026-09-07~08:
"ELYN01 과 ELYN02 는 내수가 아니고 베트남 포장했네")."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine.readers import yield_sheet as Y                       # noqa: E402


def _sheet(path, merge_spec=False):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    rows = [[None] * 7,
            [None, None, "중요공정 별 수율 현황(%)", None, None, None, None],
            [None, "공정", "조제", "충전", "포장", None, "비고"],
            ["연번", "기준", "99.5±0.48%", "97.0±3.0%", "내수", "베트남(한비돈점안액)", None],
            [None, "Lot No.", None, None, "98.0 ± 2.0%", None, None],
            [1, "ELYO01", 99.97, 97.7, 99.92, None, None],
            [2, "ELYN01", 99.97, 98.52, None, 99.75, None]]
    for r in rows:
        ws.append(r)
    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=6)   # 표 제목이 값 열 전체에 걸친다
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=6)   # '포장' 이 두 열에 걸친다
    if merge_spec:
        ws.merge_cells(start_row=5, start_column=5, end_row=5, end_column=6)
    wb.save(path)
    return path


class 병합된_머리행(unittest.TestCase):
    """머리가 여러 줄이고 칸이 병합된 표 — 표 제목 줄은 이름에서 뺀다."""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-yield-")

    def test_포장이_시장별로_갈리면_이름을_나눠_읽는다(self):
        path = _sheet(os.path.join(self.dir, "y.xlsx"))
        got = dict(Y.read_yields(path))
        self.assertEqual(got["ELYO01"]["포장(내수)"], "99.92")
        self.assertEqual(got["ELYN01"]["포장(베트남)"], "99.75")     # 예전에는 통째로 버려졌다
        self.assertNotIn("포장(내수)", got["ELYN01"])                 # 그 시장으로 포장하지 않았다

    def test_기준이_하나면_두_시장이_나눠_쓴다(self):
        path = _sheet(os.path.join(self.dir, "y2.xlsx"))
        specs = Y.read_specs(path)
        self.assertEqual(specs["포장(내수)"], "98.0 ± 2.0%")
        self.assertEqual(specs["포장(베트남)"], "98.0 ± 2.0%")
        self.assertEqual(specs["조제"], "99.5±0.48%")

    def test_병합된_기준_줄도_읽는다(self):
        path = _sheet(os.path.join(self.dir, "y3.xlsx"), merge_spec=True)
        specs = Y.read_specs(path)
        self.assertEqual(specs["포장(베트남)"], "98.0 ± 2.0%")



class 머리가_두_줄인_표(unittest.TestCase):
    """아이퓨어 수율현황표 — '포장' 아래에 '내수'·'온누리에이치엔씨' 가 따로 있다
    (담당자 2026-09-08: 세 Lot 모두 '포장 수율 값이 수율현황표에 없음')."""

    def _sheet(self, path):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        for r in [[None, "조제", "충전", "포장", "포장"],
                  [None, "조제", "충전", "내수", "온누리에이치엔씨"],
                  [None, "99.5 ± 0.5%", "92.0 ± 8.0%", "98.0 ± 2.0%", "98.0 ± 2.0%"],
                  ["LWY201", 99.97, 93.47, None, 99.12],
                  ["LWY501", 99.93, 94.93, 99.41, None]]:
            ws.append(r)
        wb.save(path)
        return path

    def test_두_줄을_이어_포장_내수로_읽는다(self):
        path = self._sheet(os.path.join(tempfile.mkdtemp(prefix="pqr-y2-"), "y.xlsx"))
        got = dict(Y.read_yields(path))
        self.assertEqual(got["LWY201"]["포장(온누리에이치엔씨)"], "99.12")
        self.assertEqual(got["LWY501"]["포장(내수)"], "99.41")
        self.assertEqual(got["LWY201"]["조제"], "99.97")     # 한 줄뿐인 공정은 그대로
        specs = Y.read_specs(path)
        self.assertEqual(specs["포장(내수)"], "98.0 ± 2.0%")
        self.assertEqual(specs["충전"], "92.0 ± 8.0%")


class 보고서_표의_두_줄_머리(unittest.TestCase):
    """담당자 2026-09-08: "온누리 값을 잘못 넣었어" — 표 머리가 '포장' 아래 '내수·온누리' 두 칸인데
    한 줄만 읽어 열이 셋으로 잡히고 값이 한 칸 밀렸다."""

    def _table(self):
        import docx
        from pqr.engine import docedit as E
        d = docx.Document()
        t = d.add_table(rows=4, cols=7)
        rows = [["연번", "중요공정 별 수율 현황 (%)", "", "", "", "", "비고"],
                ["", "공정", "조제", "충전", "포장", "", ""],
                ["", "기준\nLot No.", "99.5 ± 0.5%", "92.0 ± 8.0%", "내수", "온누리\n에이치엔씨", ""],
                ["1", "LWY201", "", "", "", "", ""]]
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                E.set_cell(t.rows[i].cells[j], v)
        E.merge_right(E.raw_cells(t.rows[1])[4], E.raw_cells(t.rows[1])[5])   # '포장' 이 두 열에 걸친다
        return t

    def test_걸친_머리_칸을_두_열에_모두_편다(self):
        from pqr.engine.recipe_ointment import _yield_columns
        self.assertEqual(_yield_columns(self._table()),
                         [(2, 3, "조제"), (3, 4, "충전"), (4, 5, "포장(내수)"), (5, 6, "포장(온누리에이치엔씨)")])

    def test_시장이_적힌_열은_그_시장_값만_쓴다(self):
        from pqr.engine.recipe_ointment import _yield_value
        vals = {"조제": "99.97", "포장(온누리에이치엔씨)": "99.12"}
        self.assertIsNone(_yield_value(vals, "포장(내수)"))       # 예전에는 99.12 가 내수로 들어갔다
        self.assertEqual(_yield_value(vals, "포장(온누리에이치엔씨)"), "99.12")
        self.assertEqual(_yield_value(vals, "조제"), "99.97")

    def test_시장이_없는_열은_괄호를_뗀_이름으로도_찾는다(self):
        from pqr.engine.recipe_ointment import _yield_value
        self.assertEqual(_yield_value({"포장(내수)": "99.41"}, "포장"), "99.41")

    def test_그_시장으로_포장하지_않은_칸은_사선(self):
        from pqr.engine.recipe_ointment import _yield_columns, _yield_value, _other_market
        cols = _yield_columns(self._table())
        vals = [_yield_value({"포장(온누리에이치엔씨)": "99.12"}, one[-1]) for one in cols]
        self.assertTrue(_other_market(cols, vals, 2))            # 내수 칸 — 사선
        self.assertFalse(_other_market(cols, vals, 3))           # 온누리 칸 — 값이 있다


class 표_제목이_값_열에_걸친_표(unittest.TestCase):
    """담당자 2026-09-08: "수율 작성 안 됐어" — 아이퓨어 7항 표는 제목
    '중요공정 별 수율 현황 (%)' 이 값 열 전체에 가로로 걸쳐 있고 그 줄 끝에 '비고' 가 있다.
    비고까지 함께 보아 제목 줄로 걸러지지 않았고, 열 이름이
    '중요공정별수율현황(%)(조제)' 가 되어 세 Lot 아홉 칸이 모두 '확인 필요' 로 나왔다."""

    def _table(self):
        import docx
        from pqr.engine import docedit as E
        d = docx.Document()
        t = d.add_table(rows=4, cols=7)
        rows = [["연번", "중요공정 별 수율 현황 (%)", "", "", "", "", "비고"],
                ["", "공정", "조제", "충전", "포장", "", ""],
                ["", "기준\nLot No.", "99.5 ± 0.5%", "92.0 ± 8.0%", "내수", "온누리\n에이치엔씨", ""],
                ["1", "LWY201", "", "", "", "", ""]]
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                E.set_cell(t.rows[i].cells[j], v)
        raw = E.raw_cells(t.rows[0])
        for _ in range(4):                                  # 제목이 공정·값 열 전체에 걸친다
            raw = E.raw_cells(t.rows[0])
            E.merge_right(raw[1], raw[2])
        E.merge_right(E.raw_cells(t.rows[1])[4], E.raw_cells(t.rows[1])[5])
        return t

    def test_제목_줄을_이름으로_쓰지_않는다(self):
        from pqr.engine.recipe_ointment import _yield_columns
        self.assertEqual(_yield_columns(self._table()),
                         [(2, 3, "조제"), (3, 4, "충전"), (4, 5, "포장(내수)"), (5, 6, "포장(온누리에이치엔씨)")])

    def test_수율현황표_이름과_그대로_맞는다(self):
        from pqr.engine.recipe_ointment import _yield_columns, _yield_value
        vals = {"조제": "99.97", "충전": "93.47", "포장(온누리 에이치엔씨)": "99.12"}
        got = [_yield_value(vals, one[-1]) for one in _yield_columns(self._table())]
        self.assertEqual(got, ["99.97", "93.47", None, "99.12"])


class 엑셀_제목_줄도_이름이_아니다(unittest.TestCase):
    """수율현황표 엑셀에서도 제목이 값 열에 걸쳐 있으면 이름에 섞이면 안 된다."""

    def test_제목이_걸쳐도_공정_이름만_읽는다(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        for r in [[None, "중요공정 별 수율 현황 (%)", None, None, None, "비고"],
                  ["공정", "조제", "충전", "포장", "포장", None],
                  [None, None, None, "내수", "온누리에이치엔씨", None],
                  ["Lot No.", "99.5 ± 0.5%", "92.0 ± 8.0%", "98.0 ± 2.0%", None, None],
                  ["LWY201", 99.97, 93.47, None, 99.12, None]]:
            ws.append(r)
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=5)
        ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=5)
        path = os.path.join(tempfile.mkdtemp(prefix="pqr-y3-"), "y.xlsx")
        wb.save(path)
        got = dict(Y.read_yields(path))
        self.assertEqual(got["LWY201"]["조제"], "99.97")
        self.assertEqual(got["LWY201"]["포장(온누리에이치엔씨)"], "99.12")


class 아이퓨어_수율현황표_그대로(unittest.TestCase):
    """담당자가 보여 준 실제 파일(7. 수율현황표 - 개인.xlsx) 짜임 그대로 —
    '조제'·'충전' 은 두 줄에 걸쳐 병합, '포장' 은 내수·온누리 두 열에 걸쳐 병합,
    기준 줄의 '98.0 ± 2.0%' 도 두 열에 걸쳐 있다."""

    def _sheet(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        for r in [[None, "조제", "충전", "포장", None],
                  [None, None, None, "내수", "온누리 에이치엔씨"],
                  [None, "99.5 ± 0.5%", "92.0 ± 8.0%", "98.0 ± 2.0%", None],
                  ["LWY201", 99.97, 93.47, None, 99.12],
                  ["LWY501", 99.93, 94.93, 99.41, None],
                  ["LWYO01", 99.88, 94.3, 98.65, None]]:
            ws.append(r)
        for spot in ("A1:A3", "B1:B2", "C1:C2", "D1:E1", "D3:E3"):
            ws.merge_cells(spot)
        path = os.path.join(tempfile.mkdtemp(prefix="pqr-ipure-"), "7. 수율현황표.xlsx")
        wb.save(path)
        return path

    def test_세_Lot_아홉_칸을_모두_읽는다(self):
        got = dict(Y.read_yields(self._sheet()))
        self.assertEqual(got["LWY201"]["포장(온누리 에이치엔씨)"], "99.12")
        self.assertEqual(got["LWY501"]["포장(내수)"], "99.41")
        self.assertEqual(got["LWYO01"]["충전"], "94.30")
        self.assertNotIn("포장(내수)", got["LWY201"])          # 그 시장으로 포장하지 않았다 — 사선

    def test_보고서_표_열_이름과_이어진다(self):
        """엑셀은 '온누리 에이치엔씨', 보고서 표는 줄바꿈이 붙은 '온누리에이치엔씨' — 빈칸은 무시한다."""
        from pqr.engine.recipe_ointment import _yield_value
        vals = dict(Y.read_yields(self._sheet()))["LWY501"]
        self.assertEqual(_yield_value(vals, "포장(내수)"), "99.41")
        self.assertIsNone(_yield_value(vals, "포장(온누리에이치엔씨)"))
        self.assertEqual(_yield_value(vals, "조제"), "99.93")


class 머리_칸이_두_열에_걸친_공양식(unittest.TestCase):

    def assertSame(self, a, b, msg=None):
        """python-docx 는 부를 때마다 새 칸 객체를 만든다 — 속 알맹이로 견준다."""
        self.assertIs(getattr(a, "_tc", a), getattr(b, "_tc", b), msg)

    """담당자 2026-09-08: 아이퓨어 공양식 7항은 '온누리 에이치엔씨' 머리 칸이 두 열에 걸쳐 있고,
    자료 줄은 '내수' 가 두 열, '온누리' 가 한 열로 갈려 머리와 경계가 다르다.
    열 번호로 짚으면 온누리 값(99.12)이 통째로 버려진다."""

    def _table(self):
        import docx
        from pqr.engine import docedit as E
        d = docx.Document()
        t = d.add_table(rows=6, cols=8)
        rows = [["연번", "중요공정 별 수율 현황 (%)", "", "", "", "", "", "비고"],
                ["", "공정", "조제", "충전", "포장", "", "", ""],
                ["", "", "", "", "내수", "온누리\n에이치엔씨", "", ""],
                ["", "기준\nLot No.", "99.5 ± 0.5%", "92.0 ± 8.0%", "98.0 ± 2.0%", "", "", ""],
                ["1", "", "", "", "", "", "", ""],
                ["최댓값", "", "", "", "", "", "", ""]]
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                E.set_cell(t.rows[i].cells[j], v)
        for _ in range(5):                                   # 0줄: 제목이 1~6열에 걸친다
            raw = E.raw_cells(t.rows[0]); E.merge_right(raw[1], raw[2])
        for _ in range(2):                                   # 1줄: '포장' 이 4~6열
            raw = E.raw_cells(t.rows[1]); E.merge_right(raw[4], raw[5])
        raw = E.raw_cells(t.rows[2]); E.merge_right(raw[5], raw[6])   # '온누리' 가 5~6열
        for _ in range(2):                                   # 3줄: 기준이 4~6열
            raw = E.raw_cells(t.rows[3]); E.merge_right(raw[4], raw[5])
        raw = E.raw_cells(t.rows[4]); E.merge_right(raw[4], raw[5])   # 자료 줄은 4~5열 / 6열
        raw = E.raw_cells(t.rows[5]); E.merge_right(raw[0], raw[1])   # 요약 줄은 '최댓값' 이 0~1열
        return t

    def test_걸친_머리_칸은_한_공정이다(self):
        from pqr.engine.recipe_ointment import _yield_columns
        self.assertEqual(_yield_columns(self._table()),
                         [(2, 3, "조제"), (3, 4, "충전"), (4, 5, "포장(내수)"), (5, 7, "포장(온누리에이치엔씨)")])

    def test_자료_줄의_칸을_겹침으로_짚는다(self):
        from pqr.engine.recipe_ointment import _yield_columns, _cells_for
        from pqr.engine import docedit as E
        t = self._table()
        cols = _yield_columns(t)
        값칸, 비고칸 = _cells_for(t.rows[4], cols)
        raw = E.raw_cells(t.rows[4])
        self.assertSame(값칸[0], raw[2])                 # 조제
        self.assertSame(값칸[1], raw[3])                 # 충전
        self.assertSame(값칸[2], raw[4])                 # 내수 — 4~5열에 걸친 칸
        self.assertSame(값칸[3], raw[5])                 # 온누리 — 6열 칸 (예전에는 못 찾았다)
        self.assertSame(비고칸, raw[6])                   # 비고는 맨 끝 칸 (예전에는 r[5]=온누리를 지웠다)

    def test_요약_줄도_같은_자리에_쓴다(self):
        from pqr.engine.recipe_ointment import _yield_columns, _cells_for
        from pqr.engine import docedit as E
        t = self._table()
        값칸 = _cells_for(t.rows[5], _yield_columns(t))[0]
        raw = E.raw_cells(t.rows[5])
        self.assertSame(값칸[0], raw[1])                 # '최댓값' 이 0~1열을 덮어 한 칸씩 밀린다
        self.assertSame(값칸[3], raw[4])


class 수율_NA(unittest.TestCase):
    """수율현황표의 'N/A' 는 값이 아니라 '그 공정(시장)을 거치지 않음' — 사선 (퀴노비드점안액 2026-09-10:
    float('N/A') 로 작성 전체가 멈췄다)."""

    def test_NA_글_판별(self):
        from pqr.engine import recipe_ointment as R
        for t in ("N/A", "n/a", " NA ", "-", "해당없음", "해당 없음"):
            self.assertTrue(R.NA_TEXT.match(t), t)
        for t in ("99.92", "0", "확인 중"):
            self.assertFalse(R.NA_TEXT.match(t), t)


if __name__ == "__main__":
    unittest.main()
