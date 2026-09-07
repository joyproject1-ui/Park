"""6항 공 기록서(제조·충전·포장 워드) — 제조단위·포장단위·수율 기준·원/자재 코드 읽기와 6·7·8.1항 반영 (2026-09-06).

담당자: "제조단위와 포장단위, 수율 기준, 주원료(원료코드가 R 로 시작하면 주원료 / 그 외는 부원료), 부원료,
포장자재 정보 등은 여기 업로드된 공 기록서를 참고하면 돼. 모든 PQR 작성할 때 참고해."
"""
import os
import sys
import tempfile
import unittest

import docx

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import docedit as E                                  # noqa: E402
from pqr.engine.readers import batch_record as B                     # noqa: E402
from pqr.engine.recipe_ointment import _add_material_rows            # noqa: E402


def _save(folder, name, tables, paras=()):
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    for rows in tables:
        t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                t.rows[i].cells[j].text = v
    path = os.path.join(folder, name)
    d.save(path)
    return path


class 공기록서_읽기(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="pqr-batch-")

    def test_제조기록서_제조단위_수율_원료(self):
        p = _save(self.dir, "6. EHLF-21 MO-QUIO3(Rev026) 퀴노비드안연고 제조기록서.docx",
                  [[["제품명", "퀴노비드안연고", "제조단위", "204,000 g (40,800 Tube)"]],
                   [["원료코드", "원료명", "규격", "배합량"], ["RBO101", "오플록사신", "USP", "612 g"],
                    ["EPP116", "정제라놀린", "KP", "10,200 g"]],
                   [["조제 수율", "기준 : 95.0% 이상", "실측", ""]]])
        got = B.read(p)
        self.assertEqual(got["kind"], "조제")
        self.assertEqual(got["batch_size"], "204,000g")
        self.assertEqual(got["yield_spec"], "95.0% 이상")
        self.assertEqual([(m["code"], m["group"], m["spec"]) for m in got["materials"]],
                         [("RBO101", "주원료", "USP"), ("EPP116", "부원료", "KP")])

    def test_충전_포장기록서와_합치기(self):
        a = _save(self.dir, "6. FI-QUIO3 충전기록서.docx", [[["충전 수율 (%)", "91.0 ± 4.0%", ""]]])
        b = _save(self.dir, "6. PK-QUIBTH 포장기록서.docx",
                  [[["포장단위", "5g x 1Tube/Case"], ["포장 수율", "98 ± 2%"]],
                   [["코드", "품명"], ["P17039", "AL-Tube(내수)"]]],
                  paras=["제조단위 : 204,000 g"])
        got = B.merge([B.read(a), B.read(b)])
        self.assertEqual(got["pack_unit"], "5g x 1Tube/Case")
        self.assertEqual(got["batch_size"], "204,000g")               # 표 밖 글줄에서도 읽는다
        self.assertEqual(got["yield_specs"], {"충전": "91.0 ± 4.0%", "포장": "98 ± 2%"})
        self.assertEqual([m["code"] for m in got["materials"]["포장자재"]], ["P17039"])

    def test_머리행_없는_표에서는_코드_옆_글자를_이름으로(self):
        p = _save(self.dir, "6. 제조기록서.docx", [[["1", "ELL101", "유동파라핀", "20,400 g"]]])
        got = B.read(p)
        self.assertEqual([(m["code"], m["name"], m["group"]) for m in got["materials"]], [("ELL101", "유동파라핀", "부원료")])


class 표에_보태기(unittest.TestCase):
    HEAD = ["연번", "관리번호", "원/자재명", "규격", "제조원"]

    def _table(self, rows):
        d = docx.Document()
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                E.set_cell(t.rows[i].cells[j], v)
        return t

    def test_없는_코드만_줄로_보태고_이름은_노랑(self):
        t = self._table([self.HEAD, ["1", "EPP116", "정제라놀린", "KP", "Nippon"], ["2", "P17039", "튜브", "자사규격", "린하르트"]])
        n = _add_material_rows(t, [{"code": "P17039", "name": "AL-Tube", "spec": ""},
                                   {"code": "P20111", "name": "갑(내수)", "spec": ""}])
        self.assertEqual(n, 1)
        rows = [[E.cell_text(c) for c in E.raw_cells(r)] for r in t.rows]
        self.assertEqual(rows[2][1:3], ["P17039", "튜브"])                 # 있던 줄은 그대로
        self.assertEqual(rows[3][:4], ["3", "P20111", "갑(내수)", "자사규격"])  # P 코드 규격은 자사규격
        self.assertEqual(rows[3][4], "")

    def test_빈_공양식이면_그_줄들이_표가_된다(self):
        t = self._table([self.HEAD, ["1", "", "", "", ""]])
        n = _add_material_rows(t, [{"code": "RBO101", "name": "오플록사신", "spec": "USP"}])
        self.assertEqual(n, 1)
        rows = [[E.cell_text(c) for c in E.raw_cells(r)] for r in t.rows]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:4], ["1", "RBO101", "오플록사신", "USP"])


    def test_한_자_다른_관리번호는_오기로_보고_고친다(self):
        # 담당자 2026-09-07: "케이스는 내수 케이스야, 기존 P38033 관리번호 오기라서 P38003 으로 수정해 줘"
        t = self._table([self.HEAD, ["1", "P38033", "케이스(내수)", "자사규격", "한신인쇄"],
                         ["2", "P33603", "케이스(수출)", "자사규격", "한신인쇄"]])
        오기 = []
        n = _add_material_rows(t, [{"code": "P38003", "name": "케 이 스", "spec": ""}], 오기)
        self.assertEqual((n, 오기), (0, [("P38033", "P38003")]))
        rows = [[E.cell_text(c) for c in E.raw_cells(r)] for r in t.rows]
        self.assertEqual(len(rows), 3)                                  # 줄을 보태지 않는다
        self.assertEqual(rows[1][1:3], ["P38003", "케이스(내수)"])
        self.assertEqual(rows[2][1:3], ["P33603", "케이스(수출)"])


class 자재가_아닌_것(unittest.TestCase):
    def test_저울과_하조용_상자는_자재로_보지_않는다(self):
        self.assertFalse(B.is_material("FAB5069", "전자저울"))
        self.assertFalse(B.is_material("P34732", "하조용 종이상자 732호"))
        self.assertTrue(B.is_material("P38003", "케이스"))
        self.assertTrue(B.is_material("RBO101", "오플록사신"))


if __name__ == "__main__":
    unittest.main()
