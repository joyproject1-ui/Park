"""빈 쪽·쪽 맨 위 빈 줄, 14.2 한 줄, 16항 내어쓰기, 판독 파일에 없는 시험일지 읽기 (담당자 2026-09-06)."""
import os
import sys
import unittest

import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pqr.engine import docedit as E, collect                       # noqa: E402


def _break_para(d):
    p = d.add_paragraph()
    r = p.add_run()
    r._r.append(r._r.makeelement(qn("w:br"), {qn("w:type"): "page"}))
    return p


def _pbb(el):
    pr = el.find(qn("w:pPr"))
    return pr is not None and pr.find(qn("w:pageBreakBefore")) is not None


class 쪽나눔_정리(unittest.TestCase):
    def test_쪽나눔_뒤_빈_문단을_건너뛰어_제목에_앞에서_쪽나눔(self):
        d = docx.Document()
        d.add_paragraph("각주 1)")
        _break_para(d); d.add_paragraph(""); d.add_paragraph("")
        d.add_paragraph("9.2.2.3 포장 완료 후")
        n = E.hard_breaks_to_page_break_before(d)
        self.assertEqual(n, 1)
        texts = [p.text for p in d.paragraphs]
        self.assertEqual(texts, ["각주 1)", "9.2.2.3 포장 완료 후"])
        self.assertTrue(_pbb(d.paragraphs[1]._p))

    def test_표_앞_쪽나눔은_표_첫_문단의_앞에서_쪽나눔으로(self):
        d = docx.Document()
        d.add_paragraph("앞 글")
        _break_para(d); d.add_paragraph("")
        t = d.add_table(rows=1, cols=1); t.rows[0].cells[0].text = "표"
        n = E.hard_breaks_to_page_break_before(d)
        self.assertEqual(n, 1)
        self.assertEqual([p.text for p in d.paragraphs], ["앞 글"])
        self.assertTrue(_pbb(t._tbl.find(".//" + qn("w:p"))))


class 불만_없음_한_줄(unittest.TestCase):
    def test_빈_줄_셋을_한_줄로(self):
        d = docx.Document()
        t = d.add_table(rows=5, cols=3)
        for j, h in enumerate(("연번", "Lot No.", "문서번호")):
            t.rows[0].cells[j].text = h
        t.rows[1].cells[0].text = "1"
        for j in (0, 1):
            E.set_vmerge(t.rows[1].cells[j], "restart"); E.set_vmerge(t.rows[2].cells[j], None); E.set_vmerge(t.rows[3].cells[j], None)
        t.rows[4].cells[0].text = "특이사항 (Comment)"
        self.assertEqual(E.single_blank_row(t), 2)
        self.assertEqual(len(t.rows), 3)
        # 내역이 없으면 연번도 비운다 (담당자 2026-09-07: "불만이 없는데 연번에 1 표시됨")
        self.assertEqual(E.cell_text(E.raw_cells(t.rows[1])[0]), "")
        self.assertIsNone(E.raw_cells(t.rows[1])[0]._tc.find(qn("w:tcPr")).find(qn("w:vMerge")))

    def test_글이_있는_줄은_줄이지_않는다(self):
        d = docx.Document()
        t = d.add_table(rows=4, cols=2)
        t.rows[0].cells[0].text = "연번"; t.rows[1].cells[1].text = "불만 1건"; t.rows[3].cells[0].text = "특이사항 (Comment)"
        self.assertEqual(E.single_blank_row(t), 0)
        self.assertEqual(len(t.rows), 4)


class 결론_내어쓰기(unittest.TestCase):
    def test_번호_문단은_번호_폭만큼_내어쓴다(self):
        d = docx.Document()
        d.add_paragraph("16. 결론")
        d.add_paragraph("16.1 본 제품 품질 평가를 통해")
        d.add_paragraph("보통 문단")
        d.add_paragraph("17. 참고")
        self.assertEqual(E.set_section_indent(d, 16, 2), 2)
        ind = d.paragraphs[1]._p.find(qn("w:pPr")).find(qn("w:ind"))
        # 왼쪽 들여쓰기는 2칸 그대로이고 번호만 그만큼 내어쓴다 — 첫 줄이 여백 밖으로 나가지 않게
        # (담당자 2026-09-07: "16.1~16.4 는 단락 들여쓰기 왼쪽 2글자로 변경해 줘")
        self.assertEqual(ind.get(qn("w:leftChars")), "200")
        self.assertEqual(ind.get(qn("w:hangingChars")), "200")
        ind = d.paragraphs[2]._p.find(qn("w:pPr")).find(qn("w:ind"))
        self.assertEqual(ind.get(qn("w:leftChars")), "200")
        self.assertIsNone(ind.get(qn("w:hangingChars")))


class 판독_파일에_없는_일지(unittest.TestCase):
    def test_source_가_없는_일지만_새로_읽는다(self):
        logs = [{"lot": "OEV301", "source": "13 시판후 안정성시험일지.pdf"}]
        scanned = ["/x/13 시판후 안정성시험일지.pdf", "/x/13 장기 안정성시험일지(내수용).pdf"]
        self.assertEqual(collect.unread_scans(logs, scanned), ["/x/13 장기 안정성시험일지(내수용).pdf"])

    def test_어느_일지를_읽었는지_모르는_옛_판독_파일은_그대로(self):
        self.assertEqual(collect.unread_scans([{"lot": "OEV301"}], ["/x/a.pdf"]), [])


if __name__ == "__main__":
    unittest.main()


class 제목_뒤_쪽_나눔(unittest.TestCase):
    """담당자 2026-09-07: "16항 결론은 왜 다음 페이지로 밀렸는지? 21페이지에 작성되도록 해"."""

    def _doc(self):
        import docx
        from docx.oxml.ns import qn
        from pqr.engine.docedit import get_or_add
        d = docx.Document()
        d.add_paragraph("16. 결론")
        d.add_paragraph("")
        body = d.add_paragraph("한림포비돈점안액에 대한 제품품질평가 결과 …")
        get_or_add(body._p.get_or_add_pPr(), "pageBreakBefore")
        d.add_paragraph("17. 참고 자료")
        other = d.add_paragraph("- 제품표준서")
        d.add_paragraph("아무 글")
        late = d.add_paragraph("떨어진 문단")
        get_or_add(late._p.get_or_add_pPr(), "pageBreakBefore")
        return d, body, late

    def test_제목_바로_뒤_본문의_쪽_나눔만_뗀다(self):
        from docx.oxml.ns import qn
        from pqr.engine import docedit as E
        d, body, late = self._doc()
        self.assertEqual(E.drop_break_after_headings(d), 1)
        self.assertIsNone(body._p.pPr.find(qn("w:pageBreakBefore")))
        self.assertIsNotNone(late._p.pPr.find(qn("w:pageBreakBefore")))   # 제목 뒤가 아닌 것은 그대로


class 머리글전체쪽수(unittest.TestCase):
    """목차 쪽 번호와 머리글 'Page n / 전체' 는 같은 PDF 쪽수에서 (담당자 2026-09-10: 목차 22쪽, 머리글 '1 / 21')."""

    def test_NUMPAGES_캐시를_적는다(self):
        import re as _re
        from pqr.engine import toc as T
        xml = ('<w:r><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>'
               '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>#</w:t></w:r>'
               '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
        out, k = T.NUMPAGES_CACHE.subn(lambda m: m.group(1) + "22" + m.group(3), xml)
        self.assertEqual(k, 1)
        self.assertIn("<w:t>22</w:t>", out)
        self.assertNotIn("#", _re.sub(r"<[^>]+>", "", out))


class 목차_쪽수_출처(unittest.TestCase):
    """LibreOffice 로 센 쪽수를 적을 때는 dirty 표시를 남긴다 — Word 로 열면 제 쪽 나눔으로 다시 센다.

    담당자 2026-09-11: 목차 33쪽인데 머리글은 '1 / 32'. 머리글 NUMPAGES 만 Word 가 다시 세고
    목차는 LibreOffice 값에 굳어 둘이 어긋났다 (올로원스 25/23 · 나조린 22/21 도 같은 원인).
    """

    def _문서(self, path):
        import docx
        from docx.oxml.ns import qn as _qn
        from pqr.engine import toc as T
        doc = docx.Document()
        head = doc.add_paragraph("1. 목 적")
        head.style = doc.styles["Heading 1"]
        T._bookmark(head._p, "_pqr_toc_1")
        para = doc.add_paragraph()
        run = para.add_run("#")._r
        for el in T._field_runs(run, "PAGEREF _pqr_toc_1 \\h", "#"):
            run.addprevious(el)
        run.getparent().remove(run)
        doc.save(path)
        return path

    def _세어적기(self, how):
        import tempfile, os as _os
        from pqr.engine import toc as T, convert
        path = self._문서(_os.path.join(tempfile.mkdtemp(), "a.docx"))
        pdf_to, pages = convert.to_pdf, T._pdf_pages
        def _가짜pdf(src, dst):
            with open(dst, "wb") as handle:
                handle.write(b"%PDF-1.4")
            return how
        convert.to_pdf = _가짜pdf
        T._pdf_pages = lambda p: ["", "1.목적"]
        try:
            T.fill_page_numbers(path)
        finally:
            convert.to_pdf, T._pdf_pages = pdf_to, pages
        import zipfile
        with zipfile.ZipFile(path) as z:
            return z.read("word/document.xml").decode("utf-8")

    def test_LibreOffice_쪽수면_dirty_를_남긴다(self):
        xml = self._세어적기("soffice")
        self.assertIn("<w:t>2</w:t>", xml)          # 어림수는 적어 둔다(제한된 보기용)
        self.assertIn('w:dirty="true"', xml)        # Word 로 열면 다시 센다

    def test_Word_쪽수면_dirty_를_뗀다(self):
        xml = self._세어적기("word")
        self.assertIn("<w:t>2</w:t>", xml)
        self.assertNotIn('w:dirty="true"', xml)
